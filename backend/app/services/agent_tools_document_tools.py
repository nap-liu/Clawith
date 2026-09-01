"""Document readers and conversion tools extracted from agent_tools."""

from __future__ import annotations

import asyncio
import multiprocessing as mp
import queue
import re
import uuid
from pathlib import Path
from typing import Any

from loguru import logger

from app.services.agent_tools_file_support import (
    _exact_storage_source_error,
    _resolve_exact_storage_source_path,
    _resolve_tool_source_path,
    _resolve_tool_target_path,
)
from app.services.document_conversion import (
    convert_html_to_pdf as convert_html_file_to_pdf,
    convert_html_to_pptx as convert_html_file_to_pptx,
)
from app.services.storage import get_storage_backend

_READ_DOCUMENT_MAX_FILE_BYTES = 50 * 1024 * 1024
_READ_DOCUMENT_TIMEOUT_SECONDS = 25
_READ_DOCUMENT_FALLBACK_TIMEOUT_SECONDS = 10
_READ_DOCUMENT_MAX_COLUMNS = 80
_READ_DOCUMENT_MAX_ROWS = 10_000
_READ_DOCUMENT_MAX_SHEETS = 10
_READ_DOCUMENT_MAX_PAGES = 200
_READ_DOCUMENT_MAX_SLIDES = 200
_READ_DOCUMENT_HARD_CHAR_CEILING = 2_000_000


class _LazyPrepareTempWorkspace:
    async def __call__(self, *args, **kwargs):
        from app.services.agent_tools_temp_workspace import _prepare_temp_workspace

        return await _prepare_temp_workspace(*args, **kwargs)


_prepare_temp_workspace = _LazyPrepareTempWorkspace()


def _safe_document_cell_text(value: Any) -> str:
    """Stringify a spreadsheet/table cell faithfully — NO content-length truncation.

    read_document must return cell content in full so embedded SQL / long口径 text is
    never silently chopped mid-statement (which is unrecoverable for the agent). Output
    that is genuinely too large is handled non-lossily by the unified overflow-to-file
    path (llm.tool_output_store.finalize_tool_output), and total memory stays bounded by
    the file-size gate + per-sheet row/column caps + the 2M _READ_DOCUMENT_HARD_CHAR_CEILING.

    The only guard kept here is the degenerate huge-integer case: str() on a multi-thousand
    digit int is pathologically slow and can raise under CPython's int_max_str_digits.
    """
    if value is None:
        return ""
    if isinstance(value, int) and value.bit_length() > 4096:
        return "[large integer omitted]"
    return str(value)


def _render_xlsx_row(values) -> str:
    """Render one spreadsheet row to a tab-joined string, trimming TRAILING empty
    cells.

    openpyxl ``iter_rows(max_col=_READ_DOCUMENT_MAX_COLUMNS)`` pads every row out
    to the column cap, so a sparse row (e.g. one value in a far-right SQL column)
    would otherwise emit dozens of trailing tabs. That padding burned the
    ``max_chars`` budget and pushed real content past the truncation line. Inner
    empty cells are kept so a populated cell stays aligned with its header column.
    """
    texts = [_safe_document_cell_text(c) for c in values]
    while texts and not texts[-1].strip():
        texts.pop()
    return "\t".join(texts)


def _read_document_sync(
    ws: Path, rel_path: str, max_chars: int = _READ_DOCUMENT_HARD_CHAR_CEILING, tenant_id: str | None = None
) -> str:
    """Synchronous document extraction. Must run outside the uvicorn event loop."""
    max_chars = min(max(int(max_chars), 1), _READ_DOCUMENT_HARD_CHAR_CEILING)
    try:
        file_path = _resolve_tool_source_path(ws, rel_path, tenant_id=tenant_id)
    except ValueError as exc:
        return str(exc)

    if not file_path.exists():
        return f"File not found: {rel_path}"
    if file_path.is_dir():
        return f"Path is a directory, not a document: {rel_path}"
    try:
        file_size = file_path.stat().st_size
    except OSError:
        file_size = 0
    if file_size > _READ_DOCUMENT_MAX_FILE_BYTES:
        return (
            f"Document is too large to read safely ({file_size / 1024 / 1024:.1f} MB). "
            "Please split or convert it to a smaller text/Markdown excerpt first."
        )

    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber

            text_parts = []
            with pdfplumber.open(str(file_path)) as pdf:
                total_pages = len(pdf.pages)
                for i, page in enumerate(pdf.pages[:_READ_DOCUMENT_MAX_PAGES]):
                    page_text = page.extract_text() or ""
                    if page_text:
                        text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
                    if sum(len(part) for part in text_parts) >= max_chars:
                        break
                if total_pages > _READ_DOCUMENT_MAX_PAGES:
                    text_parts.append(
                        f"[PDF has {total_pages} pages; only the first {_READ_DOCUMENT_MAX_PAGES} were extracted]"
                    )
            content = "\n\n".join(text_parts) if text_parts else "(PDF is empty or text extraction failed)"

        elif ext == ".docx":
            from docx import Document
            from docx.oxml.ns import qn

            doc = Document(str(file_path))
            lines: list[str] = []

            def _extract_para_text(para) -> str:
                return para.text.strip()

            def _extract_table(table) -> str:
                """Flatten a table into readable text."""
                rows = []
                for row in table.rows:
                    cells = [
                        _safe_document_cell_text(cell.text).strip() for cell in row.cells[:_READ_DOCUMENT_MAX_COLUMNS]
                    ]
                    if not cells:
                        continue
                    # Remove duplicate adjacent cells (merged cells repeat)
                    deduped = [cells[0]] + [c for i, c in enumerate(cells[1:]) if c != cells[i]]
                    row_str = " | ".join(c for c in deduped if c)
                    if row_str:
                        rows.append(row_str)
                return "\n".join(rows)

            # 1. Main paragraphs
            for para in doc.paragraphs:
                t = _extract_para_text(para)
                if t:
                    lines.append(t)

            # 2. Tables in main body
            for table in doc.tables:
                t = _extract_table(table)
                if t:
                    lines.append(t)

            # 3. Text boxes / drawing shapes (wmf/shapes in body XML)
            for shape in doc.element.body.iter(qn("w:txbxContent")):
                for child in shape.iter(qn("w:t")):
                    if child.text and child.text.strip():
                        lines.append(child.text.strip())

            # 4. Headers and footers
            for section in doc.sections:
                for hf in [section.header, section.footer]:
                    if hf and hf.is_linked_to_previous is False:
                        for para in hf.paragraphs:
                            t = para.text.strip()
                            if t:
                                lines.append(t)

            content = "\n".join(lines) if lines else "(Document is empty or uses unsupported formatting)"

        elif ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(str(file_path), read_only=True, data_only=True)
            sheets = []
            total_chars = 0  # bounds peak memory; the cross-sheet content budget
            all_sheet_names = wb.sheetnames
            stop = False
            for ws_name in all_sheet_names[:_READ_DOCUMENT_MAX_SHEETS]:
                if stop:
                    break
                sheet = wb[ws_name]
                # Declared dimension (read_only): used only to SIGNPOST when a sheet has
                # more rows than the cap — iter_rows(max_row=...) bounds the actual scan.
                declared_rows = sheet.max_row
                rows = []
                for row in sheet.iter_rows(
                    max_row=_READ_DOCUMENT_MAX_ROWS, max_col=_READ_DOCUMENT_MAX_COLUMNS, values_only=True
                ):
                    row_str = _render_xlsx_row(row)
                    if not row_str.strip():
                        continue
                    rows.append(row_str)
                    total_chars += len(row_str) + 1
                    if total_chars >= max_chars:
                        rows.append(
                            f"[content limit (~{max_chars} chars) reached; remaining cells omitted — "
                            "oversized output is materialized to a .tool_results/ file you can read_file]"
                        )
                        stop = True
                        break
                if not stop and declared_rows and declared_rows > _READ_DOCUMENT_MAX_ROWS:
                    rows.append(
                        f"[sheet '{ws_name}' has {declared_rows} rows; only the first "
                        f"{_READ_DOCUMENT_MAX_ROWS} are shown — narrow the columns or query with sql_execute]"
                    )
                if rows:
                    sheets.append(f"=== Sheet: {ws_name} ===\n" + "\n".join(rows))
            if len(all_sheet_names) > _READ_DOCUMENT_MAX_SHEETS:
                sheets.append(
                    f"[workbook has {len(all_sheet_names)} sheets; only the first "
                    f"{_READ_DOCUMENT_MAX_SHEETS} are shown]"
                )
            wb.close()
            content = "\n\n".join(sheets) if sheets else "(Excel is empty)"

        elif ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(file_path))
            slides = []
            all_slides = list(prs.slides)
            for i, slide in enumerate(all_slides[:_READ_DOCUMENT_MAX_SLIDES]):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
                if texts:
                    slides.append(f"--- Slide {i + 1} ---\n" + "\n".join(texts))
            if len(all_slides) > _READ_DOCUMENT_MAX_SLIDES:
                slides.append(
                    f"[presentation has {len(all_slides)} slides; only the first {_READ_DOCUMENT_MAX_SLIDES} are shown]"
                )
            content = "\n\n".join(slides) if slides else "(PPT is empty)"

        elif ext in (".txt", ".md", ".json", ".csv", ".log"):
            content = file_path.read_text(encoding="utf-8", errors="replace")

        else:
            return f"Unsupported file format: {ext}. Supported: PDF, DOCX, XLSX, PPTX, TXT, MD, CSV"

        if len(content) > max_chars:
            content = content[:max_chars] + f"\n\n...[truncated, {len(content)} chars total]"
        return content

    except ImportError as e:
        return f"Missing dependency: {e}. Install: pip install pdfplumber python-docx openpyxl python-pptx"
    except Exception as e:
        return f"Document read failed: {str(e)[:200]}"


def _read_document_worker(
    out_queue: mp.Queue,
    ws_str: str,
    rel_path: str,
    max_chars: int,
    tenant_id: str | None,
) -> None:
    try:
        out_queue.put(("ok", _read_document_sync(Path(ws_str), rel_path, max_chars=max_chars, tenant_id=tenant_id)))
    except BaseException as exc:
        out_queue.put(("error", f"Document read failed: {str(exc)[:200]}"))


def _read_pdf_fast_sync(ws: Path, rel_path: str, max_chars: int = 8000, tenant_id: str | None = None) -> str:
    """Fast PDF text extraction fallback for files that make pdfplumber/pdfminer hang."""
    max_chars = min(max(int(max_chars), 1), 20000)
    try:
        file_path = _resolve_tool_source_path(ws, rel_path, tenant_id=tenant_id)
    except ValueError as exc:
        return str(exc)

    if not file_path.exists():
        return f"File not found: {rel_path}"
    if file_path.is_dir():
        return f"Path is a directory, not a document: {rel_path}"

    try:
        import fitz

        text_parts = []
        with fitz.open(str(file_path)) as doc:
            for i, page in enumerate(doc[:50]):
                page_text = page.get_text("text") or ""
                if page_text:
                    text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
                if sum(len(part) for part in text_parts) >= max_chars:
                    break
        content = "\n\n".join(text_parts) if text_parts else "(PDF is empty or text extraction failed)"
        if len(content) > max_chars:
            content = content[:max_chars] + f"\n\n...[truncated, {len(content)} chars total]"
        return content
    except ImportError as exc:
        return f"PDF fallback extractor unavailable: {exc}. Install: pip install PyMuPDF"
    except Exception as exc:
        return f"PDF fallback extraction failed: {str(exc)[:200]}"


def _read_pdf_fast_worker(
    out_queue: mp.Queue,
    ws_str: str,
    rel_path: str,
    max_chars: int,
    tenant_id: str | None,
) -> None:
    try:
        out_queue.put(("ok", _read_pdf_fast_sync(Path(ws_str), rel_path, max_chars=max_chars, tenant_id=tenant_id)))
    except BaseException as exc:
        out_queue.put(("error", f"PDF fallback extraction failed: {str(exc)[:200]}"))


def _read_pdf_fast_with_timeout(ws: Path, rel_path: str, max_chars: int = 8000, tenant_id: str | None = None) -> str:
    ctx = mp.get_context("spawn")
    out_queue: mp.Queue = ctx.Queue(maxsize=1)
    proc = ctx.Process(
        target=_read_pdf_fast_worker,
        args=(out_queue, str(ws), rel_path, max_chars, tenant_id),
        daemon=True,
    )
    proc.start()
    proc.join(_READ_DOCUMENT_FALLBACK_TIMEOUT_SECONDS)
    if proc.is_alive():
        proc.terminate()
        proc.join(2)
        if proc.is_alive():
            proc.kill()
            proc.join(1)
        return (
            f"Document read timed out after {_READ_DOCUMENT_TIMEOUT_SECONDS}s, "
            f"and PDF fallback also timed out after {_READ_DOCUMENT_FALLBACK_TIMEOUT_SECONDS}s. "
            "The file may be too large or too complex to extract safely."
        )
    try:
        status, payload = out_queue.get_nowait()
    except queue.Empty:
        if proc.exitcode:
            return f"PDF fallback extraction failed: extractor exited with code {proc.exitcode}"
        return "PDF fallback extraction failed: extractor returned no content"
    if status == "ok":
        return payload
    return str(payload)


def _read_document_with_timeout(ws: Path, rel_path: str, max_chars: int = 8000, tenant_id: str | None = None) -> str:
    """Run document parsing in a killable child process so one bad file cannot freeze the site."""
    ctx = mp.get_context("spawn")
    out_queue: mp.Queue = ctx.Queue(maxsize=1)
    proc = ctx.Process(
        target=_read_document_worker,
        args=(out_queue, str(ws), rel_path, max_chars, tenant_id),
        daemon=True,
    )
    proc.start()
    proc.join(_READ_DOCUMENT_TIMEOUT_SECONDS)
    if proc.is_alive():
        proc.terminate()
        proc.join(2)
        if proc.is_alive():
            proc.kill()
            proc.join(1)
        if Path(rel_path).suffix.lower() == ".pdf":
            return _read_pdf_fast_with_timeout(ws, rel_path, max_chars=max_chars, tenant_id=tenant_id)
        return (
            f"Document read timed out after {_READ_DOCUMENT_TIMEOUT_SECONDS}s. "
            "The file may be too large or too complex to extract safely. "
            "Please split it, convert it to text/Markdown, or read a smaller excerpt."
        )
    try:
        status, payload = out_queue.get_nowait()
    except queue.Empty:
        if proc.exitcode:
            return f"Document read failed: extractor exited with code {proc.exitcode}"
        return "Document read failed: extractor returned no content"
    if status == "ok":
        return payload
    return str(payload)


async def _read_document(
    ws: Path, rel_path: str, max_chars: int = _READ_DOCUMENT_HARD_CHAR_CEILING, tenant_id: str | None = None
) -> str:
    """Read content from office documents (PDF, DOCX, XLSX, PPTX)."""
    return await asyncio.to_thread(_read_document_with_timeout, ws, rel_path, max_chars, tenant_id)


async def _read_document_from_storage(
    agent_id: uuid.UUID,
    rel_path: str,
    max_chars: int = 8000,
    tenant_id: str | None = None,
) -> str:
    storage = get_storage_backend()
    try:
        resolved = await _resolve_exact_storage_source_path(agent_id, rel_path, tenant_id)
    except Exception as exc:
        return (
            "Document read was not started.\n"
            "Stage: storage_lookup\n"
            f"Requested path: {rel_path}\n"
            f"Reason: {type(exc).__name__}: {str(exc)[:200]}"
        )
    if not resolved.exists:
        return _exact_storage_source_error(resolved)

    try:
        version = await storage.get_version(resolved.storage_key)
    except Exception as exc:
        return (
            "Document read was not started.\n"
            "Stage: storage_metadata\n"
            f"Requested path: {resolved.virtual_path}\n"
            f"Reason: {type(exc).__name__}: {str(exc)[:200]}"
        )
    if not version.exists or version.is_dir:
        return _exact_storage_source_error(resolved)
    if version.size > _READ_DOCUMENT_MAX_FILE_BYTES:
        return (
            "Document read was not started.\n"
            "Stage: materialization\n"
            f"Requested path: {resolved.virtual_path}\n"
            "File exists: true\n"
            f"File size: {version.size} bytes\n"
            f"Limit: {_READ_DOCUMENT_MAX_FILE_BYTES} bytes\n"
            "Reason: file exceeds the document-processing limit."
        )

    try:
        temp_workspace = await _prepare_temp_workspace(
            agent_id,
            tenant_id=tenant_id,
            paths=[resolved.virtual_path],
            max_file_bytes=_READ_DOCUMENT_MAX_FILE_BYTES,
        )
    except Exception as exc:
        return (
            "Document read was not started.\n"
            "Stage: materialization\n"
            f"Requested path: {resolved.virtual_path}\n"
            "File exists: true\n"
            f"Reason: {type(exc).__name__}: {str(exc)[:200]}"
        )
    try:
        materialized_path = temp_workspace.root / resolved.virtual_path
        if not materialized_path.is_file():
            return (
                "Document read was not started.\n"
                "Stage: materialization\n"
                f"Requested path: {resolved.virtual_path}\n"
                "File exists: true\n"
                "Reason: storage file was not materialized into the document workspace."
            )
        content = await _read_document(
            temp_workspace.root,
            resolved.virtual_path,
            max_chars=max_chars,
            tenant_id=None,
        )
        return content
    finally:
        temp_workspace.cleanup()


async def _convert_csv_to_xlsx(agent_id: uuid.UUID, ws: Path, arguments: dict) -> str:
    source_path = arguments.get("source_path")
    target_path = arguments.get("target_path")
    if not source_path or not target_path:
        return "❌ Missing 'source_path' or 'target_path'."
    try:
        src_file = _resolve_tool_source_path(ws, source_path)
        tgt_file = _resolve_tool_target_path(ws, target_path)
    except ValueError as exc:
        return str(exc)
    if not src_file.exists():
        return f"❌ Source file not found: {source_path}"

    try:
        import csv
        from openpyxl import Workbook

        text = src_file.read_text(encoding="utf-8-sig")
        lines = [line.strip() for line in text.splitlines() if line.strip()][:10]
        candidates = [",", "，", ";", "\t", "|"]
        delimiter = ","
        if lines:
            scores = {candidate: sum(line.count(candidate) for line in lines) for candidate in candidates}
            if any(scores.values()):
                delimiter = max(scores, key=scores.get)

        wb = Workbook()
        ws_sheet = wb.active
        with src_file.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f, delimiter=delimiter)
            for row in reader:
                values = list(row)
                while values and not str(values[-1] or "").strip():
                    values.pop()
                if values:
                    ws_sheet.append(values)

        tgt_file.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(tgt_file))
        return f"✅ Successfully converted CSV to Excel: {target_path}"
    except Exception as e:
        logger.exception(f"Convert CSV to XLSX failed: {e}")
        return f"❌ Conversion failed: {e}"


async def _convert_html_to_pdf(agent_id: uuid.UUID, ws: Path, arguments: dict) -> str:
    source_path = arguments.get("source_path")
    target_path = arguments.get("target_path")
    if not source_path or not target_path:
        return "❌ Missing 'source_path' or 'target_path'."
    try:
        src_file = _resolve_tool_source_path(ws, source_path)
        tgt_file = _resolve_tool_target_path(ws, target_path)
    except ValueError as exc:
        return str(exc)
    if not src_file.exists():
        return f"❌ Source file not found: {source_path}"

    return await convert_html_file_to_pdf(src_file, tgt_file, str(target_path), arguments)


async def _convert_html_to_pptx(agent_id: uuid.UUID, ws: Path, arguments: dict) -> str:
    source_path = arguments.get("source_path")
    target_path = arguments.get("target_path")
    if not source_path or not target_path:
        return "❌ Missing paths."
    try:
        src_file = _resolve_tool_source_path(ws, source_path)
        tgt_file = _resolve_tool_target_path(ws, target_path)
    except ValueError as exc:
        return str(exc)
    if not src_file.exists():
        return "❌ Source file not found."

    return await convert_html_file_to_pptx(src_file, tgt_file, str(target_path), ws, arguments)


async def _convert_markdown_to_docx(agent_id: uuid.UUID, ws: Path, arguments: dict) -> str:
    source_path = arguments.get("source_path")
    target_path = arguments.get("target_path")
    if not source_path or not target_path:
        return "❌ Missing paths."
    try:
        src_file = _resolve_tool_source_path(ws, source_path)
        tgt_file = _resolve_tool_target_path(ws, target_path)
    except ValueError as exc:
        return str(exc)
    if not src_file.exists():
        return "❌ Source file not found."

    try:
        from docx import Document

        md_text = src_file.read_text(encoding="utf-8")
        doc = Document()

        def flush_paragraph(lines: list[str]) -> None:
            text = " ".join(line.strip() for line in lines if line.strip()).strip()
            if text:
                doc.add_paragraph(text)

        paragraph_lines: list[str] = []
        lines = md_text.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            stripped = line.strip()

            if not stripped:
                flush_paragraph(paragraph_lines)
                paragraph_lines = []
                i += 1
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if heading_match:
                flush_paragraph(paragraph_lines)
                paragraph_lines = []
                level = min(len(heading_match.group(1)), 6)
                doc.add_heading(heading_match.group(2).strip(), level=level)
                i += 1
                continue

            bullet_match = re.match(r"^[-*+]\s+(.*)$", stripped)
            ordered_match = re.match(r"^\d+\.\s+(.*)$", stripped)
            if bullet_match or ordered_match:
                flush_paragraph(paragraph_lines)
                paragraph_lines = []
                text = (bullet_match or ordered_match).group(1).strip()
                if text:
                    doc.add_paragraph(text, style="List Bullet" if bullet_match else "List Number")
                i += 1
                continue

            if "|" in stripped:
                table_lines: list[str] = []
                flush_paragraph(paragraph_lines)
                paragraph_lines = []
                while i < len(lines) and "|" in lines[i]:
                    candidate = lines[i].strip()
                    if candidate:
                        table_lines.append(candidate)
                    i += 1
                data_rows = []
                for raw in table_lines:
                    cells = [cell.strip() for cell in raw.strip("|").split("|")]
                    if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                        continue
                    if any(cell for cell in cells):
                        data_rows.append(cells)
                if data_rows:
                    table = doc.add_table(rows=len(data_rows), cols=max(len(row) for row in data_rows))
                    table.style = "Table Grid"
                    for row_idx, row in enumerate(data_rows):
                        for col_idx, cell in enumerate(row):
                            table.cell(row_idx, col_idx).text = cell
                continue

            paragraph_lines.append(stripped)
            i += 1

        flush_paragraph(paragraph_lines)

        tgt_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(tgt_file))
        return f"✅ Successfully converted Markdown to Word: {target_path}"
    except Exception as e:
        logger.exception(f"Convert MD to Docx failed: {e}")
        return f"❌ Conversion failed: {e}"


async def _convert_markdown_to_pdf(agent_id: uuid.UUID, ws: Path, arguments: dict) -> str:
    source_path = arguments.get("source_path")
    target_path = arguments.get("target_path")
    if not source_path or not target_path:
        return "❌ Missing paths."
    try:
        src_file = _resolve_tool_source_path(ws, source_path)
        tgt_file = _resolve_tool_target_path(ws, target_path)
    except ValueError as exc:
        return str(exc)
    if not src_file.exists():
        return "❌ Source file not found."

    try:
        from weasyprint import HTML

        md_text = src_file.read_text(encoding="utf-8")

        def escape_html(text: str) -> str:
            return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

        def render_inline(text: str) -> str:
            text = escape_html(text)
            text = re.sub(r"\*\*\*(.*?)\*\*\*", r"<strong><em>\1</em></strong>", text)
            text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", text)
            text = re.sub(r"__(.*?)__", r"<strong>\1</strong>", text)
            text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", text)
            text = re.sub(r"_(.*?)_", r"<em>\1</em>", text)
            text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
            text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)
            return text

        def is_table_separator(line: str) -> bool:
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)

        html_parts: list[str] = []
        lines = md_text.splitlines()
        in_list = False
        i = 0
        while i < len(lines):
            raw_line = lines[i]
            line = raw_line.rstrip()
            stripped = line.strip()
            if not stripped:
                if in_list:
                    html_parts.append("</ul>")
                    in_list = False
                i += 1
                continue

            heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if heading_match:
                if in_list:
                    html_parts.append("</ul>")
                    in_list = False
                level = len(heading_match.group(1))
                html_parts.append(f"<h{level}>{render_inline(heading_match.group(2).strip())}</h{level}>")
                i += 1
                continue

            bullet_match = re.match(r"^[-*+]\s+(.*)$", stripped)
            if bullet_match:
                if not in_list:
                    html_parts.append("<ul>")
                    in_list = True
                html_parts.append(f"<li>{render_inline(bullet_match.group(1).strip())}</li>")
                i += 1
                continue

            if "|" in stripped and i + 1 < len(lines) and is_table_separator(lines[i + 1].strip()):
                if in_list:
                    html_parts.append("</ul>")
                    in_list = False
                header_cells = [render_inline(cell.strip()) for cell in stripped.strip("|").split("|")]
                table_rows: list[list[str]] = []
                i += 2
                while i < len(lines) and "|" in lines[i].strip():
                    row = [render_inline(cell.strip()) for cell in lines[i].strip().strip("|").split("|")]
                    table_rows.append(row)
                    i += 1
                html_parts.append(
                    "<table><thead><tr>" + "".join(f"<th>{cell}</th>" for cell in header_cells) + "</tr></thead><tbody>"
                )
                html_parts.extend("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in table_rows)
                html_parts.append("</tbody></table>")
                continue

            if in_list:
                html_parts.append("</ul>")
                in_list = False
            html_parts.append(f"<p>{render_inline(stripped)}</p>")
            i += 1

        if in_list:
            html_parts.append("</ul>")

        html_text = "\n".join(html_parts)

        full_html = (
            "<html><head><meta charset='utf-8'><style>"
            "body{font-family:'WenQuanYi Micro Hei','Noto Sans CJK SC',sans-serif;line-height:1.65;padding:2em;color:#111827;}"
            "h1,h2,h3{line-height:1.25;margin:1.2em 0 .55em;}"
            "p{margin:.55em 0;}"
            "table{width:100%;border-collapse:collapse;margin:1em 0;font-size:12px;}"
            "th,td{border:1px solid #d8dee9;padding:7px 9px;text-align:left;vertical-align:top;}"
            "th{background:#f3f4f6;font-weight:700;}"
            "code{background:#f3f4f6;padding:1px 4px;border-radius:4px;}"
            "a{color:#2563eb;text-decoration:none;}"
            "</style></head><body>"
            f"{html_text}"
            "</body></html>"
        )

        tgt_file.parent.mkdir(parents=True, exist_ok=True)
        HTML(string=full_html, base_url=str(ws.resolve())).write_pdf(str(tgt_file))
        return f"✅ Successfully converted Markdown to PDF: {target_path}"
    except Exception as e:
        logger.exception(f"Convert MD to PDF failed: {e}")
        return f"❌ Conversion failed: {e}"
