"""Workspace route implementations for agent file APIs."""

from __future__ import annotations

import base64
import mimetypes
import uuid
from pathlib import Path
from typing import Any

import aiofiles
from fastapi import HTTPException, Request, status
from fastapi.responses import FileResponse, Response
from fastapi.security import HTTPAuthorizationCredentials
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import decode_access_token, request_access_token
from app.models.agent import Agent
from app.models.user import User
from app.models.workspace import WorkspaceFileRevision
from app.services.agent_runtime_workspace import current_agent_runtime_workspace
from app.services.authentication_state import require_active_authentication_principal
from app.services.chat_attachments import sniff_image_mime_bytes
from app.services.focus_service import is_focus_file_path
from app.services.im_markdown_media import verify_im_image_ticket
from app.services.agent_file_urls import verify_agent_file_ticket
from app.services.llm.failure_outcome import render_message
from app.services.storage import ensure_local_path, guess_content_type
from app.services.workspace_collaboration import (
    acquire_edit_lock,
    content_hash,
    delete_workspace_file,
    list_revisions,
    read_text_if_exists,
    release_edit_lock,
    write_workspace_file,
)
from app.services.workspace_locking import workspace_locks


async def list_files_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """List files and directories in an agent's file system."""

    agent = await api._resolve_workspace_agent(db, current_user, agent_id, workspace_agent)
    is_creator = (agent.creator_id == current_user.id) or (current_user.role == "platform_admin")
    storage = api.get_storage_backend()
    storage_key, is_enterprise = api._visible_storage_key(agent_id, path, current_user.tenant_id)
    normalized_path = (path or "").strip().strip("/")
    path_exists = await storage.exists(storage_key)
    path_is_dir = await storage.is_dir(storage_key)
    if not path_exists and not path_is_dir:
        if not (
            normalized_path == ""
            or (is_enterprise and normalized_path == "enterprise_info")
        ):
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Path not found")
    elif path_exists and not path_is_dir:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Path is not a directory")

    items = []
    if not path and current_user.tenant_id:
        items.append(api.FileInfo(
            name="enterprise_info",
            path="enterprise_info",
            is_dir=True,
            size=0,
            modified_at="",
            version_token=None,
            url=None,
        ))
    entries = await storage.list_dir(storage_key) if path_exists or path_is_dir else []
    for entry in entries:
        if entry.name == ".gitkeep":
            continue
        if not path and entry.name.lower() in {"focus.md", "agenda.md"}:
            continue
        if not path and entry.name == "enterprise_info":
            continue
        if entry.name in api.CREATOR_ONLY_FILES and not is_creator:
            continue
        if is_enterprise:
            rel = str(Path(entry.key).relative_to(f"enterprise_info_{current_user.tenant_id}"))
            rel_path = f"enterprise_info/{rel}" if rel != "." else "enterprise_info"
        else:
            rel_path = str(Path(entry.key).relative_to(current_agent_runtime_workspace(agent_id).storage_prefix))
        items.append(api.FileInfo(
            name=entry.name,
            path=rel_path,
            is_dir=entry.is_dir,
            size=entry.size,
            modified_at=entry.modified_at,
            version_token=api._entry_version_token(entry),
            url=f"/api/agents/{agent_id}/files/download?path={rel_path}" if not entry.is_dir else None,
        ))
    return items


async def read_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """Read the content of a file."""

    agent = await api._resolve_workspace_agent(db, current_user, agent_id, workspace_agent)
    is_creator = (agent.creator_id == current_user.id) or (current_user.role == "platform_admin")
    filename = Path(path).name
    if filename in api.CREATOR_ONLY_FILES and not is_creator:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if is_focus_file_path(path):
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Focus is stored in the system database. Use the Focus API.",
        )
    storage = api.get_storage_backend()
    key, _ = api._visible_storage_key(agent_id, path, current_user.tenant_id)
    if not await storage.exists(key) or not await storage.is_file(key):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found")
    version = await storage.get_version(key)
    try:
        content = await storage.read_text(key, encoding="utf-8", errors="replace")
        return api.FileContent(path=path, content=content, version_token=version.token)
    except UnicodeDecodeError:
        stat = await storage.stat(key)
        return api.FileContent(
            path=path,
            content=f"[二进制文件: {Path(path).name}, {stat.size} bytes]",
            version_token=version.token,
        )


async def preview_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    current_user: User,
    db: AsyncSession,
):
    """Return a browser-friendly preview payload for Workspace files."""

    await api.check_agent_access(db, current_user, agent_id)
    storage = api.get_storage_backend()
    key, _ = api._visible_storage_key(agent_id, path, current_user.tenant_id)
    if not await storage.exists(key) or not await storage.is_file(key):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found")

    kind = api._file_kind(path)
    mime_type = mimetypes.guess_type(Path(path).name)[0] or "application/octet-stream"
    download_url = f"/api/agents/{agent_id}/files/download?path={path}"
    local_target: Path | None = None

    if kind in {"markdown", "html", "text"}:
        content = await storage.read_text(key, encoding="utf-8", errors="replace")
        return {
            "path": path,
            "kind": kind,
            "mime_type": mime_type,
            "content": content or "",
            "content_hash": content_hash(content or ""),
            "download_url": download_url,
        }
    if kind == "csv":
        content = await storage.read_text(key, encoding="utf-8", errors="replace")
        rows = api._parse_csv_rows(content)
        return {
            "path": path,
            "kind": kind,
            "mime_type": mime_type,
            "content": content,
            "content_hash": content_hash(content),
            "rows": rows[:500],
            "download_url": download_url,
        }
    if kind == "pdf":
        return {
            "path": path,
            "kind": kind,
            "mime_type": mime_type,
            "url": download_url,
            "download_url": download_url,
        }
    if kind == "xlsx":
        try:
            target = await ensure_local_path(key)
            local_target = target
            from openpyxl import load_workbook

            wb = load_workbook(target, read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets[:5]:
                rows = []
                for row in ws.iter_rows(max_row=120, max_col=30, values_only=True):
                    values = ["" if cell is None else str(cell) for cell in row]
                    while values and not str(values[-1] or "").strip():
                        values.pop()
                    if any(value.strip() for value in values):
                        rows.append(values)
                sheets.append({
                    "title": ws.title,
                    "rows": rows,
                })
            wb.close()
            return {
                "path": path,
                "kind": kind,
                "mime_type": mime_type,
                "text": api._extract_document_text(target, kind),
                "sheets": sheets,
                "download_url": download_url,
            }
        except Exception as exc:
            return {
                "path": path,
                "kind": kind,
                "mime_type": mime_type,
                "text": f"Preview extraction failed: {str(exc)[:200]}",
                "download_url": download_url,
            }
    if kind in {"docx", "pptx"}:
        target = await ensure_local_path(key)
        local_target = target
        extracted_text = api._extract_document_text(target, kind)
        companion = api._find_companion_text_preview(target)
        companion_content = await read_text_if_exists(companion) if companion is not None else None
        return {
            "path": path,
            "kind": kind,
            "mime_type": mime_type,
            "text": companion_content or extracted_text,
            "companion_path": str(companion.resolve().relative_to(api._agent_base_dir(agent_id).resolve())) if companion is not None and not path.startswith("enterprise_info") else None,
            "download_url": download_url,
        }
    companion = api._find_companion_text_preview(local_target) if local_target is not None else None
    if companion is not None:
        content = await read_text_if_exists(companion)
        return {
            "path": path,
            "kind": "text",
            "mime_type": "text/markdown" if companion.suffix.lower() == ".md" else "text/plain",
            "content": content or "",
            "content_hash": content_hash(content or ""),
            "companion_path": str(companion.resolve().relative_to(api._agent_base_dir(agent_id).resolve())) if not path.startswith("enterprise_info") else None,
            "download_url": download_url,
        }
    raw = await storage.read_bytes(key)
    encoded = base64.b64encode(raw[:1024 * 1024]).decode("ascii")
    return {
        "path": path,
        "kind": kind,
        "mime_type": mime_type,
        "size": len(raw),
        "base64_sample": encoded,
        "download_url": download_url,
    }


async def download_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    request: Request,
    token: str,
    im_ticket: str,
    inline: bool,
    credentials: HTTPAuthorizationCredentials | None,
    db: AsyncSession,
):
    """Download / serve a file from the agent workspace (browser-friendly)."""

    storage = api.get_storage_backend()
    if im_ticket:
        media_ticket = verify_agent_file_ticket(agent_id, path, im_ticket)
        if media_ticket is not None:
            key = media_ticket["storage_key"]
            if not await storage.is_file(key):
                raise HTTPException(status_code=404, detail=render_message("mediaAI.fileUnavailable"))
            local_path = await storage.local_path_for(key)
            headers = {"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"}
            if local_path is not None:
                return FileResponse(str(local_path), media_type=media_ticket["mime_type"], headers=headers)
            from fastapi.responses import StreamingResponse

            entry = await storage.stat(key)

            async def chunks():
                for offset in range(0, entry.size, 1024 * 1024):
                    yield await storage.read_range(key, offset, min(entry.size - 1, offset + 1024 * 1024 - 1))

            return StreamingResponse(chunks(), media_type=media_ticket["mime_type"], headers=headers)
        image_ticket = verify_im_image_ticket(agent_id, path, im_ticket)
        if image_ticket is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Image not found")
        key = image_ticket.storage_key
        if not await storage.is_file(key):
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Image not found")
        mime_type = sniff_image_mime_bytes(await storage.read_range(key, 0, 511))
        if mime_type is None:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Not an image")
        headers = {
            "Cache-Control": "private, no-store",
            "Referrer-Policy": "no-referrer",
            "X-Content-Type-Options": "nosniff",
        }
        local_path = await storage.local_path_for(key)
        if local_path is not None:
            return FileResponse(
                path=str(local_path),
                filename=Path(image_ticket.path).name,
                media_type=mime_type,
                content_disposition_type="inline",
                headers=headers,
            )
        return Response(
            content=await storage.read_bytes(key),
            media_type=mime_type,
            headers={
                **headers,
                "Content-Disposition": f'inline; filename="{Path(image_ticket.path).name}"',
            },
        )

    jwt_token = request_access_token(request, credentials, query_token=token)
    if not jwt_token:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")
    payload = decode_access_token(jwt_token)
    user_id = payload.get("sub")
    if not user_id:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")
    result = await db.execute(select(User).where(User.id == uuid.UUID(user_id)))
    user = await require_active_authentication_principal(
        db,
        result.scalar_one_or_none(),
        status_code=status.HTTP_401_UNAUTHORIZED,
    )
    agent, _access = await api.check_agent_access(db, user, agent_id)
    is_creator = (agent.creator_id == user.id) or (user.role == "platform_admin")
    filename = Path(path).name
    if filename in api.CREATOR_ONLY_FILES and not is_creator:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    key, _ = api._visible_storage_key(
        agent_id,
        path,
        user.tenant_id,
        workspace=api._runtime_workspace(agent),
    )
    if not await storage.exists(key) or not await storage.is_file(key):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found")
    presigned = await storage.presign_download_url(key, filename=Path(path).name, inline=inline)
    if presigned:
        return Response(status_code=302, headers={"Location": presigned})
    local_path = await storage.local_path_for(key)
    if local_path is not None:
        return FileResponse(
            path=str(local_path),
            filename=Path(path).name,
            content_disposition_type="inline" if inline else "attachment",
        )
    data = await storage.read_bytes(key)
    disposition = "inline" if inline else "attachment"
    return Response(
        content=data,
        media_type=guess_content_type(Path(path).name),
        headers={"Content-Disposition": f'{disposition}; filename="{Path(path).name}"'},
    )


async def write_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    data: Any,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """Write content to a file (create or overwrite)."""

    agent = await api._resolve_workspace_agent(db, current_user, agent_id, workspace_agent)
    is_creator = (agent.creator_id == current_user.id) or (current_user.role == "platform_admin")
    filename = Path(path).name
    if filename in api.CREATOR_ONLY_FILES and not is_creator:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if is_focus_file_path(path):
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Focus is stored in the system database. Use the Focus API.",
        )
    if path.startswith("enterprise_info"):
        if current_user.role not in ("platform_admin", "org_admin"):
            raise HTTPException(status_code=403, detail="Only admins can edit enterprise knowledge base")
        if path.strip("/") == "enterprise_info":
            raise HTTPException(status_code=400, detail="Cannot overwrite enterprise_info root")
        target, _, _ = api._visible_path(agent_id, path, current_user.tenant_id)
        target.parent.mkdir(parents=True, exist_ok=True)
        async with aiofiles.open(target, "w", encoding="utf-8") as f:
            await f.write(data.content)
        return {"status": "ok", "path": path, "revision_id": None}
    result = await write_workspace_file(
        db,
        agent_id=agent_id,
        base_dir=api._agent_base_dir(agent_id),
        path=path,
        content=data.content,
        actor_type="user",
        actor_id=current_user.id,
        operation="autosave" if data.autosave else "write",
        session_id=data.session_id,
        enforce_human_lock=False,
        merge_user_autosave=data.autosave,
        expected_version_token=data.expected_version_token,
    )
    if not result.ok:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=result.message)
    await api._record_project_skill_change(db, agent, current_user, result.path)
    await db.commit()
    return {"status": "ok", "path": result.path, "revision_id": result.revision_id}


async def lock_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    data: Any,
    current_user: User,
    db: AsyncSession,
):
    """Acquire or refresh a short-lived human editing lock for a file."""

    await api.check_agent_access(db, current_user, agent_id)
    if is_focus_file_path(data.path):
        raise HTTPException(status_code=status.HTTP_410_GONE, detail="Focus is stored in the system database.")
    lock = await acquire_edit_lock(
        db,
        agent_id=agent_id,
        path=data.path,
        user_id=current_user.id,
        session_id=data.session_id,
    )
    await db.commit()
    return {"status": "ok", "path": lock.path, "expires_at": lock.expires_at.isoformat()}


async def unlock_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    current_user: User,
    db: AsyncSession,
):
    """Release the current user's edit lock for a file."""

    await api.check_agent_access(db, current_user, agent_id)
    await release_edit_lock(db, agent_id=agent_id, path=path, user_id=current_user.id)
    await db.commit()
    return {"status": "ok", "path": path}


async def get_file_revisions_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    current_user: User,
    db: AsyncSession,
):
    """List version history for the currently opened Workspace file."""

    await api.check_agent_access(db, current_user, agent_id)
    if is_focus_file_path(path) or path.startswith("enterprise_info"):
        return []
    revisions = await list_revisions(db, agent_id=agent_id, path=path)
    return [
        {
            "id": str(rev.id),
            "path": rev.path,
            "operation": rev.operation,
            "actor_type": rev.actor_type,
            "actor_id": str(rev.actor_id) if rev.actor_id else None,
            "session_id": rev.session_id,
            "before_content": rev.before_content,
            "after_content": rev.after_content,
            "created_at": rev.created_at.isoformat() if rev.created_at else None,
            "updated_at": rev.updated_at.isoformat() if rev.updated_at else None,
        }
        for rev in revisions
    ]


async def restore_file_revision_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    data: Any,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """Restore a file to a previous revision's after-content."""

    await api.check_agent_access(db, current_user, agent_id)
    result = await db.execute(
        select(WorkspaceFileRevision).where(
            WorkspaceFileRevision.id == data.revision_id,
            WorkspaceFileRevision.agent_id == agent_id,
        )
    )
    revision = result.scalar_one_or_none()
    if not revision:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Revision not found")
    if revision.after_content is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot restore an empty/deleted revision")
    restored = await write_workspace_file(
        db,
        agent_id=agent_id,
        base_dir=api._agent_base_dir(agent_id),
        path=revision.path,
        content=revision.after_content,
        actor_type="user",
        actor_id=current_user.id,
        operation="restore",
        enforce_human_lock=False,
        expected_version_token=data.expected_version_token,
    )
    if not restored.ok:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=restored.message)
    await api._record_project_skill_change(db, workspace_agent, current_user, revision.path)
    await db.commit()
    return {"status": "ok", "path": revision.path, "revision_id": restored.revision_id}


async def delete_file_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    path: str,
    expected_version_token: str | None,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """Delete a file."""

    await api._require_agent_file_delete_access(db, current_user, agent_id)
    agent = await api._resolve_workspace_agent(db, current_user, agent_id, workspace_agent)
    is_creator = (agent.creator_id == current_user.id) or (current_user.role == "platform_admin")
    filename = Path(path).name
    if filename in api.CREATOR_ONLY_FILES and not is_creator:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    if is_focus_file_path(path):
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Focus is stored in the system database. Use the Focus API.",
        )
    if path.startswith("enterprise_info") and current_user.role not in ("platform_admin", "org_admin"):
        raise HTTPException(status_code=403, detail="Only admins can delete enterprise knowledge base files")
    if path.strip("/") == "enterprise_info":
        raise HTTPException(status_code=400, detail="Cannot delete enterprise_info root")
    project_skill_result = await api._delete_bound_project_skill(db, agent, current_user, path)
    if project_skill_result is not None:
        await db.commit()
        return project_skill_result
    result = await delete_workspace_file(
        db,
        agent_id=agent_id,
        base_dir=api._agent_base_dir(agent_id),
        path=path,
        actor_type="user",
        actor_id=current_user.id,
        enforce_human_lock=False,
        expected_version_token=expected_version_token,
    )
    if not result.ok:
        if "not found" in result.message.lower():
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=result.message)
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=result.message)
    await api._record_project_skill_change(db, agent, current_user, path)
    await db.commit()
    return {"status": "ok", "path": path}


async def import_skill_to_agent_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    body: Any,
    current_user: User,
    db: AsyncSession,
    workspace_agent: Agent,
):
    """Import a global skill into this agent's skills/ workspace folder."""

    agent = await api._resolve_workspace_agent(db, current_user, agent_id, workspace_agent)
    _agent, access_level = await api.check_agent_access(db, current_user, agent_id)
    if access_level != "manage":
        raise HTTPException(status_code=403, detail="需要数字员工管理权限")
    from sqlalchemy import or_
    from sqlalchemy.orm import selectinload
    from app.models.skill import Skill

    result = await db.execute(
        select(Skill)
        .where(
            Skill.id == body.skill_id,
            or_(Skill.tenant_id.is_(None), Skill.tenant_id == agent.tenant_id),
        )
        .options(selectinload(Skill.files))
    )
    skill = result.scalar_one_or_none()
    if not skill:
        raise HTTPException(status_code=404, detail="Skill not found")
    if agent.scope == "project":
        from app.models.project import Project
        from app.services.project_capability_options import load_project_capability_options
        from app.services.project_skill_assets import bind_library_skill_to_project_agent

        project = await db.get(Project, agent.project_id)
        if project is None or project.tenant_id != agent.tenant_id:
            raise HTTPException(status_code=409, detail="Project Agent workspace is unavailable")
        source_agent = await db.get(Agent, agent.source_agent_id) if agent.source_agent_id is not None else None
        if (
            source_agent is not None
            and (
                source_agent.tenant_id != project.tenant_id
                or source_agent.scope != "standard"
                or source_agent.is_deleted
            )
        ):
            source_agent = None
        options = await load_project_capability_options(
            db,
            project.tenant_id,
            [source_agent] if source_agent is not None else [],
        )
        allowed = (
            options.allows(source_agent.id, "skill", skill.id)
            if source_agent is not None
            else options.allows_shared("skill", skill.id)
        )
        if not allowed:
            raise HTTPException(status_code=422, detail="Selected project Skill is unavailable")
        binding = await bind_library_skill_to_project_agent(
            db,
            project,
            skill_id=skill.id,
            project_agent_id=agent.id,
            is_enabled=True,
            scope={},
            actor_user_id=current_user.id,
            actor_display_name=current_user.display_name,
        )
        await db.commit()
        return {
            "status": "ok",
            "skill_name": skill.name,
            "folder_name": skill.folder_name,
            "files_written": len(skill.files),
            "files": [file.path for file in skill.files],
            "project_skill_binding_id": str(binding.id),
        }
    if skill.status != "draft":
        from app.services.skill_market import install_market_skill

        return await install_market_skill(
            db,
            agent=agent,
            skill_id=skill.id,
            actor_user_id=current_user.id,
        )
    if not skill.files:
        raise HTTPException(status_code=400, detail="Skill has no files")
    storage = api.get_storage_backend()
    written = []
    async with workspace_locks(agent_id, []):
        for file_entry in skill.files:
            skill_key = api._agent_storage_key(agent_id, f"skills/{skill.folder_name}/{file_entry.path}")
            await storage.write_text(skill_key, file_entry.content, encoding="utf-8")
            written.append(file_entry.path)
    return {
        "status": "ok",
        "skill_name": skill.name,
        "folder_name": skill.folder_name,
        "files_written": len(written),
        "files": written,
    }


async def upload_file_to_workspace_impl(
    api: Any,
    *,
    agent_id: uuid.UUID,
    file: Any,
    path: str,
    current_user: User,
    db: AsyncSession,
):
    """Upload a binary file to agent workspace."""

    await api.check_agent_access(db, current_user, agent_id)
    normalized_path = (path or "").strip().strip("/")
    if not normalized_path or normalized_path == ".":
        normalized_path = api.DEFAULT_UPLOAD_DIR
    if normalized_path not in {"workspace", "skills"} and not normalized_path.startswith(("workspace/", "skills/")):
        raise HTTPException(status_code=400, detail="右侧根目录视图是 agent 根目录；上传文件时请放到 workspace/ 或 skills/ 目录下")
    filename = (file.filename or "unnamed").replace("/", "_").replace("\\", "_")
    storage = api.get_storage_backend()
    file_key = api._agent_storage_key(agent_id, f"{normalized_path}/{filename}")
    content = await file.read()
    extracted_path = None
    async with workspace_locks(agent_id, []):
        await storage.write_bytes(file_key, content, content_type=guess_content_type(filename))
        from app.services.text_extractor import needs_extraction, save_extracted_text
        if needs_extraction(filename):
            save_path = await ensure_local_path(file_key)
            txt_file = save_extracted_text(save_path, content, filename)
            if txt_file:
                extracted_path = f"{normalized_path}/{txt_file.name}"
                extracted_key = api._agent_storage_key(agent_id, extracted_path)
                await storage.write_bytes(
                    extracted_key,
                    txt_file.read_bytes(),
                    content_type="text/plain; charset=utf-8",
                )
    return {
        "status": "ok",
        "path": f"{normalized_path}/{filename}",
        "url": f"/api/agents/{agent_id}/files/download?path={normalized_path}/{filename}",
        "filename": filename,
        "size": len(content),
        "extracted_text_path": extracted_path,
    }
