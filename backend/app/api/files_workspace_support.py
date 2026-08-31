"""Shared helper implementations for agent workspace file APIs."""

from __future__ import annotations

import csv
import io
import uuid
from pathlib import Path
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.agent import Agent
from app.models.user import User
from app.services.storage import normalize_storage_key
from app.services.storage_runtime.base import StorageEntry
from app.services.workspace_paths import WorkspacePathError, resolve_agent_visible_path


async def resolve_workspace_agent_impl(
    api: Any,
    db: AsyncSession,
    current_user: User,
    agent_id: uuid.UUID,
    candidate: object,
) -> Agent:
    """Keep endpoint functions callable outside FastAPI dependency injection."""

    if getattr(candidate, "id", None) == agent_id:
        return candidate  # type: ignore[return-value]
    agent, _access = await api.check_agent_access(db, current_user, agent_id)
    return agent


def runtime_workspace_impl(api: Any, agent: Agent):
    if getattr(agent, "scope", "standard") != "project":
        return api.standard_agent_runtime_workspace(agent.id)
    if agent.project_id is None or agent.tenant_id is None:
        raise HTTPException(status_code=409, detail="Project Agent workspace is unavailable")
    return api.project_agent_runtime_workspace(
        agent_id=agent.id,
        tenant_id=agent.tenant_id,
        project_id=agent.project_id,
    )


async def record_project_skill_change_impl(
    api: Any,
    db: AsyncSession,
    agent: Agent,
    current_user: User,
    path: str,
) -> None:
    normalized = normalize_storage_key(path)
    parts = Path(normalized).parts
    if getattr(agent, "scope", "standard") != "project" or len(parts) < 2 or parts[0] != "skills":
        return
    from app.models.project import Project
    from app.services.project_git_service import commit_project_changes, project_user_git_email
    from app.services.project_skill_assets import register_project_workspace_skill

    project = await db.get(Project, agent.project_id)
    if project is None or project.tenant_id != agent.tenant_id:
        raise HTTPException(status_code=409, detail="Project Agent workspace is unavailable")
    skill_root = api._agent_base_dir(agent.id) / "skills" / parts[1]
    if (skill_root / "SKILL.md").is_file():
        await register_project_workspace_skill(
            db,
            project,
            project_agent_id=agent.id,
            folder_name=parts[1],
            actor_user_id=current_user.id,
            actor_display_name=current_user.display_name,
        )
        return
    await commit_project_changes(
        project,
        f"Update project Skill files: {parts[1]}",
        [f".agents/{agent.id}/skills/{parts[1]}"],
        author_name=current_user.display_name,
        author_email=project_user_git_email(current_user.id),
    )


async def delete_bound_project_skill_impl(
    api: Any,
    db: AsyncSession,
    agent: Agent,
    current_user: User,
    path: str,
) -> dict[str, Any] | None:
    """Route a project Skill root deletion through its binding lifecycle."""

    normalized = normalize_storage_key(path)
    parts = Path(normalized).parts
    is_skill_root = len(parts) == 2 and parts[0] == "skills"
    is_skill_manifest = len(parts) == 3 and parts[0] == "skills" and parts[2] == "SKILL.md"
    if getattr(agent, "scope", "standard") != "project" or not (is_skill_root or is_skill_manifest):
        return None

    from app.models.project import Project, ProjectCapabilityBinding
    from app.services.project_service import add_event
    from app.services.project_skill_assets import (
        delete_project_skill_asset,
        project_skill_deletion_impact,
    )

    project = await db.get(Project, agent.project_id)
    if (
        project is None
        or project.tenant_id != agent.tenant_id
        or project.owner_user_id != current_user.id
    ):
        raise HTTPException(status_code=404, detail="Project not found")
    bindings = list(
        (
            await db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == project.id,
                    ProjectCapabilityBinding.tenant_id == project.tenant_id,
                    ProjectCapabilityBinding.capability_type == "skill",
                    ProjectCapabilityBinding.inherited_from_agent_id == agent.id,
                )
            )
        ).scalars()
    )
    expected_path = f"skills/{parts[1]}"
    binding = next(
        (
            item
            for item in bindings
            if isinstance(item.config, dict)
            and isinstance(item.config.get("skill_asset"), dict)
            and item.config["skill_asset"].get("path") == expected_path
        ),
        None,
    )
    if binding is None:
        return None

    impact = await project_skill_deletion_impact(db, project, binding)
    deleted = await delete_project_skill_asset(
        db,
        project,
        binding,
        actor_user_id=current_user.id,
        actor_display_name=current_user.display_name,
    )
    add_event(
        db,
        project,
        "capability.deleted",
        f"Deleted project Skill {deleted['skill_name']}",
        actor_user_id=current_user.id,
        metadata={
            "asset_id": deleted["asset_id"],
            "affected_member_count": impact["affected_member_count"],
            "source": "agent_files",
        },
    )
    await db.flush()
    return {
        "status": "ok",
        "path": expected_path,
        "project_skill_deleted": True,
        "affected_member_count": impact["affected_member_count"],
        "affected_members": impact["affected_members"],
    }


def safe_path_impl(api: Any, agent_id: uuid.UUID, rel_path: str) -> Path:
    """Ensure the path is within the agent's directory (no path traversal)."""

    base = api._agent_base_dir(agent_id)
    full = (base / rel_path).resolve()
    if not str(full).startswith(str(base.resolve())):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Path traversal not allowed")
    return full


def visible_path_impl(
    api: Any,
    agent_id: uuid.UUID,
    rel_path: str,
    tenant_id: uuid.UUID | None,
) -> tuple[Path, Path, bool]:
    """Resolve an agent-visible path, including virtual enterprise_info/."""

    try:
        resolved = resolve_agent_visible_path(
            api._agent_base_dir(agent_id),
            rel_path,
            workspace_root=Path(api.settings.AGENT_DATA_DIR),
            tenant_id=str(tenant_id) if tenant_id else None,
        )
    except WorkspacePathError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=str(exc)) from exc
    return resolved.path, resolved.relative_root, resolved.is_enterprise


def is_enterprise_visible_path_impl(rel_path: str) -> bool:
    normalized = (rel_path or "").strip().strip("/")
    return normalized == "enterprise_info" or normalized.startswith("enterprise_info/")


def visible_storage_key_impl(
    api: Any,
    agent_id: uuid.UUID,
    rel_path: str,
    tenant_id: uuid.UUID | None,
    *,
    workspace: Any = None,
) -> tuple[str, bool]:
    normalized = (rel_path or "").strip().strip("/")
    if api._is_enterprise_visible_path(normalized):
        if not tenant_id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No tenant associated")
        sub_path = normalized[len("enterprise_info"):].lstrip("/")
        return api._enterprise_storage_key(str(tenant_id), sub_path), True
    key = workspace.storage_key(normalized) if workspace is not None else api._agent_storage_key(agent_id, normalized)
    return key, False


async def require_agent_file_delete_access_impl(
    api: Any,
    db: AsyncSession,
    current_user: User,
    agent_id: uuid.UUID,
) -> None:
    """Allow destructive workspace file operations only for managers/admins."""

    _agent, access_level = await api.check_agent_access(db, current_user, agent_id)
    if access_level == "manage" or current_user.role in ("platform_admin", "org_admin", "super_admin"):
        return
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Only agent managers or admins can delete files",
    )


def entry_version_token_impl(entry: StorageEntry) -> str | None:
    token = entry.version_id or entry.etag or entry.content_hash
    if token:
        return token
    if entry.is_dir:
        return None
    if entry.modified_at or entry.size:
        return f"{entry.modified_at}:{entry.size}"
    return None


def file_kind_impl(api: Any, path: str) -> str:
    file_path = Path(path)
    ext = file_path.suffix.lower()
    name = file_path.name.lower()
    if ext in {".md", ".markdown"}:
        return "markdown"
    if ext == ".csv":
        return "csv"
    if ext in {".html", ".htm"}:
        return "html"
    if ext == ".pdf":
        return "pdf"
    if ext in {".xlsx", ".xls"}:
        return "xlsx"
    if ext in {".docx", ".doc"}:
        return "docx"
    if ext in {".pptx", ".ppt"}:
        return "pptx"
    if ext in {".txt", ".log", ".json"} or ext in api.TEXT_PREVIEW_EXTENSIONS or name in api.TEXT_PREVIEW_FILENAMES:
        return "text"
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}:
        return "image"
    return "binary"


def find_companion_text_preview_impl(target: Path) -> Path | None:
    for suffix in (".md", ".txt"):
        candidate = target.with_suffix(suffix)
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def extract_document_text_impl(target: Path, kind: str) -> str:
    """Best-effort rich document text extraction for lightweight previews."""

    try:
        if kind == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(target, read_only=True, data_only=True)
            sheets: list[str] = []
            for ws in wb.worksheets[:5]:
                rows = []
                for row in ws.iter_rows(max_row=80, max_col=20, values_only=True):
                    rows.append("\t".join("" if cell is None else str(cell) for cell in row))
                sheets.append(f"Sheet: {ws.title}\n" + "\n".join(rows))
            return "\n\n".join(sheets)
        if kind == "docx":
            from docx import Document

            doc = Document(str(target))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if kind == "pptx":
            from pptx import Presentation

            prs = Presentation(str(target))
            slides = []
            for idx, slide in enumerate(prs.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                slides.append(f"Slide {idx}\n" + "\n".join(texts))
            return "\n\n".join(slides)
    except ImportError as exc:
        return f"Missing preview dependency: {exc}"
    except Exception as exc:
        return f"Preview extraction failed: {str(exc)[:200]}"
    return ""


def detect_csv_delimiter_impl(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()][:10]
    if not lines:
        return ","
    candidates = [",", "，", ";", "\t", "|"]
    scores = {
        candidate: sum(line.count(candidate) for line in lines)
        for candidate in candidates
    }
    return max(scores, key=scores.get) if any(scores.values()) else ","


def parse_csv_rows_impl(api: Any, text: str) -> list[list[str]]:
    delimiter = api._detect_csv_delimiter(text)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    normalized: list[list[str]] = []
    for row in rows[:500]:
        values = list(row)
        while values and not str(values[-1] or "").strip():
            values.pop()
        if values:
            normalized.append(values)
    return normalized
