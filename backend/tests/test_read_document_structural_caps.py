"""read_document structural caps must be generous + VISIBLY signposted — never a
silent content drop.

Round 3 of the "read_document 还是有内容被截断" thread. After removing the per-cell
500-char truncation, the remaining structural caps still silently dropped content
past their limit with NO marker, unrecoverable for the agent:
  - xlsx per-sheet rows capped at 200 (a 口径 sheet with >200 rows lost the rest)
  - pdf pages capped at 50, pptx slides capped at 50
  - xlsx workbook sheets capped at 10

Caps are now generous memory-safety bounds (rows 10000, pages/slides 200); when one
actually trims a pathologically large doc it appends an explicit, visible marker so
the agent knows content remains and can recover it (read_file / sql_execute).
"""

from __future__ import annotations

from pathlib import Path

import app.services.agent_tools as at
from app.services.agent_tools import _read_document_sync


def _make_narrow_xlsx(path: Path, n_rows: int) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "模块")
    ws.cell(1, 2, "口径")
    for r in range(2, 2 + n_rows):
        ws.cell(r, 1, f"row{r}")
        ws.cell(r, 2, f"SELECT {r} FROM yk.t")
    wb.save(str(path))


def test_xlsx_returns_far_more_than_200_rows(tmp_path):
    # Regression: the old max_row=200 silently dropped rows 201+. 300 rows survive.
    _make_narrow_xlsx(tmp_path / "doc.xlsx", n_rows=300)
    out = _read_document_sync(tmp_path, "doc.xlsx")
    assert "SELECT 250 FROM yk.t" in out  # past the old 200 cap
    assert "SELECT 301 FROM yk.t" in out  # last data row
    assert "only the first" not in out  # no cap signpost for a normal-size sheet


def test_xlsx_row_cap_is_signposted_not_silent(tmp_path, monkeypatch):
    monkeypatch.setattr(at, "_READ_DOCUMENT_MAX_ROWS", 5)
    _make_narrow_xlsx(tmp_path / "doc.xlsx", n_rows=40)  # 40 data rows => max_row 41
    out = _read_document_sync(tmp_path, "doc.xlsx")
    assert "only the first 5 are shown" in out  # visible marker
    assert "sql_execute" in out  # tells the agent how to recover the rest
    assert "SELECT 2 FROM yk.t" in out  # the first rows are still present


def test_xlsx_sheet_cap_is_signposted_not_silent(tmp_path, monkeypatch):
    monkeypatch.setattr(at, "_READ_DOCUMENT_MAX_SHEETS", 2)
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.cell(1, 1, "s0")
    for i in range(1, 4):
        wb.create_sheet(f"extra{i}").cell(1, 1, f"data{i}")
    wb.save(str(tmp_path / "doc.xlsx"))  # 4 sheets total
    out = _read_document_sync(tmp_path, "doc.xlsx")
    assert "workbook has 4 sheets; only the first 2 are shown" in out


def test_pptx_slide_cap_is_signposted_not_silent(tmp_path, monkeypatch):
    monkeypatch.setattr(at, "_READ_DOCUMENT_MAX_SLIDES", 2)
    from pptx import Presentation

    prs = Presentation()
    title_only = prs.slide_layouts[5]
    for i in range(4):
        slide = prs.slides.add_slide(title_only)
        slide.shapes.title.text = f"slide {i}"
    prs.save(str(tmp_path / "d.pptx"))  # 4 slides total
    out = _read_document_sync(tmp_path, "d.pptx")
    assert "presentation has 4 slides; only the first 2 are shown" in out
    assert "slide 0" in out  # the first slides are still present
