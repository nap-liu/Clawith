"""API-level acceptance tests for the AI-native project closed loop.

These tests intentionally use the real SQLAlchemy models and managed Git
implementation behind a small FastAPI app.  They verify observable API
contracts and durable state, rather than mocking the project service itself.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from typing import Any, AsyncIterator

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import app.models.agent  # noqa: F401
import app.models.audit  # noqa: F401
import app.models.chat_compaction  # noqa: F401
import app.models.chat_session  # noqa: F401
import app.models.llm  # noqa: F401
import app.models.mcp_server  # noqa: F401
import app.models.org  # noqa: F401
import app.models.participant  # noqa: F401
import app.models.project  # noqa: F401
import app.models.skill  # noqa: F401
import app.models.subagent_run  # noqa: F401
import app.models.tenant  # noqa: F401
import app.models.tool  # noqa: F401
import app.models.user  # noqa: F401
from app.api import files as files_api
from app.api import mcp_servers as mcp_servers_api
from app.api import projects as projects_api
from app.core.security import get_current_user
from app.database import Base, get_db
from app.models.agent import Agent
from app.models.audit import ChatMessage
from app.models.chat_session import ChatSession
from app.models.llm import LLMModel
from app.models.mcp_server import MCPServer
from app.models.org import OrgDepartment, OrgMember
from app.models.project import (
    Project,
    ProjectAccessGrant,
    ProjectCapabilityBinding,
    ProjectEvent,
    ProjectMemberSnapshot,
    ProjectTemplate,
)
from app.models.skill import Skill, SkillFile, SkillInstall
from app.models.tenant import Tenant
from app.models.tool import AgentTool, Tool
from app.models.user import Identity, User
from app.services.project_git_service import project_repo_path

pytestmark = pytest.mark.asyncio

TABLES = [
    "llm_models",
    "identities",
    "tenants",
    "users",
    "org_departments",
    "org_members",
    "agent_templates",
    "agents",
    "agent_permissions",
    "mcp_servers",
    "mcp_server_overrides",
    "tools",
    "agent_tools",
    "agent_agent_relationships",
    "participants",
    "skills",
    "skill_files",
    "skill_installs",
    "project_templates",
    "projects",
    "project_repository_operations",
    "project_access_grants",
    "project_member_snapshots",
    "project_capability_bindings",
    "project_work_items",
    "project_runs",
    "project_run_member_snapshots",
    "project_events",
    "chat_sessions",
    "chat_messages",
    "chat_compactions",
    "subagent_runs",
    "workspace_file_revisions",
]


@dataclass
class ProjectApiEnv:
    client: AsyncClient
    db: AsyncSession
    session_factory: Any
    tenant_id: uuid.UUID
    owner_id: uuid.UUID
    viewer_id: uuid.UUID
    source_leader_id: uuid.UUID
    source_worker_id: uuid.UUID
    source_reviewer_id: uuid.UUID
    leader_id: uuid.UUID
    worker_id: uuid.UUID
    reviewer_id: uuid.UUID
    owner: User
    viewer: User
    leader: Agent
    worker: Agent
    reviewer: Agent
    active_user: dict[str, uuid.UUID]
    storage_root: Path

    def authenticate_as(self, user_id: uuid.UUID) -> None:
        self.active_user["value"] = user_id


async def _user(db: AsyncSession, tenant: Tenant, name: str) -> User:
    suffix = uuid.uuid4().hex[:8]
    identity = Identity(
        username=f"{name.lower()}-{suffix}",
        email=f"{name.lower()}-{suffix}@project.test",
        password_hash="test-only",
    )
    db.add(identity)
    await db.flush()
    user = User(
        identity_id=identity.id,
        tenant_id=tenant.id,
        display_name=name,
        role="member",
        is_active=True,
    )
    db.add(user)
    await db.flush()
    return user


async def _agent(db: AsyncSession, tenant: Tenant, owner: User, name: str, role: str) -> Agent:
    agent = Agent(
        name=name,
        role_description=role,
        creator_id=owner.id,
        tenant_id=tenant.id,
        status="running",
        access_mode="private",
        autonomy_policy={"write": "confirm"},
        max_tool_rounds=17,
    )
    db.add(agent)
    await db.flush()
    return agent


@pytest.fixture
async def project_api(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> AsyncIterator[ProjectApiEnv]:
    database_url = os.environ.get("PROJECT_TEST_DATABASE_URL", "sqlite+aiosqlite:///:memory:")
    engine = create_async_engine(database_url)
    table_names = list(TABLES)
    async with engine.begin() as connection:
        if os.environ.get("PROJECT_TEST_DATABASE_URL"):
            await connection.exec_driver_sql("DROP SCHEMA public CASCADE")
            await connection.exec_driver_sql("CREATE SCHEMA public")
        await connection.run_sync(
            lambda sync_connection: Base.metadata.create_all(
                sync_connection,
                tables=[Base.metadata.tables[name] for name in table_names],
            )
        )
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    session = session_factory()
    # Durable subagent helpers intentionally open their own transaction. Point
    # that runtime at the same isolated test database instead of the process
    # default database configured for production.
    monkeypatch.setattr("app.services.subagent_runtime.async_session", session_factory)
    monkeypatch.setattr(
        "app.services.project_group_turn_lifecycle.async_session",
        session_factory,
    )
    monkeypatch.setattr("app.services.project_runtime_tools.async_session", session_factory)
    monkeypatch.setattr("app.api.websocket.async_session", session_factory)
    monkeypatch.setattr("app.services.channel_llm.async_session", session_factory)

    tenant = Tenant(name="Project API", slug=f"project-api-{uuid.uuid4().hex[:8]}")
    session.add(tenant)
    await session.flush()
    tenant_model = LLMModel(
        tenant_id=tenant.id,
        provider="openai",
        model="project-tenant-default",
        api_key_encrypted="test-only",
        label="Project tenant default",
        enabled=True,
        context_window=128000,
    )
    session.add(tenant_model)
    await session.flush()
    tenant.default_model_id = tenant_model.id
    owner = await _user(session, tenant, "Owner")
    viewer = await _user(session, tenant, "Viewer")
    leader = await _agent(session, tenant, owner, "Leader", "Drive outcomes")
    worker = await _agent(session, tenant, owner, "Worker", "Build deliverables")
    reviewer = await _agent(session, tenant, owner, "Reviewer", "Review evidence")
    session.add(
        Tool(
            name="send_message_to_parent",
            display_name="Message Parent Agent",
            description="Return an interim result to the parent session",
            type="builtin",
            category="subagent",
            parameters_schema={"type": "object", "properties": {"message": {"type": "string"}}},
            enabled=True,
            is_default=True,
            source="builtin",
        )
    )
    await session.commit()

    tenant_id = tenant.id
    owner_id = owner.id
    viewer_id = viewer.id
    leader_id = leader.id
    worker_id = worker.id
    reviewer_id = reviewer.id
    active_user = {"value": owner_id}

    async def override_db() -> AsyncIterator[AsyncSession]:
        try:
            yield session
            await session.commit()
        except Exception:
            await session.rollback()
            raise

    async def override_user() -> User:
        # Production auth loads a user on every request. Resolve by immutable
        # ID so an expected authorization rollback cannot leave the fixture
        # holding an expired ORM identity.
        result = await session.execute(select(User).where(User.id == active_user["value"]))
        return result.scalar_one()

    async def skip_external_delivery(_run_id: uuid.UUID) -> None:
        # The endpoint's durable queue/run/event behavior is in scope.  The
        # external agent transport is independently covered by bridge tests.
        return None

    async def skip_external_audit(**_kwargs) -> None:
        return None

    monkeypatch.setattr(
        "app.services.project_git_service.get_settings",
        lambda: SimpleNamespace(STORAGE_LOCAL_ROOT=str(tmp_path)),
    )
    monkeypatch.setattr(
        "app.services.agent_runtime_workspace.get_settings",
        lambda: SimpleNamespace(STORAGE_LOCAL_ROOT=str(tmp_path), AGENT_DATA_DIR=str(tmp_path)),
    )
    monkeypatch.setattr(projects_api, "deliver_project_a2a", skip_external_delivery)
    monkeypatch.setattr(mcp_servers_api, "write_audit_log", skip_external_audit)

    test_app = FastAPI()
    test_app.include_router(projects_api.router, prefix="/api")
    test_app.include_router(files_api.router, prefix="/api")
    test_app.include_router(mcp_servers_api.router, prefix="/api")
    test_app.dependency_overrides[get_db] = override_db
    test_app.dependency_overrides[get_current_user] = override_user

    async with AsyncClient(transport=ASGITransport(app=test_app), base_url="http://project.test") as client:
        yield ProjectApiEnv(
            client=client,
            db=session,
            session_factory=session_factory,
            tenant_id=tenant_id,
            owner_id=owner_id,
            viewer_id=viewer_id,
            source_leader_id=leader_id,
            source_worker_id=worker_id,
            source_reviewer_id=reviewer_id,
            leader_id=leader_id,
            worker_id=worker_id,
            reviewer_id=reviewer_id,
            owner=owner,
            viewer=viewer,
            leader=leader,
            worker=worker,
            reviewer=reviewer,
            active_user=active_user,
            storage_root=tmp_path,
        )

    await session.close()
    await engine.dispose()


async def _create_project(
    env: ProjectApiEnv,
    *,
    name: str = "Closed loop",
    leader_id: uuid.UUID | None = None,
) -> dict[str, Any]:
    response = await env.client.post(
        "/api/projects",
        json={
            "name": name,
            "description": "API acceptance project",
            "goal": "Deliver a traceable result",
            "success_criteria": ["Evidence is committed", "Review is recorded"],
            "members": [
                {"agent_id": str(leader_id or env.source_leader_id), "is_leader": True},
                {"agent_id": str(env.source_worker_id)},
                {"agent_id": str(env.source_reviewer_id)},
            ],
            "capabilities": [
                {
                    "capability_type": "tool",
                    "capability_name": "project-shell",
                    "source": "shared",
                    "scope": {"paths": ["workspace/**"]},
                },
                {
                    "capability_type": "tool",
                    "capability_name": "worker-private-tool",
                    "source": "inherited",
                    "inherited_from_agent_id": str(env.source_worker_id),
                    "scope": {"paths": ["deliverables/**"]},
                },
            ],
            "settings": {
                "git": {"repository_mode": "managed", "branch_policy": "work_item"},
                "runtime": {"mode": "balanced", "monthly_budget": 300},
                "policies": {"approval": "risk", "max_parallel_runs": 3},
            },
        },
    )
    assert response.status_code == 201, response.text
    project = response.json()
    member_ids_by_name = {
        member["agent_name"]: uuid.UUID(member["agent_id"])
        for member in project["members"]
    }
    env.leader_id = member_ids_by_name.get("Leader", env.leader_id)
    env.worker_id = member_ids_by_name.get("Worker", env.worker_id)
    env.reviewer_id = member_ids_by_name.get("Reviewer", env.reviewer_id)
    return project


async def _create_work_item(
    env: ProjectApiEnv,
    project_id: str | uuid.UUID,
    *,
    title: str,
    assignee_agent_id: uuid.UUID | None = None,
    dependency_ids: list[str] | None = None,
) -> dict[str, Any]:
    response = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={
            "title": title,
            "assignee_agent_id": str(assignee_agent_id) if assignee_agent_id else None,
            "dependency_ids": dependency_ids or [],
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


async def _mark_project_running(env: ProjectApiEnv, project_id: str | uuid.UUID) -> None:
    project = await env.db.get(Project, uuid.UUID(str(project_id)))
    assert project is not None
    project.status = "running"
    await env.db.commit()


async def test_project_create_validates_before_managed_storage(project_api: ProjectApiEnv):
    env = project_api
    response = await env.client.post(
        "/api/projects",
        json={"name": "Invalid share", "visibility": "shared", "shared_with_user_ids": []},
    )

    assert response.status_code == 422
    tenant_storage = env.storage_root / "_projects" / str(env.tenant_id)
    assert not tenant_storage.exists() or not any(tenant_storage.iterdir())


async def test_project_create_removes_managed_storage_after_service_failure(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    env = project_api

    async def fail_agent_copy(*_args, **_kwargs):
        raise RuntimeError("forced project Agent failure")

    monkeypatch.setattr("app.services.project_agent_service.create_project_agent", fail_agent_copy)
    with pytest.raises(RuntimeError, match="forced project Agent failure"):
        await env.client.post(
            "/api/projects",
            json={
                "name": "Compensated failure",
                "members": [{"agent_id": str(env.leader_id), "is_leader": True}],
            },
        )

    tenant_storage = env.storage_root / "_projects" / str(env.tenant_id)
    assert not tenant_storage.exists() or not any(tenant_storage.iterdir())
    assert await env.db.scalar(select(func.count()).select_from(Project)) == 0


async def test_project_create_removes_managed_storage_after_commit_failure(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    env = project_api

    async def fail_commit():
        raise RuntimeError("forced commit failure")

    monkeypatch.setattr(env.db, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="forced commit failure"):
        await env.client.post(
            "/api/projects",
            json={
                "name": "Commit failure",
                "members": [{"agent_id": str(env.leader_id), "is_leader": True}],
            },
        )

    tenant_storage = env.storage_root / "_projects" / str(env.tenant_id)
    assert not tenant_storage.exists() or not any(tenant_storage.iterdir())
    assert await env.db.scalar(select(func.count()).select_from(Project)) == 0


async def test_project_create_applies_member_settings_only_to_project_agent(
    project_api: ProjectApiEnv,
):
    from app.models.mcp_server import MCPServerOverride

    env = project_api
    model = LLMModel(
        tenant_id=env.tenant_id,
        provider="openai",
        model="project-member-model",
        api_key_encrypted="test-only",
        label="Project member model",
        enabled=True,
        context_window=128000,
    )
    tool = Tool(
        name=f"project_member_tool_{uuid.uuid4().hex[:8]}",
        display_name="Project member tool",
        description="Project-local test tool",
        type="builtin",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        config_schema={"fields": [{"key": "mode", "type": "text"}]},
        enabled=True,
        is_default=False,
        source="builtin",
        tenant_id=env.tenant_id,
    )
    mcp = MCPServer(
        tenant_id=env.tenant_id,
        name=f"project-member-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Project member MCP",
        base_url_template="https://mcp.project.test",
        headers_template={},
        created_by_user_id=env.viewer_id,
    )
    skill = Skill(
        tenant_id=env.tenant_id,
        name="Project member Skill",
        description="Project-local Skill",
        category="general",
        folder_name=f"project-member-skill-{uuid.uuid4().hex[:8]}",
        visibility="tenant",
        status="published",
        publisher_agent_id=env.source_leader_id,
    )
    env.db.add_all([model, tool, mcp, skill])
    await env.db.flush()
    mcp_tool = Tool(
            name=f"project_member_mcp_tool_{uuid.uuid4().hex[:8]}",
            display_name="Project member MCP tool",
            description="Project-local MCP tool",
            type="mcp",
            category="general",
            parameters_schema={"type": "object", "properties": {}},
            enabled=True,
            is_default=False,
            source="admin",
            tenant_id=env.tenant_id,
            mcp_server_id=mcp.id,
            mcp_server_name=mcp.display_name,
            mcp_tool_name="project_member_action",
    )
    disabled_mcp_tool = Tool(
        name=f"project_member_mcp_disabled_{uuid.uuid4().hex[:8]}",
        display_name="Project member disabled MCP tool",
        description="Second independently selectable MCP tool",
        type="mcp",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        is_default=False,
        source="admin",
        tenant_id=env.tenant_id,
        mcp_server_id=mcp.id,
        mcp_server_name=mcp.display_name,
        mcp_tool_name="project_member_disabled_action",
    )
    env.db.add_all([mcp_tool, disabled_mcp_tool])
    env.db.add(
        SkillFile(
            skill_id=skill.id,
            path="SKILL.md",
            content="---\nname: Project member Skill\ndescription: Project-local Skill\n---\n",
        )
    )
    env.db.add(
        AgentTool(
            agent_id=env.source_leader_id,
            tool_id=tool.id,
            enabled=False,
            config={"mode": "source"},
            source="user_installed",
        )
    )
    await env.db.commit()
    mcp_id = mcp.id

    response = await env.client.post(
        "/api/projects",
        json={
            "name": "Create settings isolation",
            "members": [
                {
                    "agent_id": str(env.source_leader_id),
                    "is_leader": True,
                    "settings": {
                        "config_snapshot": {
                            "primary_model_id": str(model.id),
                            "max_tool_rounds": 33,
                            "project_instruction": "Only for this project",
                        },
                        "tools": [
                            {
                                "tool_id": str(tool.id),
                                "enabled": True,
                                "config": {"mode": "project"},
                            },
                            {
                                "tool_id": str(mcp_tool.id),
                                "enabled": True,
                                "config": {"scope": "project-only"},
                            },
                            {
                                "tool_id": str(disabled_mcp_tool.id),
                                "enabled": False,
                                "config": {},
                            },
                        ],
                        "mcp_server_overrides": [
                            {
                                "server_id": str(mcp.id),
                                "system_prompt_block": "Use the project evidence scope only.",
                                "headers_template": {"X-Project-Scope": "${agent.id}"},
                                "credential_template": "project-only-test-token",
                            }
                        ],
                        "skill_capability_ids": [str(skill.id)],
                    },
                }
            ],
        },
    )
    assert response.status_code == 201, response.text
    project_id = uuid.UUID(response.json()["id"])
    created_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(ProjectMemberSnapshot.project_id == project_id)
        )
    ).scalar_one()
    project_agent_id = created_member.agent_id
    assert created_member.agent_id != env.source_leader_id
    assert created_member.config_snapshot["primary_model_id"] == str(model.id)
    assert created_member.config_snapshot["max_tool_rounds"] == 33
    assert created_member.config_snapshot["project_instruction"] == "Only for this project"

    source_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == env.source_leader_id,
                AgentTool.tool_id == tool.id,
            )
        )
    ).scalar_one()
    project_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == created_member.agent_id,
                AgentTool.tool_id == tool.id,
            )
        )
    ).scalar_one()
    assert source_assignment.enabled is False
    assert source_assignment.config == {"mode": "source"}
    assert project_assignment.enabled is True
    assert project_assignment.config == {"mode": "project"}
    project_mcp_assignments = {
        assignment.tool_id: assignment
        for assignment in (
            await env.db.execute(
                select(AgentTool).where(
                    AgentTool.agent_id == created_member.agent_id,
                    AgentTool.tool_id.in_([mcp_tool.id, disabled_mcp_tool.id]),
                )
            )
        ).scalars()
    }
    assert project_mcp_assignments[mcp_tool.id].enabled is True
    assert project_mcp_assignments[mcp_tool.id].config == {"scope": "project-only"}
    assert project_mcp_assignments[disabled_mcp_tool.id].enabled is False
    project_override = (
        await env.db.execute(
            select(MCPServerOverride).where(
                MCPServerOverride.mcp_server_id == mcp.id,
                MCPServerOverride.scope_type == "agent",
                MCPServerOverride.scope_id == created_member.agent_id,
            )
        )
    ).scalar_one()
    assert project_override.system_prompt_block == "Use the project evidence scope only."
    assert project_override.headers_template == {"X-Project-Scope": "${agent.id}"}
    assert project_override.credential_template == "project-only-test-token"

    source_override = (
        await env.db.execute(
            select(MCPServerOverride).where(
                MCPServerOverride.mcp_server_id == mcp.id,
                MCPServerOverride.scope_type == "agent",
                MCPServerOverride.scope_id == env.source_leader_id,
            )
        )
    ).scalar_one_or_none()
    assert source_override is None
    bindings = (
        await env.db.execute(
            select(ProjectCapabilityBinding).where(
                ProjectCapabilityBinding.project_id == project_id,
                ProjectCapabilityBinding.inherited_from_agent_id == created_member.agent_id,
            )
        )
    ).scalars().all()
    assert {(binding.capability_type, binding.capability_id) for binding in bindings} >= {
        ("mcp", mcp.id),
        ("skill", skill.id),
    }
    mcp_binding = next(
        binding
        for binding in bindings
        if binding.capability_type == "mcp" and binding.capability_id == mcp.id
    )
    assert mcp_binding.is_enabled is True

    disable_selected = await env.client.put(
        f"/api/projects/{project_id}/members/{created_member.id}/tools",
        json=[{"tool_id": str(mcp_tool.id), "enabled": False}],
    )
    assert disable_selected.status_code == 200, disable_selected.text
    await env.db.refresh(mcp_binding)
    assert mcp_binding.is_enabled is False

    enable_other = await env.client.put(
        f"/api/projects/{project_id}/members/{created_member.id}/tools",
        json=[{"tool_id": str(disabled_mcp_tool.id), "enabled": True}],
    )
    assert enable_other.status_code == 200, enable_other.text
    await env.db.refresh(mcp_binding)
    assert mcp_binding.is_enabled is True
    project_mcp_state = {
        assignment.tool_id: assignment.enabled
        for assignment in (
            await env.db.execute(
                select(AgentTool).where(
                    AgentTool.agent_id == created_member.agent_id,
                    AgentTool.tool_id.in_([mcp_tool.id, disabled_mcp_tool.id]),
                )
            )
        ).scalars()
    }
    assert project_mcp_state == {
        mcp_tool.id: False,
        disabled_mcp_tool.id: True,
    }
    source_mcp_assignments = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == env.source_leader_id,
                AgentTool.tool_id.in_([mcp_tool.id, disabled_mcp_tool.id]),
            )
        )
    ).scalars().all()
    assert source_mcp_assignments == []
    invalid_override = await env.client.post(
        "/api/projects",
        json={
            "name": "Reject detached MCP configuration",
            "members": [
                {
                    "agent_id": str(env.source_leader_id),
                    "is_leader": True,
                    "settings": {
                        "tools": [
                            {
                                "tool_id": str(mcp_tool.id),
                                "enabled": False,
                                "config": {},
                            }
                        ],
                        "mcp_server_overrides": [
                            {
                                "server_id": str(mcp.id),
                                "headers_template": {"X-Project-Scope": "blocked"},
                            }
                        ],
                    },
                }
            ],
        },
    )
    assert invalid_override.status_code == 422, invalid_override.text
    assert "requires at least one enabled tool" in invalid_override.json()["detail"]

    # Project editors can maintain the isolated Agent override without gaining
    # authority over the company MCP definition or its shared tool catalog.
    agent_override_url = (
        f"/api/admin/mcp-servers/{mcp_id}/overrides/agent/{project_agent_id}"
    )
    override_view = await env.client.get(
        f"/api/admin/mcp-servers/{mcp_id}/overrides",
        params={"agent_id": str(project_agent_id)},
    )
    assert override_view.status_code == 200, override_view.text
    assert [item["scope_id"] for item in override_view.json()["agent"]] == [
        str(project_agent_id)
    ]
    override_update = await env.client.put(
        agent_override_url,
        json={"headers_template": {"X-Project-Scope": "updated-project-only"}},
    )
    assert override_update.status_code == 200, override_update.text
    assert override_update.json()["headers_template"] == {
        "X-Project-Scope": "updated-project-only"
    }
    global_patch = await env.client.patch(
        f"/api/admin/mcp-servers/{mcp_id}",
        json={"display_name": "Must not change globally"},
    )
    assert global_patch.status_code == 403, global_patch.text
    refresh_global_catalog = await env.client.post(
        f"/api/admin/mcp-servers/{mcp_id}/refresh-tools",
        params={"agent_id": str(project_agent_id)},
    )
    assert refresh_global_catalog.status_code == 403, refresh_global_catalog.text

    detached_mcp = MCPServer(
        tenant_id=env.tenant_id,
        name=f"detached-project-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Detached project MCP",
        base_url_template="https://detached.project.test",
        headers_template={},
        created_by_user_id=env.viewer_id,
    )
    env.db.add(detached_mcp)
    await env.db.commit()
    detached_mcp_id = detached_mcp.id
    detached_override = await env.client.put(
        f"/api/admin/mcp-servers/{detached_mcp_id}/overrides/agent/{project_agent_id}",
        json={"headers_template": {"X-Project-Scope": "forged"}},
    )
    assert detached_override.status_code == 404, detached_override.text
    detached_row = (
        await env.db.execute(
            select(MCPServerOverride).where(
                MCPServerOverride.mcp_server_id == detached_mcp_id,
                MCPServerOverride.scope_type == "agent",
                MCPServerOverride.scope_id == project_agent_id,
            )
        )
    ).scalar_one_or_none()
    assert detached_row is None


async def test_project_private_mcp_credential_is_referenced_only_for_its_creator_and_not_exported(
    project_api: ProjectApiEnv,
):
    from app.models.mcp_server import MCPServerOverride
    from app.services.mcp_server_service import (
        lookup_overrides,
        lookup_project_source_tool_config,
    )
    from app.services.project_agent_template_service import export_project_capabilities_for_template

    env = project_api
    server = MCPServer(
        tenant_id=env.tenant_id,
        name=f"private-project-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Private project MCP",
        base_url_template="https://private.project.test/mcp",
        headers_template={},
        created_by_user_id=env.owner_id,
    )
    env.db.add(server)
    await env.db.flush()
    tool = Tool(
        name=f"private_project_mcp_tool_{uuid.uuid4().hex[:8]}",
        display_name="Private project MCP tool",
        description="Private source tool",
        type="mcp",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        is_default=False,
        source="admin",
        tenant_id=env.tenant_id,
        mcp_server_id=server.id,
        mcp_server_name=server.display_name,
        mcp_tool_name="private_action",
    )
    env.db.add(tool)
    await env.db.flush()
    env.db.add_all(
        [
            AgentTool(
                agent_id=env.source_leader_id,
                tool_id=tool.id,
                enabled=True,
                config={"api_key": "private-agent-tool-secret"},
                source="user_installed",
                installed_by_agent_id=env.source_leader_id,
            ),
            MCPServerOverride(
                mcp_server_id=server.id,
                scope_type="agent",
                scope_id=env.source_leader_id,
                credential_template="private-source-secret",
                headers_template={"Authorization": "Bearer private-source-secret"},
                last_modified_by_user_id=env.owner_id,
            ),
        ]
    )
    await env.db.commit()

    created = await env.client.post(
        "/api/projects",
        json={
            "name": "Private MCP reference",
            "members": [
                {
                    "agent_id": str(env.source_leader_id),
                    "is_leader": True,
                    "settings": {
                        "tools": [
                            {
                                "tool_id": str(tool.id),
                                "enabled": True,
                                "config": {
                                    "scope": "project-only",
                                    "api_key": "submitted-project-secret",
                                    "token": "submitted-token",
                                    "auth": "submitted-auth",
                                    "headers": {"Authorization": "submitted-header-secret"},
                                },
                            }
                        ],
                    },
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    project_id = uuid.UUID(created.json()["id"])
    project = await env.db.get(Project, project_id)
    member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(ProjectMemberSnapshot.project_id == project_id)
        )
    ).scalar_one()
    project_override = (
        await env.db.execute(
            select(MCPServerOverride).where(
                MCPServerOverride.mcp_server_id == server.id,
                MCPServerOverride.scope_type == "agent",
                MCPServerOverride.scope_id == member.agent_id,
            )
        )
    ).scalar_one_or_none()
    assert project_override is None
    project_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == member.agent_id,
                AgentTool.tool_id == tool.id,
            )
        )
    ).scalar_one()
    assert project_assignment.config == {"scope": "project-only"}
    assert await lookup_project_source_tool_config(
        env.db,
        project_agent_id=member.agent_id,
        tool_id=tool.id,
        execution_user_id=env.owner_id,
    ) == {"api_key": "private-agent-tool-secret"}
    assert await lookup_project_source_tool_config(
        env.db,
        project_agent_id=member.agent_id,
        tool_id=tool.id,
        execution_user_id=env.viewer_id,
    ) == {}
    env.db.add(
        MCPServerOverride(
            mcp_server_id=server.id,
            scope_type="agent",
            scope_id=member.agent_id,
            credential_template="legacy-copied-secret",
            headers_template={"Authorization": "Bearer legacy-copied-secret"},
            last_modified_by_user_id=env.owner_id,
        )
    )
    await env.db.commit()

    _tenant_override, creator_override = await lookup_overrides(
        env.db,
        server.id,
        env.tenant_id,
        member.agent_id,
        execution_user_id=env.owner_id,
        allow_project_source_reference=True,
    )
    assert creator_override is not None
    assert creator_override.scope_id == env.source_leader_id
    assert creator_override.credential_template == "private-source-secret"

    _tenant_override, other_override = await lookup_overrides(
        env.db,
        server.id,
        env.tenant_id,
        member.agent_id,
        execution_user_id=env.viewer_id,
        allow_project_source_reference=True,
    )
    assert other_override is None

    exported = await export_project_capabilities_for_template(env.db, project)
    assert any(item.get("capability_id") == str(server.id) for item in exported)
    assert "private-source-secret" not in json.dumps(exported, ensure_ascii=False)
    assert "legacy-copied-secret" not in json.dumps(exported, ensure_ascii=False)
    assert "private-agent-tool-secret" not in json.dumps(exported, ensure_ascii=False)
    assert "submitted-project-secret" not in json.dumps(exported, ensure_ascii=False)


async def test_project_rejects_copying_private_mcp_credentials(project_api: ProjectApiEnv):
    env = project_api
    server = MCPServer(
        tenant_id=env.tenant_id,
        name=f"private-copy-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Private copy MCP",
        base_url_template="https://private-copy.project.test/mcp",
        headers_template={},
        created_by_user_id=env.owner_id,
    )
    env.db.add(server)
    await env.db.flush()
    tool = Tool(
        name=f"private_copy_mcp_tool_{uuid.uuid4().hex[:8]}",
        display_name="Private copy MCP tool",
        description="Private source tool",
        type="mcp",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        is_default=False,
        source="agent",
        tenant_id=env.tenant_id,
        mcp_server_id=server.id,
        mcp_server_name=server.display_name,
        mcp_tool_name="private_copy_action",
    )
    env.db.add(tool)
    await env.db.flush()
    env.db.add(
        AgentTool(
            agent_id=env.source_leader_id,
            tool_id=tool.id,
            enabled=True,
            source="user_installed",
            installed_by_agent_id=env.source_leader_id,
        )
    )
    await env.db.commit()

    response = await env.client.post(
        "/api/projects",
        json={
            "name": "Reject private MCP copy",
            "members": [
                {
                    "agent_id": str(env.source_leader_id),
                    "is_leader": True,
                    "settings": {
                        "tools": [{"tool_id": str(tool.id), "enabled": True, "config": {}}],
                        "mcp_server_overrides": [
                            {
                                "server_id": str(server.id),
                                "credential_template": "must-not-copy",
                            }
                        ],
                    },
                }
            ],
        },
    )
    assert response.status_code == 422, response.text
    assert "cannot be copied" in response.json()["detail"]


async def test_project_capability_options_are_the_create_authorization_boundary(
    project_api: ProjectApiEnv,
):
    env = project_api
    shared_mcp = MCPServer(
        tenant_id=env.tenant_id,
        name=f"shared-project-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Shared Project MCP",
        base_url_template="https://shared.project.test/mcp",
        headers_template={},
    )
    worker_mcp = MCPServer(
        tenant_id=env.tenant_id,
        name=f"worker-project-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Worker Project MCP",
        base_url_template="https://worker.project.test/mcp",
        headers_template={},
    )
    market_skill = Skill(
        tenant_id=env.tenant_id,
        name="Market Project Skill",
        description="A published project Skill.",
        category="general",
        folder_name=f"market-project-skill-{uuid.uuid4().hex[:8]}",
        visibility="public",
        status="published",
    )
    worker_skill = Skill(
        tenant_id=env.tenant_id,
        name="Worker Project Skill",
        description="A Skill installed for one digital employee.",
        category="general",
        folder_name=f"worker-project-skill-{uuid.uuid4().hex[:8]}",
        visibility="tenant",
        status="published",
        publisher_agent_id=env.source_reviewer_id,
    )
    draft_market_skill = Skill(
        tenant_id=env.tenant_id,
        name="Draft Market Project Skill",
        description="Not published.",
        category="general",
        folder_name=f"draft-market-project-skill-{uuid.uuid4().hex[:8]}",
        visibility="public",
        status="draft",
    )
    offline_worker_skill = Skill(
        tenant_id=env.tenant_id,
        name="Offline Worker Project Skill",
        description="No longer available.",
        category="general",
        folder_name=f"offline-worker-project-skill-{uuid.uuid4().hex[:8]}",
        visibility="tenant",
        status="offline",
        publisher_agent_id=env.source_worker_id,
    )
    draft_worker_skill = Skill(
        tenant_id=env.tenant_id,
        name="Draft Worker Project Skill",
        description="Not ready for project use.",
        category="general",
        folder_name=f"draft-worker-project-skill-{uuid.uuid4().hex[:8]}",
        visibility="tenant",
        status="draft",
        publisher_agent_id=env.source_worker_id,
    )
    env.db.add_all(
        [
            shared_mcp,
            worker_mcp,
            market_skill,
            worker_skill,
            draft_market_skill,
            offline_worker_skill,
            draft_worker_skill,
        ]
    )
    await env.db.flush()
    shared_tool = Tool(
        name=f"shared_project_mcp_tool_{uuid.uuid4().hex[:8]}",
        display_name="Shared project action",
        description="Perform an approved shared action.",
        type="mcp",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="admin",
        tenant_id=env.tenant_id,
        mcp_server_id=shared_mcp.id,
    )
    worker_tool = Tool(
        name=f"worker_project_mcp_tool_{uuid.uuid4().hex[:8]}",
        display_name="Worker project action",
        description="Perform an action installed by one digital employee.",
        type="mcp",
        category="general",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="agent",
        tenant_id=env.tenant_id,
        mcp_server_id=worker_mcp.id,
    )
    env.db.add_all([shared_tool, worker_tool])
    await env.db.flush()
    env.db.add_all(
        [
            AgentTool(
                agent_id=env.source_worker_id,
                tool_id=worker_tool.id,
                enabled=True,
                source="user_installed",
            ),
            SkillInstall(
                tenant_id=env.tenant_id,
                skill_id=worker_skill.id,
                agent_id=env.source_worker_id,
                installed_by_user_id=env.owner_id,
                is_active=True,
            ),
            SkillFile(
                skill_id=market_skill.id,
                path="SKILL.md",
                content="---\nname: Market Project Skill\ndescription: Published\n---\n",
            ),
            SkillFile(
                skill_id=worker_skill.id,
                path="SKILL.md",
                content="---\nname: Worker Project Skill\ndescription: Installed\n---\n",
            ),
        ]
    )
    await env.db.commit()
    shared_mcp_id = shared_mcp.id
    worker_mcp_id = worker_mcp.id
    worker_tool_id = worker_tool.id
    market_skill_id = market_skill.id
    worker_skill_id = worker_skill.id
    draft_market_skill_id = draft_market_skill.id
    offline_worker_skill_id = offline_worker_skill.id
    draft_worker_skill_id = draft_worker_skill.id

    bootstrap = await env.client.get("/api/projects/bootstrap-options")
    assert bootstrap.status_code == 200, bootstrap.text
    bootstrap_payload = bootstrap.json()
    assert "skills" not in bootstrap_payload
    assert "mcp_servers" not in bootstrap_payload
    capabilities = bootstrap_payload["capabilities"]
    tool_options = bootstrap_payload["tools"]
    shared_mcp_tool_option = next(
        item for item in tool_options if item["id"] == str(shared_tool.id)
    )
    worker_mcp_tool_option = next(
        item
        for item in tool_options
        if item["id"] == str(worker_tool_id)
        and item["installed_by_agent_id"] == str(env.source_worker_id)
    )
    worker_skill_option = next(
        item
        for item in capabilities
        if item["capability_id"] == str(worker_skill_id)
        and item["owner_agent_id"] == str(env.source_worker_id)
    )
    assert shared_mcp_tool_option["source"] == "admin"
    assert shared_mcp_tool_option["installed_by_agent_id"] is None
    assert shared_mcp_tool_option["enabled"] is False
    assert worker_mcp_tool_option["source"] == "agent"
    assert worker_mcp_tool_option["agent_tool_source"] == "user_installed"
    assert worker_mcp_tool_option["enabled"] is False
    assert all(item["type"] != "mcp" for item in capabilities)
    assert worker_skill_option["owner_agent_id"] == str(env.source_worker_id)
    assert all(item["capability_id"] != str(draft_market_skill_id) for item in capabilities)
    assert all(item["capability_id"] != str(offline_worker_skill_id) for item in capabilities)
    assert all(item["capability_id"] != str(draft_worker_skill_id) for item in capabilities)
    assert all("key" not in item for item in capabilities)

    async def create_with(
        *,
        member_id: uuid.UUID,
        mcp_ids: list[uuid.UUID] | None = None,
        skill_ids: list[uuid.UUID] | None = None,
        tool_ids: list[uuid.UUID] | None = None,
    ):
        return await env.client.post(
            "/api/projects",
            json={
                "name": f"Capability boundary {uuid.uuid4().hex[:8]}",
                "members": [
                    {
                        "agent_id": str(member_id),
                        "is_leader": True,
                        "settings": {
                            "tools": [
                                {"tool_id": str(value), "enabled": True}
                                for value in (tool_ids or [])
                            ],
                            "mcp_capability_ids": [str(value) for value in (mcp_ids or [])],
                            "skill_capability_ids": [str(value) for value in (skill_ids or [])],
                        },
                    }
                ],
            },
        )

    before_projects = await env.db.scalar(select(func.count()).select_from(Project))
    before_agents = await env.db.scalar(select(func.count()).select_from(Agent))
    before_bindings = await env.db.scalar(
        select(func.count()).select_from(ProjectCapabilityBinding)
    )
    rejected = [
        await create_with(member_id=env.source_leader_id, tool_ids=[worker_tool_id]),
        await create_with(member_id=env.source_leader_id, mcp_ids=[worker_mcp_id]),
        await create_with(member_id=env.source_leader_id, skill_ids=[worker_skill_id]),
        await create_with(member_id=env.source_leader_id, skill_ids=[draft_market_skill_id]),
        await create_with(member_id=env.source_worker_id, skill_ids=[offline_worker_skill_id]),
        await create_with(member_id=env.source_worker_id, skill_ids=[draft_worker_skill_id]),
        await env.client.post(
            "/api/projects",
            json={
                "name": f"Forged inherited capability {uuid.uuid4().hex[:8]}",
                "members": [{"agent_id": str(env.source_leader_id), "is_leader": True}],
                "capabilities": [
                    {
                        "capability_type": "skill",
                        "capability_id": str(worker_skill_id),
                        "source": "inherited",
                        "inherited_from_agent_id": str(env.source_leader_id),
                    }
                ],
            },
        ),
        await env.client.post(
            "/api/projects",
            json={
                "name": f"Forged shared capability {uuid.uuid4().hex[:8]}",
                "members": [{"agent_id": str(env.source_leader_id), "is_leader": True}],
                "shared_capability_ids": [str(worker_skill_id)],
            },
        ),
    ]
    assert [response.status_code for response in rejected] == [422] * 8
    assert await env.db.scalar(select(func.count()).select_from(Project)) == before_projects
    assert await env.db.scalar(select(func.count()).select_from(Agent)) == before_agents
    assert (
        await env.db.scalar(select(func.count()).select_from(ProjectCapabilityBinding)
        ) == before_bindings
    )
    tenant_storage = env.storage_root / "_projects" / str(env.tenant_id)
    assert not tenant_storage.exists() or not any(tenant_storage.iterdir())

    leader_project_response = await create_with(member_id=env.source_leader_id)
    assert leader_project_response.status_code == 201, leader_project_response.text
    leader_project_id = uuid.UUID(leader_project_response.json()["id"])
    leader_project_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == leader_project_id
            )
        )
    ).scalar_one()
    leader_project_agent_id = leader_project_member.agent_id
    binding_count = await env.db.scalar(
        select(func.count()).select_from(ProjectCapabilityBinding)
    )
    cross_member_adds = [
        await env.client.post(
            f"/api/projects/{leader_project_id}/capabilities",
            json={
                "capability_type": capability_type,
                "capability_id": str(capability_id),
                "source": "inherited",
                "inherited_from_agent_id": str(leader_project_agent_id),
            },
        )
        for capability_type, capability_id in (
            ("mcp", worker_mcp_id),
            ("skill", worker_skill_id),
        )
    ]
    assert [response.status_code for response in cross_member_adds] == [422, 422]
    cross_member_import = await env.client.post(
        f"/api/agents/{leader_project_agent_id}/files/import-skill",
        json={"skill_id": str(worker_skill_id)},
    )
    assert cross_member_import.status_code == 422
    assert (
        await env.db.scalar(select(func.count()).select_from(ProjectCapabilityBinding))
        == binding_count
    )

    allowed = await create_with(
        member_id=env.source_worker_id,
        mcp_ids=[shared_mcp_id, worker_mcp_id],
        skill_ids=[market_skill_id, worker_skill_id],
    )
    assert allowed.status_code == 201, allowed.text
    project_id = uuid.UUID(allowed.json()["id"])
    project_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(ProjectMemberSnapshot.project_id == project_id)
        )
    ).scalar_one()
    bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == project_id,
                    ProjectCapabilityBinding.inherited_from_agent_id == project_member.agent_id,
                )
            )
        ).scalars()
    )
    assert {(binding.capability_type, binding.capability_id) for binding in bindings} >= {
        ("mcp", shared_mcp_id),
        ("mcp", worker_mcp_id),
        ("skill", market_skill_id),
        ("skill", worker_skill_id),
    }
    source_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == env.source_worker_id,
                AgentTool.tool_id == worker_tool_id,
            )
        )
    ).scalar_one()
    source_install = (
        await env.db.execute(
            select(SkillInstall).where(
                SkillInstall.agent_id == env.source_worker_id,
                SkillInstall.skill_id == worker_skill_id,
            )
        )
    ).scalar_one()
    assert source_assignment.enabled is True
    assert source_install.is_active is True


async def test_private_share_settings_and_audit_are_a_real_api_round_trip(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Private by default")
    project_id = project["id"]

    assert project["visibility"] == "private"
    assert project["shared_with"] == []
    assert project["status"] == "planning"
    assert project["access_role"] == "owner"
    assert project["execution_user_id"] == str(env.owner_id)
    assert project["execution_user_name"] == "Owner"

    env.authenticate_as(env.viewer_id)
    hidden = await env.client.get(f"/api/projects/{project_id}")
    assert hidden.status_code == 404

    env.authenticate_as(env.owner_id)
    settings_response = await env.client.get(f"/api/projects/{project_id}/settings")
    assert settings_response.status_code == 200, settings_response.text
    settings = settings_response.json()
    assert settings["policies"]["approval"] == "risk"
    assert settings["runtime"]["monthly_budget"] == 300

    saved = await env.client.patch(
        f"/api/projects/{project_id}/settings",
        json={
            "policies": {"approval": "all_writes", "max_parallel_runs": 2},
            "runtime": {"mode": "quality", "monthly_budget": 450},
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["policies"] == {"approval": "all_writes", "max_parallel_runs": 2}
    assert saved.json()["runtime"]["monthly_budget"] == 450
    rejected_mixed_settings = await env.client.patch(
        f"/api/projects/{project_id}/settings",
        json={
            "runtime": {"mode": "must-not-save", "monthly_budget": 999},
            "git": {"repository_mode": "external"},
        },
    )
    assert rejected_mixed_settings.status_code == 422
    settings_after_rejection = (await env.client.get(f"/api/projects/{project_id}/settings")).json()
    assert settings_after_rejection["runtime"] == saved.json()["runtime"]

    shared = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert shared.status_code == 200, shared.text
    assert shared.json()["visibility"] == "shared"
    assert shared.json()["execution_user_id"] == str(env.viewer_id)
    assert shared.json()["execution_user_name"] == "Viewer"
    assert shared.json()["shared_with_user_ids"] == [str(env.viewer_id)]
    assert [entry["user_id"] for entry in shared.json()["shared_with"]] == [str(env.viewer_id)]

    env.authenticate_as(env.viewer_id)
    visible = await env.client.get(f"/api/projects/{project_id}")
    assert visible.status_code == 200
    assert visible.json()["access_role"] == "view"
    assert (await env.client.get(f"/api/projects/{project_id}/settings")).status_code == 200
    viewer_settings_write = await env.client.patch(
        f"/api/projects/{project_id}/settings",
        json={"runtime": {"mode": "viewer-must-not-save"}},
    )
    assert viewer_settings_write.status_code == 404
    cannot_edit = await env.client.patch(f"/api/projects/{project_id}", json={"name": "Not allowed"})
    assert cannot_edit.status_code == 404

    grant = (
        await env.db.execute(
            select(ProjectAccessGrant).where(
                ProjectAccessGrant.project_id == uuid.UUID(project_id),
                ProjectAccessGrant.user_id == env.viewer_id,
            )
        )
    ).scalar_one()
    grant.role = "edit"
    await env.db.commit()
    editor_contract = await env.client.get(f"/api/projects/{project_id}")
    assert editor_contract.status_code == 200
    assert editor_contract.json()["access_role"] == "edit"
    normal_editor_update = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"description": "Editors may update ordinary project fields"},
    )
    assert normal_editor_update.status_code == 200
    editor_settings_update = await env.client.patch(
        f"/api/projects/{project_id}/settings",
        json={"current_signal": "Editor-visible delivery signal"},
    )
    assert editor_settings_update.status_code == 200
    editor_cannot_unshare = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "private"},
    )
    assert editor_cannot_unshare.status_code == 404
    editor_cannot_reshare = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"shared_with_user_ids": [str(env.viewer_id)]},
    )
    assert editor_cannot_reshare.status_code == 404

    env.authenticate_as(env.owner_id)
    retained_editor = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert retained_editor.status_code == 200, retained_editor.text
    assert retained_editor.json()["shared_with"] == [
        {
            "user_id": str(env.viewer_id),
            "display_name": "Viewer",
            "role": "edit",
        }
    ]
    env.authenticate_as(env.viewer_id)
    assert (await env.client.patch(f"/api/projects/{project_id}", json={"description": "Still editable"})).status_code == 200

    env.authenticate_as(env.owner_id)
    empty_shared = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "shared", "shared_with_user_ids": []},
    )
    assert empty_shared.status_code == 422

    outsider_tenant = Tenant(name="Outsider", slug=f"outsider-{uuid.uuid4().hex[:8]}")
    env.db.add(outsider_tenant)
    await env.db.flush()
    outsider = await _user(env.db, outsider_tenant, "Outsider")
    await env.db.commit()
    cross_tenant_share = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "shared", "shared_with_user_ids": [str(outsider.id)]},
    )
    assert cross_tenant_share.status_code == 422

    private = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "private"},
    )
    assert private.status_code == 200
    assert private.json()["visibility"] == "private"
    assert private.json()["shared_with"] == []
    assert private.json()["execution_user_id"] == str(env.owner_id)
    assert private.json()["execution_user_name"] == "Owner"
    remaining_grants = (
        (await env.db.execute(select(ProjectAccessGrant).where(ProjectAccessGrant.project_id == uuid.UUID(project_id))))
        .scalars()
        .all()
    )
    assert remaining_grants == []
    env.authenticate_as(env.viewer_id)
    assert (await env.client.get(f"/api/projects/{project_id}")).status_code == 404
    assert (await env.client.get(f"/api/projects/{project_id}/settings")).status_code == 404
    env.authenticate_as(env.owner_id)

    events_response = await env.client.get(f"/api/projects/{project_id}/events")
    assert events_response.status_code == 200
    event_types = {event["event_type"] for event in events_response.json()}
    assert {"project.created", "project.initialized", "project.settings.updated", "project.updated"} <= event_types


async def test_shared_project_execution_user_is_validated_atomically(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Shared execution identity")
    project_id = project["id"]

    missing_selection = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "shared", "shared_with_user_ids": [str(env.viewer_id)]},
    )
    assert missing_selection.status_code == 422
    unchanged = (await env.client.get(f"/api/projects/{project_id}")).json()
    assert unchanged["visibility"] == "private"
    assert unchanged["shared_with"] == []
    assert unchanged["execution_user_id"] == str(env.owner_id)

    configured = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert configured.status_code == 200, configured.text

    stored = await env.db.get(Project, uuid.UUID(project_id))
    assert stored is not None
    stored.execution_user_id = None
    await env.db.commit()
    legacy_acl_change = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"shared_with_user_ids": [str(env.viewer_id)]},
    )
    assert legacy_acl_change.status_code == 422
    legacy_summary = (await env.client.get(f"/api/projects/{project_id}")).json()
    assert legacy_summary["execution_user_id"] == str(env.owner_id)
    reconfigured = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert reconfigured.status_code == 200, reconfigured.text

    tenant = await env.db.get(Tenant, env.tenant_id)
    assert tenant is not None
    alternate = await _user(env.db, tenant, "Alternate")
    await env.db.commit()
    invalid_replacement = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"shared_with_user_ids": [str(alternate.id)]},
    )
    assert invalid_replacement.status_code == 422
    still_configured = (await env.client.get(f"/api/projects/{project_id}")).json()
    assert [entry["user_id"] for entry in still_configured["shared_with"]] == [str(env.viewer_id)]
    assert still_configured["execution_user_id"] == str(env.viewer_id)

    alternate.is_active = False
    await env.db.commit()
    inactive_selection = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "shared_with_user_ids": [str(env.viewer_id), str(alternate.id)],
            "execution_user_id": str(alternate.id),
        },
    )
    assert inactive_selection.status_code == 422
    alternate.is_active = True
    await env.db.commit()

    replaced = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "shared_with_user_ids": [str(alternate.id)],
            "execution_user_id": str(alternate.id),
        },
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["execution_user_id"] == str(alternate.id)
    assert replaced.json()["execution_user_name"] == "Alternate"


async def test_project_owner_directory_is_tenant_scoped_and_excludes_owner(
    project_api: ProjectApiEnv,
):
    env = project_api
    project = await _create_project(env, name="Project share directory")
    project_id = project["id"]

    root = OrgDepartment(
        tenant_id=env.tenant_id,
        name="Product",
        path="Product",
        status="active",
    )
    env.db.add(root)
    await env.db.flush()
    child = OrgDepartment(
        tenant_id=env.tenant_id,
        name="Research",
        path="Product/Research",
        parent_id=root.id,
        status="active",
    )
    env.db.add(child)
    await env.db.flush()
    env.db.add_all(
        [
            OrgMember(
                tenant_id=env.tenant_id,
                user_id=env.owner_id,
                name="Owner",
                email="owner@project.test",
                department_id=root.id,
                department_path=root.path,
                status="active",
            ),
            OrgMember(
                tenant_id=env.tenant_id,
                user_id=env.viewer_id,
                name="Viewer",
                email="viewer@project.test",
                department_id=child.id,
                department_path=child.path,
                status="active",
            ),
        ]
    )
    outsider_tenant = Tenant(
        name="Directory outsider",
        slug=f"directory-outsider-{uuid.uuid4().hex[:8]}",
    )
    env.db.add(outsider_tenant)
    await env.db.flush()
    outsider = await _user(env.db, outsider_tenant, "DirectoryOutsider")
    outsider_department = OrgDepartment(
        tenant_id=outsider_tenant.id,
        name="Other company",
        path="Other company",
        status="active",
    )
    env.db.add(outsider_department)
    await env.db.flush()
    env.db.add(
        OrgMember(
            tenant_id=outsider_tenant.id,
            user_id=outsider.id,
            name="Directory outsider",
            department_id=outsider_department.id,
            department_path=outsider_department.path,
            status="active",
        )
    )
    await env.db.commit()

    departments = await env.client.get(
        f"/api/projects/{project_id}/directory/departments"
    )
    assert departments.status_code == 200, departments.text
    assert [item["id"] for item in departments.json()["items"]] == [str(root.id)]
    assert departments.json()["items"][0]["has_children"] is True
    assert departments.json()["my_department"]["id"] == str(root.id)

    members = await env.client.get(
        f"/api/projects/{project_id}/directory/members",
        params={"department_id": str(root.id), "include_descendants": "true"},
    )
    assert members.status_code == 200, members.text
    assert [item["id"] for item in members.json()["items"]] == [
        str(env.viewer_id)
    ]
    assert str(env.owner_id) not in {item["id"] for item in members.json()["items"]}
    assert str(outsider.id) not in {item["id"] for item in members.json()["items"]}

    cross_tenant_department = await env.client.get(
        f"/api/projects/{project_id}/directory/members",
        params={"department_id": str(outsider_department.id)},
    )
    assert cross_tenant_department.status_code == 404

    env.authenticate_as(env.viewer_id)
    denied_departments = await env.client.get(
        f"/api/projects/{project_id}/directory/departments"
    )
    denied_members = await env.client.get(
        f"/api/projects/{project_id}/directory/members",
        params={"search": "Owner"},
    )
    assert denied_departments.status_code == 404
    assert denied_members.status_code == 404


async def test_project_owner_directory_lists_active_users_without_synced_org_profiles(
    project_api: ProjectApiEnv,
):
    env = project_api
    project = await _create_project(env, name="Project share without org sync")

    departments = await env.client.get(
        f"/api/projects/{project['id']}/directory/departments"
    )
    assert departments.status_code == 200, departments.text
    assert departments.json() == {"items": [], "my_department": None}

    members = await env.client.get(
        f"/api/projects/{project['id']}/directory/members"
    )
    assert members.status_code == 200, members.text
    assert members.json()["total"] == 1
    assert members.json()["items"] == [
        {
            "id": str(env.viewer_id),
            "member_id": None,
            "name": "Viewer",
            "nickname": None,
            "department_id": None,
            "department_path": "",
            "title": "",
            "avatar_url": None,
            "email": env.viewer.email,
        }
    ]
    assert str(env.owner_id) not in {
        item["id"] for item in members.json()["items"]
    }


async def test_member_and_run_snapshots_are_isolated_and_a2a_bypasses_leader(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Snapshots and mesh")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    worker_id = env.worker_id
    reviewer_id = env.reviewer_id
    source_before = {
        "name": env.worker.name,
        "role_description": env.worker.role_description,
        "autonomy_policy": dict(env.worker.autonomy_policy or {}),
        "max_tool_rounds": env.worker.max_tool_rounds,
    }
    review_item = await _create_work_item(
        env,
        project_id,
        title="Review acceptance evidence",
        assignee_agent_id=reviewer_id,
    )

    members_response = await env.client.get(f"/api/projects/{project_id}/members")
    assert members_response.status_code == 200
    members = members_response.json()
    worker_snapshot = next(member for member in members if member["agent_id"] == str(worker_id))
    changed_snapshot = {
        **worker_snapshot["config_snapshot"],
        "autonomy_policy": {"write": "allow-in-project"},
        "project_instruction": "Only exists in this project",
    }
    patched_member = await env.client.patch(
        f"/api/projects/{project_id}/members/{worker_snapshot['id']}",
        json={"config_snapshot": changed_snapshot},
    )
    assert patched_member.status_code == 200, patched_member.text
    assert patched_member.json()["config_snapshot"]["project_instruction"] == "Only exists in this project"

    env.db.expire(env.worker)
    source_worker = await env.db.get(Agent, worker_id)
    assert source_worker is not None
    assert {
        "name": source_worker.name,
        "role_description": source_worker.role_description,
        "autonomy_policy": dict(source_worker.autonomy_policy or {}),
        "max_tool_rounds": source_worker.max_tool_rounds,
    } == source_before

    # Worker wakes Reviewer directly. Neither side is the project Leader and
    # no global Agent relationship is created or required.
    a2a_response = await env.client.post(
        f"/api/projects/{project_id}/a2a",
        json={
            "from_agent_id": str(worker_id),
            "to_agent_id": str(reviewer_id),
            "title": "Review acceptance evidence",
            "message": "Review the acceptance evidence directly",
            "mode": "review",
            "expected_output": "A cited pass/fail review decision",
            "work_item_id": review_item["id"],
        },
    )
    assert a2a_response.status_code == 202, a2a_response.text
    assert a2a_response.json()["from_agent_id"] == str(worker_id)
    assert a2a_response.json()["to_agent_id"] == str(reviewer_id)

    run_response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(worker_id), "trigger_type": "manual", "input": {"objective": "Build v1"}},
    )
    assert run_response.status_code == 201, run_response.text
    run_id = run_response.json()["id"]
    assert run_response.json()["input"]["title"] == "Build v1"
    assert run_response.json()["output"]["subagent_session_id"]

    frozen_response = await env.client.get(f"/api/projects/{project_id}/runs/{run_id}/member-snapshots")
    assert frozen_response.status_code == 200
    frozen = frozen_response.json()
    assert len(frozen) == 3
    assert sum(item["is_leader"] for item in frozen) == 1
    worker_frozen = next(item for item in frozen if item["agent_id"] == str(worker_id))
    reviewer_frozen = next(item for item in frozen if item["agent_id"] == str(reviewer_id))
    assert {item["name"] for item in worker_frozen["member_snapshot"]["capabilities"]["items"]} == {
        "project-shell",
        "worker-private-tool",
    }
    assert {item["name"] for item in reviewer_frozen["member_snapshot"]["capabilities"]["items"]} == {
        "project-shell"
    }
    assert "member_config_snapshot" not in worker_frozen
    assert "capability_snapshot" not in worker_frozen

    capabilities = (await env.client.get(f"/api/projects/{project_id}/capabilities")).json()
    shared_capability = next(item for item in capabilities if item["capability_name"] == "project-shell")
    disabled = await env.client.patch(
        f"/api/projects/{project_id}/capabilities/{shared_capability['id']}",
        json={"is_enabled": False},
    )
    assert disabled.status_code == 200
    frozen_again = (await env.client.get(f"/api/projects/{project_id}/runs/{run_id}/member-snapshots")).json()
    assert frozen_again == frozen

    events = (await env.client.get(f"/api/projects/{project_id}/events")).json()
    a2a_event = next(event for event in events if event["event_type"] == "a2a.queued")
    assert a2a_event["from_agent_id"] == str(worker_id)
    assert a2a_event["to_agent_id"] == str(reviewer_id)
    assert {"run.queued", "capability.updated", "member.snapshot.updated"} <= {event["event_type"] for event in events}


async def test_run_event_and_dashboard_expose_frozen_product_summary(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Product snapshot summary")
    project_id = project["id"]
    await _mark_project_running(env, project_id)

    members = (await env.client.get(f"/api/projects/{project_id}/members")).json()
    worker = next(member for member in members if member["agent_id"] == str(env.worker_id))
    model = (
        await env.db.execute(select(LLMModel).where(LLMModel.tenant_id == env.tenant_id))
    ).scalar_one()
    updated_config = {
        **worker["config_snapshot"],
        "primary_model_id": str(model.id),
        "max_tool_rounds": 23,
        "project_instruction": "Use the project acceptance criteria.",
    }
    patched = await env.client.patch(
        f"/api/projects/{project_id}/members/{worker['id']}",
        json={"config_snapshot": updated_config},
    )
    assert patched.status_code == 200, patched.text

    created = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(env.worker_id), "input": {"objective": "Produce the delivery"}},
    )
    assert created.status_code == 201, created.text
    run = created.json()
    summary = run["member_snapshot"]
    assert summary["name"] == "Worker"
    assert summary["configuration"] == {
        "primary_model": {"name": "Project tenant default", "availability": "available"},
        "fallback_model": None,
        "max_tool_rounds": 23,
        "has_project_instruction": True,
    }
    assert summary["capabilities"]["total"] == len(summary["capabilities"]["items"])
    assert {item["source"] for item in summary["capabilities"]["items"]} <= {"project", "member"}
    assert "member_config_snapshot" not in summary
    assert "capability_snapshot" not in summary

    snapshots = (
        await env.client.get(f"/api/projects/{project_id}/runs/{run['id']}/member-snapshots")
    ).json()
    worker_snapshot = next(item for item in snapshots if item["agent_id"] == str(env.worker_id))
    assert worker_snapshot["member_snapshot"] == summary
    assert "member_config_snapshot" not in worker_snapshot
    assert "capability_snapshot" not in worker_snapshot

    events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
    queued_event = next(event for event in events if event["event_type"] == "run.queued")
    assert queued_event["member_snapshot"] == summary

    dashboard = (await env.client.get(f"/api/projects/{project_id}/dashboard")).json()
    dashboard_run = next(item for item in dashboard["runs"] if item["id"] == run["id"])
    dashboard_event = next(item for item in dashboard["events"] if item["event_type"] == "run.queued")
    assert dashboard_run["member_snapshot"] == summary
    assert dashboard_event["member_snapshot"] == summary


async def test_rest_a2a_requires_action_scope_and_ready_dependencies(project_api: ProjectApiEnv):
    from app.models.project import ProjectMemberSnapshot, ProjectRun

    env = project_api
    project = await _create_project(env, name="Actionable REST A2A")
    await _mark_project_running(env, project["id"])
    parent = await _create_work_item(
        env,
        project["id"],
        title="Produce reviewed input evidence",
        assignee_agent_id=env.worker_id,
    )
    child = await _create_work_item(
        env,
        project["id"],
        title="Make release decision",
        assignee_agent_id=env.reviewer_id,
        dependency_ids=[parent["id"]],
    )
    request = {
        "from_agent_id": str(env.worker_id),
        "to_agent_id": str(env.reviewer_id),
        "title": "Make release decision",
        "message": "Evaluate the reviewed input evidence against the release criteria.",
        "mode": "review",
        "expected_output": "A cited approve/reject decision with material risks",
        "work_item_id": child["id"],
    }

    passive = await env.client.post(
        f"/api/projects/{project['id']}/a2a",
        json={**request, "mode": "notify"},
    )
    assert passive.status_code == 422

    blocked = await env.client.post(f"/api/projects/{project['id']}/a2a", json=request)
    assert blocked.status_code == 409
    assert "该任务的前置任务尚未完成" in blocked.text
    assert "Produce reviewed input evidence" in blocked.text

    completed = await env.client.patch(
        f"/api/projects/{project['id']}/work-items/{parent['id']}",
        json={"status": "done"},
    )
    assert completed.status_code == 200, completed.text
    queued = await env.client.post(f"/api/projects/{project['id']}/a2a", json=request)
    assert queued.status_code == 202, queued.text

    run = await env.db.get(ProjectRun, uuid.UUID(queued.json()["run_id"]))
    assert run is not None
    assert run.work_item_id == uuid.UUID(child["id"])
    assert run.input["title"] == request["title"]
    assert run.input["expected_output"] == request["expected_output"]
    assert run.input["message"].endswith(f"Expected output: {request['expected_output']}")
    event = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.run_id == run.id,
                ProjectEvent.event_type == "a2a.queued",
            )
        )
    ).scalar_one()
    assert event.summary == "Queued project collaboration action: Make release decision"
    assert event.event_metadata["expected_output"] == request["expected_output"]


async def test_member_departure_is_audited_revocation_and_restore_starts_a_fresh_child(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.api.websocket import WebSocketChatHandler
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectEvent, ProjectRun
    from app.models.subagent_run import SubagentRun
    from app.services.project_runtime_tools import load_project_runtime_scope
    from app.services.project_service import project_session_access_mode
    from app.services.conversation_turn_lifecycle import (
        get_conversation_turn_snapshot,
    )

    env = project_api
    project = await _create_project(env, name="Member lifecycle")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    await env.db.commit()
    blocked_item = await _create_work_item(
        env,
        project_id,
        title="Blocked member action",
        assignee_agent_id=env.worker_id,
    )

    members = (await env.client.get(f"/api/projects/{project_id}/members")).json()
    leader = next(row for row in members if row["agent_id"] == str(env.leader_id))
    worker = next(row for row in members if row["agent_id"] == str(env.worker_id))

    leader_removal = await env.client.post(
        f"/api/projects/{project_id}/members/{leader['id']}/remove",
        json={"reason": "cannot remove current leader"},
    )
    assert leader_removal.status_code == 422

    first_run_response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(env.worker_id), "trigger_type": "manual", "input": {"objective": "First"}},
    )
    assert first_run_response.status_code == 201, first_run_response.text
    first_run_payload = first_run_response.json()
    first_child_id = uuid.UUID(first_run_payload["output"]["subagent_session_id"])
    first_project_run_id = uuid.UUID(first_run_payload["id"])
    first_child_session = await env.db.get(ChatSession, first_child_id)
    assert first_child_session is not None
    owner_user = await env.db.get(User, env.owner_id)
    assert owner_user is not None
    assert await project_session_access_mode(env.db, owner_user, first_child_session) == "edit"
    live_handler = WebSocketChatHandler.__new__(WebSocketChatHandler)
    live_handler.project_session_access = "edit"
    live_handler.conv_id = str(first_child_id)
    live_handler.user_id = env.owner_id
    live_handler.read_only = False
    assert await live_handler._project_session_still_writable() is True

    # Sending from the exact project child drawer is a durable inbox append,
    # not a second generic Web LLM turn.  Client retry is idempotent, keeps the
    # active ProjectRun association explicit and returns the persisted receipt.
    live_handler.agent_id = env.worker_id
    live_handler.source_channel = "subagent"
    receipts: list[dict] = []

    async def _capture_ws(_agent_id: str, conversation_id: str, payload: dict):
        if conversation_id == str(first_child_id):
            receipts.append(payload)

    monkeypatch.setattr("app.api.websocket.manager.send_to_session", _capture_ws)
    for _retry in range(2):
        assert (
            await live_handler._enqueue_project_subagent_message(
                content="Continue from the project drawer",
                display_content="Continue from the project drawer",
                file_name="evidence.txt",
                client_message_id="drawer-client-1",
                attachments=[{"type": "file", "name": "evidence.txt", "url": "/evidence.txt"}],
            )
            is True
        )

    env.db.expire_all()
    drawer_inputs = (
        (
            await env.db.execute(
                select(ChatMessage).where(
                    ChatMessage.conversation_id == str(first_child_id),
                    ChatMessage.content == "Continue from the project drawer",
                )
            )
        )
        .scalars()
        .all()
    )
    durable_child = await env.db.get(SubagentRun, first_child_id)
    canonical_turn_anchor_id = await env.db.scalar(
        select(ChatMessage.id)
        .where(
            ChatMessage.conversation_id == str(first_child_id),
            ChatMessage.message_meta["kind"].as_string() == "subagent_input",
            ChatMessage.message_meta["subagent_input_state"]
            .as_string()
            .in_(["pending", "processing"]),
        )
        .order_by(ChatMessage.created_at, ChatMessage.id)
        .limit(1)
    )
    assert len(drawer_inputs) == 1
    assert drawer_inputs[0].message_meta["kind"] == "subagent_input"
    assert drawer_inputs[0].message_meta["subagent_input_state"] == "pending"
    assert drawer_inputs[0].message_meta["project_run_id"] == str(first_project_run_id)
    assert drawer_inputs[0].message_meta["attachments"][0]["name"] == "evidence.txt"
    assert durable_child is not None and durable_child.status == "queued"
    assert durable_child.lease_owner is None
    assert durable_child.lease_expires_at is None
    committed = [row for row in receipts if row.get("type") == "user_message_committed"]
    assert len(committed) == 2
    assert {row["message_id"] for row in committed} == {str(drawer_inputs[0].id)}
    assert all(row["turn"]["phase"] == "active" for row in committed)
    # The receipt identifies the newly committed row separately, while its
    # lifecycle projection truthfully keeps the oldest queued input as the one
    # session owner. A later drawer append cannot jump the queue.
    assert {row["turn"]["turn_anchor_id"] for row in committed} == {
        str(canonical_turn_anchor_id)
    }

    removed = await env.client.post(
        f"/api/projects/{project_id}/members/{worker['id']}/remove",
        json={"reason": "staffing change"},
    )
    assert removed.status_code == 200, removed.text
    removed_payload = removed.json()
    assert removed_payload["id"] == worker["id"]
    assert removed_payload["is_enabled"] is False
    assert removed_payload["config_snapshot"]["membership"]["state"] == "departed"

    env.db.expire_all()
    first_child = await env.db.get(SubagentRun, first_child_id)
    first_project_run = await env.db.get(ProjectRun, first_project_run_id)
    first_child_session = await env.db.get(ChatSession, first_child_id)
    current_turn = await get_conversation_turn_snapshot(
        env.db,
        agent_id=env.worker_id,
        conversation_id=str(first_child_id),
    )
    assert first_child is not None and first_child.status == "cancelled"
    assert first_project_run is not None and first_project_run.status == "cancelled"
    assert first_child_session is not None
    assert first_child_session.im_config["membership_revoked"] is True
    assert current_turn.anchor_id == canonical_turn_anchor_id
    assert current_turn.status == "cancelled"
    assert current_turn.revision >= 2
    owner_user = await env.db.get(User, env.owner_id)
    assert owner_user is not None
    assert await project_session_access_mode(env.db, owner_user, first_child_session) == "read"
    assert await live_handler._project_session_still_writable() is False
    with pytest.raises(ValueError, match="active runtime member|permanently read-only"):
        await load_project_runtime_scope(
            env.db,
            session_id=first_child_id,
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
        )

    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    mention = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "@Worker continue", "mentions": [str(env.worker_id)]},
    )
    assert mention.status_code == 422
    assigned_run = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(env.worker_id), "trigger_type": "manual", "input": {"objective": "Blocked"}},
    )
    assert assigned_run.status_code == 422
    a2a = await env.client.post(
        f"/api/projects/{project_id}/a2a",
        json={
            "from_agent_id": str(env.leader_id),
            "to_agent_id": str(env.worker_id),
            "title": "Complete blocked member action",
            "message": "Blocked",
            "mode": "delegate",
            "expected_output": "A completed deliverable with evidence",
            "work_item_id": blocked_item["id"],
        },
    )
    assert a2a.status_code == 422

    duplicate = await env.client.post(
        f"/api/projects/{project_id}/members",
        json={"agent_id": str(env.worker_id)},
    )
    assert duplicate.status_code == 409
    assert "restore" in duplicate.json()["detail"].lower()

    restored = await env.client.post(
        f"/api/projects/{project_id}/members/{worker['id']}/restore",
        json={"reason": "return to project"},
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["id"] == worker["id"]
    assert restored.json()["is_enabled"] is True
    assert restored.json()["config_snapshot"]["membership"]["generation"] == 2

    env.db.expire_all()
    first_child_session = await env.db.get(ChatSession, first_child_id)
    assert first_child_session is not None
    owner_user = await env.db.get(User, env.owner_id)
    assert owner_user is not None
    assert await project_session_access_mode(env.db, owner_user, first_child_session) == "read"
    assert await live_handler._project_session_still_writable() is False

    second_run_response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(env.worker_id), "trigger_type": "manual", "input": {"objective": "Fresh"}},
    )
    assert second_run_response.status_code == 201, second_run_response.text
    second_child_id = uuid.UUID(second_run_response.json()["output"]["subagent_session_id"])
    assert second_child_id != first_child_id
    second_child_session = await env.db.get(ChatSession, second_child_id)
    assert second_child_session is not None
    assert second_child_session.im_config["project_membership_generation"] == 2
    owner_user = await env.db.get(User, env.owner_id)
    assert owner_user is not None
    assert await project_session_access_mode(env.db, owner_user, second_child_session) == "edit"

    events = (await env.db.execute(select(ProjectEvent).where(ProjectEvent.project_id == project_id))).scalars().all()
    lifecycle_events = [row for row in events if row.event_type in {"member.departed", "member.restored"}]
    assert [row.event_type for row in lifecycle_events] == ["member.departed", "member.restored"]
    assert lifecycle_events[0].event_metadata["snapshot_retained"] is True
    assert lifecycle_events[1].event_metadata["old_sessions_remain_read_only"] is True


async def test_only_project_owner_can_remove_and_restore_project_agent_members(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.services.project_service import project_session_access_mode

    env = project_api
    project = await _create_project(env, name="Editor membership controls")
    project_id = project["id"]
    stored_project = await env.db.get(Project, uuid.UUID(project_id))
    assert stored_project is not None
    stored_project.status = "running"
    await env.db.commit()
    members = (await env.client.get(f"/api/projects/{project_id}/members")).json()
    reviewer = next(row for row in members if row["agent_id"] == str(env.reviewer_id))
    reviewer_run = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"agent_id": str(env.reviewer_id), "trigger_type": "manual", "input": {"objective": "Review"}},
    )
    assert reviewer_run.status_code == 201, reviewer_run.text
    reviewer_child_id = uuid.UUID(reviewer_run.json()["output"]["subagent_session_id"])
    reviewer_session = await env.db.get(ChatSession, reviewer_child_id)
    assert reviewer_session is not None

    grant = await env.client.post(
        f"/api/projects/{project_id}/access-grants",
        json={"user_id": str(env.viewer_id), "role": "view"},
    )
    assert grant.status_code == 201, grant.text
    viewer_user = await env.db.get(User, env.viewer_id)
    assert viewer_user is not None
    assert await project_session_access_mode(env.db, viewer_user, reviewer_session) == "read"
    env.authenticate_as(env.viewer_id)
    forbidden = await env.client.post(
        f"/api/projects/{project_id}/members/{reviewer['id']}/remove",
        json={"reason": "viewer cannot"},
    )
    assert forbidden.status_code == 404

    env.authenticate_as(env.owner_id)
    grant_id = grant.json()["id"]
    grant_row = await env.db.get(ProjectAccessGrant, uuid.UUID(grant_id))
    assert grant_row is not None
    grant_row.role = "edit"
    await env.db.commit()
    viewer_user = await env.db.get(User, env.viewer_id)
    reviewer_session = await env.db.get(ChatSession, reviewer_child_id)
    assert viewer_user is not None and reviewer_session is not None
    assert await project_session_access_mode(env.db, viewer_user, reviewer_session) == "edit"
    env.authenticate_as(env.viewer_id)
    removed = await env.client.post(
        f"/api/projects/{project_id}/members/{reviewer['id']}/remove",
        json={"reason": "editor staffing"},
    )
    assert removed.status_code == 404

    env.authenticate_as(env.owner_id)
    removed = await env.client.post(
        f"/api/projects/{project_id}/members/{reviewer['id']}/remove",
        json={"reason": "owner staffing"},
    )
    assert removed.status_code == 200, removed.text
    viewer_user = await env.db.get(User, env.viewer_id)
    reviewer_session = await env.db.get(ChatSession, reviewer_child_id)
    assert viewer_user is not None and reviewer_session is not None
    assert await project_session_access_mode(env.db, viewer_user, reviewer_session) == "read"
    restored = await env.client.post(
        f"/api/projects/{project_id}/members/{reviewer['id']}/restore",
        json={"reason": "editor restore"},
    )
    assert restored.status_code == 200, restored.text
    viewer_user = await env.db.get(User, env.viewer_id)
    reviewer_session = await env.db.get(ChatSession, reviewer_child_id)
    assert viewer_user is not None and reviewer_session is not None
    assert await project_session_access_mode(env.db, viewer_user, reviewer_session) == "read"


async def test_manual_run_requires_kickoff_and_dispatches_default_leader(project_api: ProjectApiEnv):
    from app.models.chat_session import ChatSession
    from app.models.project import Project, ProjectRun
    from app.models.subagent_run import SubagentRun

    env = project_api
    project = await _create_project(env, name="Manual execution contract")
    project_id = uuid.UUID(project["id"])

    before_kickoff = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"trigger_type": "manual", "input": {"objective": "Continue delivery"}},
    )
    assert before_kickoff.status_code == 409

    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    await env.db.commit()

    empty = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"trigger_type": "manual", "input": {}},
    )
    assert empty.status_code == 422
    assert "explicit task/objective/message" in empty.text

    response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"trigger_type": "manual", "input": {"objective": "Continue delivery"}},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["agent_id"] == str(env.leader_id)
    assert payload["status"] in {"queued", "running"}
    assert payload["input"]["dispatch"]["task"] == "Continue delivery"
    assert payload["output"]["subagent_session_id"]

    run = await env.db.get(ProjectRun, uuid.UUID(payload["id"]))
    child = await env.db.get(SubagentRun, uuid.UUID(payload["output"]["subagent_session_id"]))
    assert child is not None
    child_session = await env.db.get(ChatSession, child.id)
    assert child_session is not None
    assert child_session.agent_id == env.leader_id
    assert run is not None and run.output["subagent_run_id"] == str(child.id)
    assert child.project_id == project_id
    tenant = await env.db.get(Tenant, env.tenant_id)
    anchor = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child.id),
                ChatMessage.message_meta["project_run_id"].as_string() == str(run.id),
            )
        )
    ).scalar_one()
    assert tenant is not None and tenant.default_model_id is not None
    tenant_model = await env.db.get(LLMModel, tenant.default_model_id)
    assert tenant_model is not None
    assert child.model is None
    from app.services.chat_model_selection import resolve_project_member_runtime_models

    project_agent = await env.db.get(Agent, env.leader_id)
    assert project_agent is not None
    resolved_models = await resolve_project_member_runtime_models(
        env.db,
        agent=project_agent,
        member_config=child_session.im_config["member_config_snapshot"],
        project_settings=stored_project.settings,
    )
    assert resolved_models.primary_model is not None
    assert resolved_models.primary_model.id == tenant_model.id
    assert "model_id" not in anchor.message_meta


async def test_project_run_and_leader_batch_use_frozen_shared_execution_user(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import Project, ProjectRun
    from app.models.subagent_run import SubagentRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Frozen shared execution")
    project_id = uuid.UUID(project["id"])
    configured = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert configured.status_code == 200, configured.text
    await _mark_project_running(env, project_id)

    real_created = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"trigger_type": "manual", "input": {"objective": "Use the selected shared identity"}},
    )
    assert real_created.status_code == 201, real_created.text
    real_child = await env.db.get(
        SubagentRun,
        uuid.UUID(real_created.json()["output"]["subagent_session_id"]),
    )
    assert real_child is not None and real_child.execution_user_id == env.viewer_id

    real_dispatch = subagent_runtime.dispatch_project_run

    async def defer_dispatch(_run_id: uuid.UUID) -> dict[str, str]:
        return {"status": "queued"}

    monkeypatch.setattr(subagent_runtime, "dispatch_project_run", defer_dispatch)
    created = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={
            "agent_id": str(env.worker_id),
            "trigger_type": "manual",
            "input": {"objective": "Freeze this execution identity"},
        },
    )
    assert created.status_code == 201, created.text
    run_id = uuid.UUID(created.json()["id"])
    frozen_run = await env.db.get(ProjectRun, run_id)
    assert frozen_run is not None
    assert frozen_run.initiated_by_user_id == env.owner_id
    assert frozen_run.execution_user_id == env.viewer_id

    tenant = await env.db.get(Tenant, env.tenant_id)
    assert tenant is not None
    alternate = await _user(env.db, tenant, "Dispatch Alternate")
    await env.db.commit()
    alternate_id = alternate.id
    changed = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "shared_with_user_ids": [str(env.viewer_id), str(alternate_id)],
            "execution_user_id": str(alternate_id),
        },
    )
    assert changed.status_code == 200, changed.text

    dispatched_users: list[uuid.UUID] = []

    async def capture_subagent(**kwargs):
        dispatched_users.append(kwargs["execution_user_id"])
        return SimpleNamespace(id=uuid.uuid4(), status="queued"), True

    monkeypatch.setattr(subagent_runtime, "create_subagent", capture_subagent)
    dispatched = await real_dispatch(run_id)
    assert dispatched["status"] == "queued"
    assert dispatched_users == [env.viewer_id]

    env.db.expire_all()
    frozen_run = await env.db.get(ProjectRun, run_id)
    assert frozen_run is not None
    group_id = uuid.UUID(frozen_run.input["dispatch"]["group_session_id"])
    group = await env.db.get(ChatSession, group_id)
    assert group is not None
    env.db.add(
        ChatMessage(
            agent_id=group.agent_id,
            sender_agent_id=env.worker_id,
            role="assistant",
            content="Participant evidence is ready",
            conversation_id=str(group.id),
            message_meta={
                "kind": "project_subagent_reply",
                "leader_batch_state": "pending",
                "source_project_run_ids": [str(run_id)],
            },
            created_at=datetime.now(UTC) - timedelta(seconds=1),
        )
    )
    await env.db.commit()

    assert await subagent_runtime._dispatch_project_leader_batch(group.id, debounce_seconds=0) is True
    assert dispatched_users == [env.viewer_id, alternate_id]
    batch_run = (
        await env.db.execute(
            select(ProjectRun).where(
                ProjectRun.project_id == project_id,
                ProjectRun.trigger_type == "leader_reply_batch",
            )
        )
    ).scalar_one()
    assert batch_run.initiated_by_user_id == env.owner_id
    assert batch_run.execution_user_id == alternate_id
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None and stored_project.execution_user_id == alternate_id


async def test_project_run_without_agent_model_uses_exact_project_model(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import ProjectRun
    from app.models.subagent_run import SubagentRun

    env = project_api
    selected = LLMModel(
        tenant_id=env.tenant_id,
        provider="openai",
        model="project-explicit-model",
        api_key_encrypted="project-test-only",
        label="Project explicit model",
        enabled=True,
        context_window=64000,
    )
    env.db.add(selected)
    await env.db.flush()
    project = await _create_project(env, name="Project model fallback")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "runtime": {"model": str(selected.id)},
    }
    await env.db.commit()

    response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={
            "agent_id": str(env.worker_id),
            "trigger_type": "manual",
            "input": {"objective": "Use the project model"},
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    child_id = uuid.UUID(payload["output"]["subagent_session_id"])
    child = await env.db.get(SubagentRun, child_id)
    project_run = await env.db.get(ProjectRun, uuid.UUID(payload["id"]))
    anchor = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["project_run_id"].as_string() == payload["id"],
            )
        )
    ).scalar_one()
    assert child is not None and child.project_id == project_id
    assert project_run is not None and project_run.status in {"queued", "running"}
    assert child.model is None
    assert "model_id" not in anchor.message_meta

    # Historical project child inputs did not carry a per-turn model snapshot.
    # The unified channel path must still resolve the same project model rather
    # than falling back to an absent Agent-level configuration.
    anchor.message_meta = {key: value for key, value in dict(anchor.message_meta or {}).items() if key != "model_id"}
    await env.db.commit()
    captured: dict = {}

    async def _fake_scene(*_args, **_kwargs):
        return {}

    async def _fake_llm(**kwargs):
        captured.update(kwargs)
        return "project model works"

    monkeypatch.setattr("app.services.scene_service.load_turn_scene_context", _fake_scene)
    monkeypatch.setattr("app.services.llm.call_llm_with_failover", _fake_llm)
    monkeypatch.setattr("app.services.channel_llm.is_agent_expired", lambda _agent: False)
    from app.services.channel_llm import _call_agent_llm
    from app.services.agent_runtime_workspace import resolve_agent_runtime_workspace

    async with env.session_factory() as runtime_db:
        runtime_agent = await runtime_db.get(Agent, env.worker_id)
        runtime_session = await runtime_db.get(ChatSession, child_id)
        assert runtime_agent is not None and runtime_session is not None
        runtime_workspace = resolve_agent_runtime_workspace(
            agent_id=runtime_agent.id,
            agent_scope=runtime_agent.scope,
            agent_project_id=runtime_agent.project_id,
            tenant_id=runtime_agent.tenant_id,
            session_project_id=runtime_session.project_id,
            session_config=runtime_session.im_config,
        )
        reply = await _call_agent_llm(
            runtime_db,
            env.worker_id,
            anchor.content,
            session_id=str(child_id),
            user_id=env.owner_id,
            turn_anchor_id=anchor.id,
            prepared_tools=[],
            broadcast_web=False,
            runtime_session=runtime_session,
            runtime_workspace=runtime_workspace,
        )
    assert reply == "project model works"
    assert captured["primary_model"].id == selected.id


async def test_project_run_without_any_tenant_model_keeps_a_durable_unresolved_child(project_api: ProjectApiEnv):
    from app.models.project import ProjectRun
    from app.models.subagent_run import SubagentRun

    env = project_api
    tenant_models = (await env.db.execute(select(LLMModel).where(LLMModel.tenant_id == env.tenant_id))).scalars().all()
    for model in tenant_models:
        model.enabled = False

    foreign_tenant = Tenant(name="Foreign", slug=f"foreign-{uuid.uuid4().hex[:8]}")
    env.db.add(foreign_tenant)
    await env.db.flush()
    foreign_model = LLMModel(
        tenant_id=foreign_tenant.id,
        provider="openai",
        model="foreign-model",
        api_key_encrypted="must-not-cross-tenant",
        label="Foreign model",
        enabled=True,
        context_window=64000,
    )
    env.db.add(foreign_model)
    await env.db.flush()

    project = await _create_project(env, name="No model project")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "runtime": {"model": str(foreign_model.id)},
    }
    await env.db.commit()

    response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={
            "agent_id": str(env.worker_id),
            "trigger_type": "manual",
            "input": {"objective": "Must fail explicitly"},
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    project_run = await env.db.get(ProjectRun, uuid.UUID(payload["id"]))
    children = (await env.db.execute(select(SubagentRun).where(SubagentRun.project_id == project_id))).scalars().all()
    assert project_run is not None and project_run.status in {"queued", "running"}
    assert project_run.error is None
    assert payload["output"].get("subagent_session_id")
    assert len(children) == 1 and children[0].model is None


async def test_project_run_reconcile_persists_finished_terminal_state(project_api: ProjectApiEnv):
    from app.models.project import ProjectRun

    env = project_api
    project = await _create_project(env, name="Repair stale run")
    run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=uuid.UUID(project["id"]),
        initiated_by_user_id=env.owner_id,
        status="running",
        trigger_type="manual",
        finished_at=datetime.now(UTC),
        output={"result": "Already finished"},
    )
    env.db.add(run)
    await env.db.commit()

    response = await env.client.get(f"/api/projects/{project['id']}/runs")
    assert response.status_code == 200, response.text
    repaired_payload = next(item for item in response.json() if item["id"] == str(run.id))
    assert repaired_payload["status"] == "succeeded"

    # Verify the GET-owned reconciliation committed, rather than merely
    # changing the request session identity map.
    async with env.session_factory() as independent_db:
        persisted = await independent_db.get(ProjectRun, run.id)
        assert persisted is not None and persisted.status == "succeeded"


async def test_pending_project_dispatch_scan_cannot_starve_after_fifty_active_runs(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import ProjectRun
    from app.services import project_service, subagent_runtime

    env = project_api
    project = await _create_project(env, name="Fair durable dispatch scan")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    old_created_at = datetime(2026, 1, 1, tzinfo=UTC)
    new_created_at = datetime(2026, 1, 2, tzinfo=UTC)

    def dispatch_payload(task: str) -> dict[str, dict[str, str]]:
        return {
            "dispatch": {
                "group_session_id": str(uuid.uuid4()),
                "project_member_id": str(uuid.uuid4()),
                "turn_anchor_id": str(uuid.uuid4()),
                "task": task,
            }
        }

    old_active_runs = [
        ProjectRun(
            tenant_id=env.tenant_id,
            project_id=project_id,
            agent_id=env.leader_id,
            initiated_by_user_id=env.owner_id,
            status="running",
            trigger_type="leader_kickoff",
            input=dispatch_payload(f"Already dispatched {index}"),
            output={"subagent_run_id": str(uuid.uuid4())},
            created_at=old_created_at,
        )
        for index in range(55)
    ]
    pending_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="leader_kickoff",
        input=dispatch_payload("Must not be starved"),
        output={},
        created_at=new_created_at,
    )
    env.db.add_all([*old_active_runs, pending_run])
    await env.db.commit()

    # Keep this regression focused on candidate scanning. Terminal repair is
    # covered independently and must not affect active dispatched rows here.
    monkeypatch.setattr(project_service, "reconcile_project_run_terminal_state", lambda _run: False)

    pending_ids = await subagent_runtime._pending_project_dispatch_runs()

    assert pending_ids == [pending_run.id]
    assert await subagent_runtime._pending_project_dispatch_runs() == [pending_run.id]

    pending_run.output = {"subagent_run_id": str(uuid.uuid4())}
    await env.db.commit()
    assert await subagent_runtime._pending_project_dispatch_runs() == []


async def test_dispatch_does_not_regress_child_completed_project_run(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import Project, ProjectMemberSnapshot, ProjectRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Fast child completion")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    leader_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == project_id,
                ProjectMemberSnapshot.agent_id == env.leader_id,
            )
        )
    ).scalar_one()
    anchor = ChatMessage(
        id=uuid.uuid4(),
        agent_id=uuid.UUID(group["access_agent_id"]),
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Complete immediately",
        conversation_id=group["id"],
        message_meta={"kind": "project_run_request", "mentions": [str(env.leader_id)]},
    )
    run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="manual",
        input={
            "dispatch": {
                "group_session_id": group["id"],
                "project_member_id": str(leader_member.id),
                "turn_anchor_id": str(anchor.id),
                "task": "Complete immediately",
            }
        },
    )
    env.db.add_all([anchor, run])
    await env.db.commit()
    project_run_id = run.id
    child_id = uuid.uuid4()

    async def complete_before_dispatch_commit(**_kwargs):
        async with env.session_factory() as race_db:
            raced = await race_db.get(ProjectRun, project_run_id)
            assert raced is not None
            raced.status = "succeeded"
            raced.finished_at = datetime.now(UTC)
            raced.output = {"result": "Fast result"}
            await race_db.commit()
        return SimpleNamespace(id=child_id, status="completed"), True

    monkeypatch.setattr(subagent_runtime, "create_subagent", complete_before_dispatch_commit)
    result = await subagent_runtime.dispatch_project_run(project_run_id)
    assert result["status"] == "succeeded"
    env.db.expire_all()
    persisted = await env.db.get(ProjectRun, project_run_id)
    assert persisted is not None and persisted.status == "succeeded"
    assert persisted.finished_at is not None
    assert persisted.output["result"] == "Fast result"
    assert persisted.output["subagent_session_id"] == str(child_id)


async def test_dispatch_does_not_regress_fast_waiting_project_run(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import Project, ProjectMemberSnapshot, ProjectRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Fast child waiting")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    leader_member = await env.db.scalar(
        select(ProjectMemberSnapshot).where(
            ProjectMemberSnapshot.project_id == project_id,
            ProjectMemberSnapshot.agent_id == env.leader_id,
        )
    )
    assert leader_member is not None
    anchor = ChatMessage(
        id=uuid.uuid4(),
        agent_id=uuid.UUID(group["access_agent_id"]),
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Wait immediately",
        conversation_id=group["id"],
        message_meta={"kind": "project_run_request"},
    )
    run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="manual",
        input={
            "dispatch": {
                "group_session_id": group["id"],
                "project_member_id": str(leader_member.id),
                "turn_anchor_id": str(anchor.id),
                "task": "Wait immediately",
            }
        },
    )
    env.db.add_all([anchor, run])
    await env.db.commit()
    project_run_id = run.id
    child_id = uuid.uuid4()

    async def wait_before_dispatch_commit(**_kwargs):
        async with env.session_factory() as race_db:
            raced = await race_db.get(ProjectRun, project_run_id)
            assert raced is not None
            raced.status = "waiting"
            raced.started_at = datetime.now(UTC)
            race_db.add(
                ChatMessage(
                    agent_id=env.leader_id,
                    user_id=env.owner_id,
                    sender_user_id=env.owner_id,
                    role="user",
                    content="Wait immediately",
                    conversation_id=str(child_id),
                    message_meta={
                        "kind": subagent_runtime.SUBAGENT_INPUT,
                        "subagent_input_state": subagent_runtime.INPUT_PROCESSING,
                        "project_run_id": str(project_run_id),
                    },
                )
            )
            await race_db.commit()
        return SimpleNamespace(id=child_id, status=subagent_runtime.RUN_WAITING), True

    monkeypatch.setattr(subagent_runtime, "create_subagent", wait_before_dispatch_commit)
    result = await subagent_runtime.dispatch_project_run(project_run_id)

    assert result["status"] == "waiting"
    env.db.expire_all()
    persisted = await env.db.get(ProjectRun, project_run_id)
    assert persisted is not None and persisted.status == "waiting"
    assert persisted.started_at is not None


async def test_project_dispatch_respects_project_parallel_task_limit(
    project_api: ProjectApiEnv,
):
    from app.models.project import Project, ProjectMemberSnapshot, ProjectRun
    from app.services import subagent_runtime
    from app.services.project_service import freeze_run_members

    env = project_api
    project = await _create_project(env, name="Bounded project concurrency")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "running"
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "runtime": {"max_parallel_runs": 1},
    }
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    leader_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == project_id,
                ProjectMemberSnapshot.agent_id == env.leader_id,
            )
        )
    ).scalar_one()
    anchor = ChatMessage(
        id=uuid.uuid4(),
        agent_id=uuid.UUID(group["access_agent_id"]),
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Run after capacity is available",
        conversation_id=group["id"],
        message_meta={"kind": "project_run_request"},
    )
    active = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.worker_id,
        initiated_by_user_id=env.owner_id,
        status="running",
        trigger_type="leader_kickoff",
        started_at=datetime.now(UTC),
    )
    pending = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="manual",
        input={
            "dispatch": {
                "group_session_id": group["id"],
                "project_member_id": str(leader_member.id),
                "turn_anchor_id": str(anchor.id),
                "task": "Run after capacity is available",
            }
        },
    )
    env.db.add_all([anchor, active, pending])
    await env.db.flush()
    await freeze_run_members(env.db, stored_project, pending)
    await env.db.commit()
    active_id = active.id
    pending_id = pending.id

    dispatched = await subagent_runtime.dispatch_project_run(pending_id)
    assert dispatched["status"] == "queued"
    child_id = uuid.UUID(dispatched["subagent_run_id"])
    env.db.expire_all()
    still_pending = await env.db.get(ProjectRun, pending_id)
    assert still_pending is not None and still_pending.status == "queued"
    assert still_pending.started_at is None
    assert still_pending.output["subagent_run_id"] == str(child_id)
    assert await subagent_runtime._claim_subagent(child_id) is None

    stored_active = await env.db.get(ProjectRun, active_id)
    assert stored_active is not None
    stored_active.status = "succeeded"
    stored_active.finished_at = datetime.now(UTC)
    await env.db.commit()
    assert await subagent_runtime._claim_subagent(child_id) == child_id
    claimed_input = await subagent_runtime._load_or_start_input(child_id)
    assert claimed_input is not None
    claimed_anchor, _recovering = claimed_input
    env.db.expire_all()
    started = await env.db.get(ProjectRun, pending_id)
    assert started is not None and started.status == "running"
    assert started.started_at is not None

    await subagent_runtime._requeue_capacity_blocked_subagent(child_id, claimed_anchor.id)
    env.db.expire_all()
    requeued = await env.db.get(ProjectRun, pending_id)
    assert requeued is not None and requeued.status == "queued"
    assert requeued.started_at is None
    input_row = await env.db.get(ChatMessage, claimed_anchor.id)
    assert input_row is not None
    assert input_row.message_meta["subagent_input_state"] == subagent_runtime.INPUT_PENDING


async def test_active_child_inputs_durably_advance_only_their_exact_project_runs(
    project_api: ProjectApiEnv,
):
    """Worker recovery must not leave the Runs UI stuck at queued."""
    from app.models.project import ProjectMemberSnapshot, ProjectRun
    from app.models.subagent_run import SubagentRun
    from app.services import subagent_runtime
    from app.services.project_service import freeze_run_members

    env = project_api
    project = await _create_project(env, name="Project run worker state")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    leader_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == project_id,
                ProjectMemberSnapshot.agent_id == env.leader_id,
            )
        )
    ).scalar_one()
    anchor = ChatMessage(
        id=uuid.uuid4(),
        agent_id=uuid.UUID(group["access_agent_id"]),
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Run the exact durable input",
        conversation_id=group["id"],
        message_meta={"kind": "project_run_request"},
    )
    project_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="manual",
        input={
            "dispatch": {
                "group_session_id": group["id"],
                "project_member_id": str(leader_member.id),
                "turn_anchor_id": str(anchor.id),
                "task": "Run the exact durable input",
            }
        },
    )
    unrelated_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="manual",
        input={"objective": "Must remain queued until its own input runs"},
    )
    terminal_finished_at = datetime.now(UTC)
    terminal_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="succeeded",
        trigger_type="manual",
        input={"objective": "Already finished"},
        finished_at=terminal_finished_at,
    )
    env.db.add_all([anchor, project_run, unrelated_run, terminal_run])
    await env.db.flush()
    project_row = await env.db.get(Project, project_id)
    assert project_row is not None
    await freeze_run_members(env.db, project_row, project_run)
    await env.db.commit()
    project_run_id = project_run.id
    unrelated_run_id = unrelated_run.id
    terminal_run_id = terminal_run.id

    dispatched = await subagent_runtime.dispatch_project_run(project_run_id)
    child_id = uuid.UUID(dispatched["subagent_session_id"])
    assert dispatched["status"] == "queued"

    # Reproduce a process exit after the exact input became processing but
    # before an older worker updated ProjectRun.
    assert await subagent_runtime._claim_subagent(child_id) == child_id
    async with env.session_factory() as crash_db:
        child = await crash_db.get(SubagentRun, child_id, with_for_update=True)
        child_input = (
            await crash_db.execute(
                select(ChatMessage).where(
                    ChatMessage.conversation_id == str(child_id),
                    ChatMessage.message_meta["project_run_id"].as_string() == str(project_run_id),
                )
            )
        ).scalar_one()
        child_input.message_meta = {
            **dict(child_input.message_meta or {}),
            "subagent_input_state": "processing",
            "subagent_turn_anchor_id": str(child_input.id),
            "turn_status": "running",
        }
        child.lease_expires_at = datetime(2000, 1, 1, tzinfo=UTC)
        await crash_db.commit()

    assert await subagent_runtime._claim_subagent(child_id) == child_id
    recovered = await subagent_runtime._load_or_start_input(child_id)
    assert recovered is not None and recovered[1] is True

    env.db.expire_all()
    running = await env.db.get(ProjectRun, project_run_id)
    untouched = await env.db.get(ProjectRun, unrelated_run_id)
    terminal = await env.db.get(ProjectRun, terminal_run_id)
    assert running is not None and running.status == "running"
    assert running.started_at is not None
    first_started_at = running.started_at
    assert untouched is not None and untouched.status == "queued"
    assert untouched.started_at is None
    assert terminal is not None and terminal.status == "succeeded"
    assert terminal.finished_at is not None
    terminal_persisted_finished_at = terminal.finished_at

    # Inputs appended while the reusable child is already executing are
    # consumed at a round boundary. They use the same exact-id transition;
    # a referenced terminal Run must still remain terminal.
    env.db.add_all(
        [
            ChatMessage(
                agent_id=env.leader_id,
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Start the second exact run",
                conversation_id=str(child_id),
                message_meta={
                    "kind": "subagent_input",
                    "subagent_input_state": "pending",
                    "project_run_id": str(unrelated_run_id),
                },
            ),
            ChatMessage(
                agent_id=env.leader_id,
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Do not regress the terminal run",
                conversation_id=str(child_id),
                message_meta={
                    "kind": "subagent_input",
                    "subagent_input_state": "pending",
                    "project_run_id": str(terminal_run_id),
                },
            ),
        ]
    )
    await env.db.commit()
    injected = await subagent_runtime._drain_subagent_inbox(
        child_id,
        recovered[0].id,
    )
    assert {row["content"] for row in injected} == {
        "Start the second exact run",
        "Do not regress the terminal run",
    }
    env.db.expire_all()
    now_running = await env.db.get(ProjectRun, unrelated_run_id)
    terminal_after_drain = await env.db.get(ProjectRun, terminal_run_id)
    assert now_running is not None and now_running.status == "running"
    assert now_running.started_at is not None
    assert terminal_after_drain is not None
    assert terminal_after_drain.status == "succeeded"
    assert terminal_after_drain.finished_at == terminal_persisted_finished_at

    # Idempotent recovery keeps the original start timestamp and terminal fact.
    recovered_again = await subagent_runtime._load_or_start_input(child_id)
    assert recovered_again is not None and recovered_again[1] is True
    env.db.expire_all()
    running_again = await env.db.get(ProjectRun, project_run_id)
    terminal_again = await env.db.get(ProjectRun, terminal_run_id)
    assert running_again is not None and running_again.started_at == first_started_at
    assert terminal_again is not None and terminal_again.status == "succeeded"
    assert terminal_again.finished_at == terminal_persisted_finished_at


async def test_project_a2a_delivery_returns_scoped_session_identifiers(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectRun
    from app.services import agent_tools, project_service

    env = project_api
    project = await _create_project(env, name="A2A session identity")
    await _mark_project_running(env, project["id"])
    work_item = await _create_work_item(
        env,
        project["id"],
        title="Resolve exact project thread",
        assignee_agent_id=env.reviewer_id,
    )
    queued = await env.client.post(
        f"/api/projects/{project['id']}/a2a",
        json={
            "from_agent_id": str(env.worker_id),
            "to_agent_id": str(env.reviewer_id),
            "title": "Resolve exact project thread",
            "message": "Return the exact project thread",
            "mode": "consult",
            "expected_output": "The exact scoped session identifier",
            "work_item_id": work_item["id"],
        },
    )
    assert queued.status_code == 202, queued.text
    queued_body = queued.json()

    async def scoped_sender(from_agent_id, args, **_kwargs):
        project_id = uuid.UUID(args["_project_id"])
        target_id = uuid.UUID(args["agent_id"])
        access_id = min(from_agent_id, target_id, key=str)
        peer_id = max(from_agent_id, target_id, key=str)
        async with env.session_factory() as delivery_db:
            delivery_db.add(
                ChatSession(
                    project_id=project_id,
                    agent_id=access_id,
                    peer_agent_id=peer_id,
                    source_channel="agent",
                    title="Worker ↔ Reviewer",
                    external_conv_id=f"project-a2a:{project_id}:{peer_id}",
                )
            )
            await delivery_db.commit()
        return "✅ Notification sent"

    monkeypatch.setattr(project_service, "async_session", env.session_factory)
    monkeypatch.setattr(agent_tools, "_send_message_to_agent", scoped_sender)
    await project_service.deliver_project_a2a(uuid.UUID(queued_body["run_id"]))

    env.db.expire_all()
    run = await env.db.get(ProjectRun, uuid.UUID(queued_body["run_id"]))
    assert run is not None and run.status == "succeeded"
    assert run.output["group_session_id"] == queued_body["group_session_id"]
    assert run.output["session_id"]
    assert run.output["session_agent_id"] == run.output["session_access_agent_id"]
    assert run.output["session_title"] == "Worker ↔ Reviewer"
    delivered = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.run_id == run.id,
                ProjectEvent.event_type == "a2a.delivered",
            )
        )
    ).scalar_one()
    assert delivered.event_metadata["session_id"] == run.output["session_id"]
    assert delivered.event_metadata["group_session_id"] == queued_body["group_session_id"]


async def test_project_a2a_native_receipt_keeps_exact_session_when_pair_has_newer_thread(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.subagent_run import SubagentRun
    from app.services import project_service

    env = project_api
    project = await _create_project(env, name="Exact concurrent A2A receipt")
    project_id = uuid.UUID(project["id"])
    project_run_id = uuid.uuid4()
    access_id = min(env.worker_id, env.reviewer_id, key=str)
    peer_id = max(env.worker_id, env.reviewer_id, key=str)
    exact_session = ChatSession(
        project_id=project_id,
        agent_id=access_id,
        peer_agent_id=peer_id,
        source_channel="agent",
        title="Exact earlier thread",
        external_conv_id=f"a2a-{uuid.uuid4().hex}",
    )
    newer_session = ChatSession(
        project_id=project_id,
        agent_id=access_id,
        peer_agent_id=peer_id,
        source_channel="agent",
        title="Unrelated newer thread",
        external_conv_id=f"a2a-{uuid.uuid4().hex}",
    )
    child_id = uuid.uuid4()
    child_session = ChatSession(
        id=child_id,
        project_id=project_id,
        agent_id=env.reviewer_id,
        source_channel="subagent",
        title="Exact child execution",
    )
    env.db.add_all([exact_session, newer_session, child_session])
    await env.db.flush()
    env.db.add(
        SubagentRun(
            id=child_id,
            parent_session_id=exact_session.id,
            project_id=project_id,
            execution_user_id=env.owner_id,
            origin_tool_call_id=f"test-exact-{uuid.uuid4()}",
            mode="run",
            status="queued",
        )
    )
    await env.db.commit()

    receipt = json.dumps(
        {
            "status": "queued",
            "session_id": str(exact_session.id),
            "a2a_session_id": str(exact_session.id),
            "project_run_id": str(project_run_id),
            "subagent_run_id": str(child_id),
            "subagent_session_id": str(child_id),
        }
    )
    session_info, error = await project_service._resolve_project_a2a_session_info(
        env.db,
        project_id=project_id,
        project_run_id=project_run_id,
        source_agent_id=env.worker_id,
        target_agent_id=env.reviewer_id,
        result=receipt,
    )

    assert error is None
    assert session_info["session_id"] == str(exact_session.id)
    assert session_info["session_id"] != str(newer_session.id)
    assert session_info["subagent_session_id"] == str(child_id)


async def test_project_a2a_native_receipt_rejects_forged_scope_without_latest_fallback(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.subagent_run import SubagentRun
    from app.services import project_service

    env = project_api
    project = await _create_project(env, name="Reject forged A2A receipt")
    project_id = uuid.UUID(project["id"])
    project_run_id = uuid.uuid4()
    access_id = min(env.worker_id, env.reviewer_id, key=str)
    peer_id = max(env.worker_id, env.reviewer_id, key=str)
    valid_latest = ChatSession(
        project_id=project_id,
        agent_id=access_id,
        peer_agent_id=peer_id,
        source_channel="agent",
        title="Valid latest thread",
        external_conv_id=f"a2a-{uuid.uuid4().hex}",
    )
    forged_parent = ChatSession(
        project_id=project_id,
        agent_id=min(env.leader_id, env.reviewer_id, key=str),
        peer_agent_id=max(env.leader_id, env.reviewer_id, key=str),
        source_channel="agent",
        title="Wrong project member pair",
        external_conv_id=f"a2a-{uuid.uuid4().hex}",
    )
    forged_child_id = uuid.uuid4()
    forged_child = ChatSession(
        id=forged_child_id,
        project_id=project_id,
        agent_id=env.reviewer_id,
        source_channel="subagent",
        title="Wrong parent child",
    )
    env.db.add_all([valid_latest, forged_parent, forged_child])
    await env.db.flush()
    env.db.add(
        SubagentRun(
            id=forged_child_id,
            parent_session_id=forged_parent.id,
            project_id=project_id,
            execution_user_id=env.owner_id,
            origin_tool_call_id=f"test-forged-{uuid.uuid4()}",
            mode="run",
            status="queued",
        )
    )
    await env.db.commit()

    forged_receipt = json.dumps(
        {
            "status": "queued",
            "session_id": str(forged_parent.id),
            "a2a_session_id": str(forged_parent.id),
            "project_run_id": str(project_run_id),
            "subagent_run_id": str(forged_child_id),
            "subagent_session_id": str(forged_child_id),
        }
    )
    session_info, error = await project_service._resolve_project_a2a_session_info(
        env.db,
        project_id=project_id,
        project_run_id=project_run_id,
        source_agent_id=env.worker_id,
        target_agent_id=env.reviewer_id,
        result=forged_receipt,
    )

    assert session_info == {}
    assert error == "Project A2A transport receipt references an invalid collaboration session"

    legacy_info, legacy_error = await project_service._resolve_project_a2a_session_info(
        env.db,
        project_id=project_id,
        project_run_id=project_run_id,
        source_agent_id=env.worker_id,
        target_agent_id=env.reviewer_id,
        result="✅ Legacy transport delivered",
    )
    assert legacy_error is None
    assert legacy_info["session_id"] == str(valid_latest.id)


async def test_project_a2a_uses_durable_project_child_and_exact_standard_timeline(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    """A→B stays single-target while B retains project tools and exact trace."""
    from datetime import timedelta

    from app.models.chat_session import ChatSession
    from app.models.project import ProjectEvent, ProjectMemberSnapshot, ProjectRun, ProjectWorkItem
    from app.models.subagent_run import SubagentRun
    from app.services import agent_tools, project_runtime_tools, project_service, subagent_runtime
    from app.services.project_service import freeze_run_members

    env = project_api
    monkeypatch.setattr(agent_tools, "async_session", env.session_factory)
    monkeypatch.setattr(project_service, "async_session", env.session_factory)
    project = await _create_project(env, name="Durable exact project A2A")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)

    dependency = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.leader_id,
        created_by_agent_id=env.leader_id,
        title="Approve the evidence contract",
        description="Define the required evidence before implementation",
        status="done",
        priority="medium",
        acceptance_criteria=["Evidence contract approved"],
        dependency_ids=[],
    )
    env.db.add(dependency)
    await env.db.flush()
    dependency_id = dependency.id
    work_item = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.worker_id,
        created_by_agent_id=env.leader_id,
        title="Write A2A evidence",
        description="B must update this item inside its project runtime",
        status="todo",
        priority="medium",
        acceptance_criteria=["Committed evidence exists", "Evidence cites its source"],
        dependency_ids=[str(dependency_id)],
    )
    parallel_work_item = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.reviewer_id,
        created_by_agent_id=env.leader_id,
        title="Review parallel A2A evidence",
        description="A concurrent delegation must preserve this exact work-item lineage",
        status="todo",
        priority="medium",
        acceptance_criteria=["Review evidence exists"],
        dependency_ids=[],
    )
    env.db.add_all([work_item, parallel_work_item])
    await env.db.flush()
    env.db.add(
        ProjectEvent(
            tenant_id=env.tenant_id,
            project_id=project_id,
            work_item_id=work_item.id,
            actor_agent_id=env.leader_id,
            event_type="work_item.updated",
            summary="Initial evidence attached",
            event_metadata={"evidence": ["docs/evidence-contract.md", "commit-contract-1234"]},
        )
    )
    await env.db.commit()
    work_item_id = work_item.id
    parallel_work_item_id = parallel_work_item.id

    parent_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        work_item_id=work_item_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="running",
        trigger_type="manual",
        input={"title": "Coordinate A2A evidence"},
        output={},
    )
    env.db.add(parent_run)
    await env.db.flush()
    project_row = await env.db.get(Project, project_id)
    assert project_row is not None
    await freeze_run_members(env.db, project_row, parent_run)
    await env.db.commit()

    raw_result = await agent_tools._send_message_to_agent(
        env.leader_id,
        {
            "agent_id": str(env.worker_id),
            "message": "Write docs/a2a-evidence.md and mark the assigned item done.",
            "msg_type": "task_delegate",
            "force_async": True,
            "_project_id": str(project_id),
            "_parent_project_run_id": str(parent_run.id),
        },
        user_id=env.owner_id,
        origin_session_id=None,
        tool_call_id="leader-a2a-worker-1",
    )
    result = json.loads(raw_result)
    assert result["status"] == "queued"
    assert result["awakened_agent_ids"] == [str(env.worker_id)]
    a2a_session_id = uuid.UUID(result["a2a_session_id"])
    child_id = uuid.UUID(result["subagent_session_id"])
    project_run_id = uuid.UUID(result["project_run_id"])

    a2a_session = await env.db.get(ChatSession, a2a_session_id)
    child = await env.db.get(ChatSession, child_id)
    durable = await env.db.get(SubagentRun, child_id)
    run = await env.db.get(ProjectRun, project_run_id)
    assert a2a_session is not None and a2a_session.source_channel == "agent"
    assert a2a_session.project_id == project_id
    assert child is not None and child.source_channel == "subagent"
    assert child.project_id == project_id and child.agent_id == env.worker_id
    assert durable is not None and durable.parent_session_id == a2a_session_id
    assert durable.project_member_id is not None
    assert run is not None and run.status in {"queued", "running"}
    assert run.work_item_id == work_item_id
    assert run.input["parent_project_run_id"] == str(parent_run.id)
    assert run.input["title"] == "Coordinate A2A evidence"
    assert run.input["work_item_snapshot"] == {
        "id": str(work_item_id),
        "title": "Write A2A evidence",
        "description": "B must update this item inside its project runtime",
        "status": "todo",
        "acceptance_criteria": ["Committed evidence exists", "Evidence cites its source"],
        "dependencies": [
            {
                "id": str(dependency_id),
                "title": "Approve the evidence contract",
                "status": "done",
            }
        ],
        "evidence": ["docs/evidence-contract.md", "commit-contract-1234"],
    }

    # Later edits must not rewrite the exact context frozen for this A2A turn.
    current_work_item = await env.db.get(ProjectWorkItem, work_item_id)
    assert current_work_item is not None
    current_work_item.dependency_ids = []
    await env.db.commit()
    env.db.expire_all()
    persisted_run = await env.db.get(ProjectRun, project_run_id)
    assert persisted_run is not None
    assert persisted_run.input["work_item_snapshot"]["dependencies"][0]["id"] == str(dependency_id)
    run = persisted_run

    assert run.output["session_id"] == str(a2a_session_id)
    assert run.output["subagent_session_id"] == str(child_id)

    serialized_runs = (await env.client.get(f"/api/projects/{project_id}/runs")).json()
    serialized_run = next(row for row in serialized_runs if row["id"] == str(run.id))
    assert serialized_run["work_item_id"] == str(work_item_id)
    assert serialized_run["title"] == "Coordinate A2A evidence"

    tool_names = {
        item["function"]["name"]
        for item in await subagent_runtime.prepare_subagent_tools(
            env.worker_id,
            child_id,
            execution_user_id=env.owner_id,
        )
    }
    assert {"write_file", "project_update_work_item", "project_message_agent"} <= tool_names

    write_result = json.loads(
        await project_runtime_tools.execute_project_workspace_tool(
            "write_file",
            {
                "workspace": "project",
                "path": "docs/a2a-evidence.md",
                "content": "# Exact A2A evidence\n",
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-write-a2a-evidence",
            turn_anchor_id=None,
        )
    )
    assert write_result["commit"]
    await project_runtime_tools.execute_project_runtime_tool(
        "project_update_work_item",
        {
            "work_item_id": str(work_item_id),
            "status": "done",
            "progress_note": "Evidence committed through project A2A child",
            "evidence": ["docs/a2a-evidence.md", write_result["commit"]],
        },
        agent_id=env.worker_id,
        execution_user_id=env.owner_id,
        session_id=str(child_id),
        tool_call_id="worker-finish-a2a-item",
        turn_anchor_id=None,
    )

    child_input = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["project_run_id"].as_string() == str(project_run_id),
            )
        )
    ).scalar_one()
    assert "Immutable work-item snapshot" in child_input.content
    assert "Evidence cites its source" in child_input.content
    assert "Approve the evidence contract" in child_input.content
    assert "docs/evidence-contract.md" in child_input.content

    message_schema = project_runtime_tools.PROJECT_TOOL_REGISTRY["project_message_agent"]["function"]
    assert message_schema["parameters"]["properties"]["mode"]["enum"] == ["task_delegate", "consult"]
    assert set(message_schema["parameters"]["required"]) == {
        "agent_id",
        "title",
        "message",
        "mode",
        "expected_output",
    }
    assert message_schema["description"] == (
        "Send one active project member a review request or assigned task. Include the relevant context, requested "
        "work, expected result, and related work item when one exists."
    )
    assert message_schema["parameters"]["properties"]["mode"]["description"] == (
        "Choose task_delegate for assigned work and consult for a review or decision."
    )
    with pytest.raises(ValueError, match="成员协作请求需要明确任务或咨询内容。"):
        await project_runtime_tools.execute_project_runtime_tool(
            "project_message_agent",
            {
                "agent_id": str(env.reviewer_id),
                "message": "FYI: evidence is ready; acknowledge receipt and wait.",
                "mode": "notify",
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-passive-notify-blocked",
            turn_anchor_id=child_input.id,
        )
    with pytest.raises(ValueError, match="成员协作请求需要填写标题。"):
        await project_runtime_tools.execute_project_runtime_tool(
            "project_message_agent",
            {
                "agent_id": str(env.reviewer_id),
                "message": "Evaluate the committed evidence and return a decision with cited findings.",
                "mode": "task_delegate",
                "expected_output": "A cited pass/fail decision.",
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-untitled-delegation-blocked",
            turn_anchor_id=child_input.id,
        )
    with pytest.raises(ValueError, match="成员协作请求需要包含可执行的工作内容和预期结果。"):
        await project_runtime_tools.execute_project_runtime_tool(
            "project_message_agent",
            {
                "agent_id": str(env.reviewer_id),
                "title": "Progress notification",
                "message": "已完成，已更新工作项，请收到后等待。",
                "mode": "consult",
                "expected_output": "Acknowledge receipt.",
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-mechanical-handoff-blocked",
            turn_anchor_id=child_input.id,
        )

    delegated_result = json.loads(
        await project_runtime_tools.execute_project_runtime_tool(
            "project_message_agent",
            {
                "agent_id": str(env.reviewer_id),
                "title": "Review A2A evidence",
                "message": "Review the committed A2A evidence against acceptance criteria.",
                "mode": "task_delegate",
                "expected_output": "A pass/fail decision with exact evidence references.",
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-delegate-a2a-review",
            turn_anchor_id=child_input.id,
        )
    )
    delegated_run = await env.db.get(
        ProjectRun,
        uuid.UUID(delegated_result["project_run_id"]),
    )
    assert delegated_run is not None
    assert delegated_run.work_item_id == work_item_id
    assert delegated_run.input["parent_project_run_id"] == str(run.id)
    assert delegated_run.input["title"] == "Review A2A evidence"

    explicit_result = json.loads(
        await project_runtime_tools.execute_project_runtime_tool(
            "project_message_agent",
            {
                "agent_id": str(env.reviewer_id),
                "work_item_id": str(parallel_work_item_id),
                "title": "Review the parallel evidence stream",
                "message": "Review only the evidence for the explicitly referenced parallel work item.",
                "mode": "task_delegate",
                "expected_output": "An independent review decision for the parallel work item.",
                "new_conversation": True,
            },
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="worker-delegate-explicit-parallel-item",
            turn_anchor_id=child_input.id,
        )
    )
    explicit_run = await env.db.get(
        ProjectRun,
        uuid.UUID(explicit_result["project_run_id"]),
    )
    assert explicit_run is not None
    assert explicit_run.work_item_id == parallel_work_item_id
    assert explicit_run.input["parent_project_run_id"] == str(run.id)
    assert explicit_run.input["work_item_id"] == str(parallel_work_item_id)

    child_input.message_meta = {
        **dict(child_input.message_meta or {}),
        "subagent_input_state": "processing",
        "subagent_turn_anchor_id": str(child_input.id),
        "turn_status": "running",
    }
    durable.status = "running"
    durable.lease_owner = subagent_runtime.settings.INSTANCE_ID
    base_time = child_input.created_at + timedelta(seconds=1)
    env.db.add_all(
        [
            ChatMessage(
                agent_id=env.worker_id,
                sender_agent_id=env.worker_id,
                role="assistant",
                content="",
                thinking="I should write the evidence before completing the item.",
                conversation_id=str(child_id),
                message_meta={"turn_anchor_id": str(child_input.id)},
                created_at=base_time,
            ),
            ChatMessage(
                agent_id=env.worker_id,
                sender_agent_id=env.worker_id,
                role="tool_call",
                content=json.dumps(
                    {
                        "name": "write_file",
                        "call_id": "worker-write-a2a-evidence",
                        "args": {"workspace": "project", "path": "docs/a2a-evidence.md"},
                        "status": "done",
                        "result": write_result,
                    }
                ),
                conversation_id=str(child_id),
                message_meta={"turn_anchor_id": str(child_input.id)},
                created_at=base_time + timedelta(seconds=1),
            ),
        ]
    )
    await env.db.commit()

    assert (
        await subagent_runtime._finish_subagent_turn(
            run_id=child_id,
            anchor_id=child_input.id,
            reply="Evidence committed and assigned work item completed.",
            failed=False,
        )
        is True
    )
    completion = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["kind"].as_string() == "subagent_completion",
            )
        )
    ).scalar_one()
    assert completion.message_meta["subagent_wake"] is True
    assert await subagent_runtime._dispatch_parent_event(completion.id) is True

    visible_rows = (
        (
            await env.db.execute(
                select(ChatMessage)
                .where(ChatMessage.conversation_id == str(a2a_session_id))
                .order_by(ChatMessage.created_at, ChatMessage.id)
            )
        )
        .scalars()
        .all()
    )
    assert visible_rows[0].sender_agent_id == env.leader_id
    assert visible_rows[0].message_meta["target_agent_id"] == str(env.worker_id)
    assert any(row.thinking for row in visible_rows)
    assert any(row.role == "tool_call" for row in visible_rows)
    final_reply = next(row for row in visible_rows if row.content.startswith("Evidence committed"))
    assert final_reply.role == "assistant"
    assert final_reply.sender_agent_id == env.worker_id
    assert final_reply.message_meta["a2a_session_id"] == str(a2a_session_id)
    assert final_reply.message_meta["subagent_session_id"] == str(child_id)

    # An exact peer-to-peer A2A reply remains visible in its own standard Chat
    # Session and also enters the durable project coordination queue.  It must
    # not directly resume/broadcast; the existing batch dispatcher gives the
    # Leader one coalesced follow-up turn.
    group_session = (
        await env.db.execute(
            select(ChatSession).where(
                ChatSession.project_id == project_id,
                ChatSession.source_channel == "project",
            )
        )
    ).scalar_one()
    group_reply = (
        await env.db.execute(
            select(ChatMessage).where(ChatMessage.external_event_key == f"project-a2a-group-reply:{completion.id}")
        )
    ).scalar_one()
    assert group_reply.conversation_id == str(group_session.id)
    assert group_reply.sender_agent_id == env.worker_id
    assert group_reply.message_meta["source_a2a_session_id"] == str(a2a_session_id)
    assert group_reply.message_meta["leader_batch_state"] == "pending"
    group_reply_id = group_reply.id
    assert (
        await subagent_runtime._dispatch_project_leader_batch(
            group_session.id,
            debounce_seconds=0,
        )
        is True
    )
    env.db.expire_all()
    delivered_group_reply = await env.db.get(ChatMessage, group_reply_id)
    assert delivered_group_reply is not None
    assert delivered_group_reply.message_meta["leader_batch_state"] == "delivered"

    completed_run = await env.db.get(ProjectRun, project_run_id)
    completed_item = await env.db.get(ProjectWorkItem, work_item_id)
    assert completed_run is not None and completed_run.status == "succeeded"
    assert completed_run.output["session_id"] == str(a2a_session_id)
    assert completed_run.output["subagent_session_id"] == str(child_id)
    assert completed_item is not None and completed_item.status == "done"
    terminal_event = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.run_id == project_run_id,
                ProjectEvent.event_type == "run.succeeded",
            )
        )
    ).scalar_one()
    assert terminal_event.event_metadata["session_id"] == str(a2a_session_id)
    assert terminal_event.event_metadata["subagent_session_id"] == str(child_id)

    failure_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.worker_id,
        initiated_by_user_id=env.owner_id,
        status="queued",
        trigger_type="a2a",
        input={"message": "Exercise the failed terminal audit contract"},
        output={"session_id": str(a2a_session_id)},
    )
    env.db.add(failure_run)
    await env.db.flush()
    project_row = await env.db.get(Project, project_id)
    assert project_row is not None
    await freeze_run_members(env.db, project_row, failure_run)
    await env.db.commit()
    failure_run_id = failure_run.id
    await subagent_runtime.append_subagent_message(
        agent_id=env.worker_id,
        parent_session_id=str(a2a_session_id),
        subagent_id=str(child_id),
        message="This test turn intentionally fails.",
        execution_user_id=env.owner_id,
        origin_tool_call_id="worker-a2a-failure-audit",
        project_run_id=failure_run_id,
        input_metadata={"project_a2a": True, "a2a_session_id": str(a2a_session_id)},
    )
    failed_input = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["project_run_id"].as_string() == str(failure_run_id),
            )
        )
    ).scalar_one()
    failed_input.message_meta = {
        **dict(failed_input.message_meta or {}),
        "subagent_input_state": "processing",
        "subagent_turn_anchor_id": str(failed_input.id),
        "turn_status": "running",
    }
    durable = await env.db.get(SubagentRun, child_id)
    assert durable is not None
    durable.status = "running"
    durable.lease_owner = subagent_runtime.settings.INSTANCE_ID
    await env.db.commit()
    assert (
        await subagent_runtime._finish_subagent_turn(
            run_id=child_id,
            anchor_id=failed_input.id,
            reply="Intentional project A2A failure",
            failed=True,
        )
        is True
    )
    failed_terminal_event = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.run_id == failure_run_id,
                ProjectEvent.event_type == "run.failed",
            )
        )
    ).scalar_one()
    assert failed_terminal_event.event_metadata["session_id"] == str(a2a_session_id)
    assert failed_terminal_event.event_metadata["subagent_session_id"] == str(child_id)

    worker_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == project_id,
                ProjectMemberSnapshot.agent_id == env.worker_id,
            )
        )
    ).scalar_one()
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    await project_service.deactivate_project_member(
        env.db,
        stored_project,
        worker_member,
        actor_user_id=env.owner_id,
        reason="A2A membership regression",
    )
    await env.db.commit()
    rejected = await agent_tools._send_message_to_agent(
        env.leader_id,
        {
            "agent_id": str(env.worker_id),
            "message": "This departed member must not wake.",
            "msg_type": "notify",
            "force_async": True,
            "_project_id": str(project_id),
        },
        user_id=env.owner_id,
        tool_call_id="leader-a2a-departed-worker",
    )
    rejection_payload = json.loads(rejected)
    assert rejection_payload["status"] == "error"
    assert rejection_payload["code"] == "project_member_inactive"


async def test_project_group_routes_human_to_leader_and_reuses_durable_children(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectRun

    env = project_api
    project = await _create_project(env, name="Project Agent Group")
    project_id = project["id"]
    await _mark_project_running(env, project_id)

    group_response = await env.client.get(f"/api/projects/{project_id}/group-session")
    assert group_response.status_code == 200, group_response.text
    group = group_response.json()
    assert group["source_channel"] == "project"
    assert group["max_mentions"] == 3

    passive = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Visible update only",
            "mentions": [],
            "attachments": [],
            "client_message_id": "passive-confirmation-anchor",
        },
    )
    assert passive.status_code == 201, passive.text
    passive_body = passive.json()
    assert passive_body["awakened_agent_ids"] == [str(env.leader_id)]
    assert passive_body["default_leader_agent_id"] == str(env.leader_id)
    assert len(passive_body["subagent_runs"]) == 1
    assert passive_body["subagent_runs"][0]["agent_id"] == str(env.leader_id)
    assert passive_body["message"]["message_meta"]["visible_to_group"] is True
    assert passive_body["message"]["message_meta"]["wake_policy"] == "default_leader_plus_structured_mentions"
    assert passive_body["turn"]["phase"] == "active"
    assert passive_body["turn"]["generation"] == 1
    cohort_anchor_id = passive_body["turn"]["turn_anchor_id"]
    passive_run = await env.db.get(
        ProjectRun,
        uuid.UUID(passive_body["subagent_runs"][0]["project_run_id"]),
    )
    assert passive_run is not None
    passive_run.status = "waiting"
    passive_child_session_id = passive_body["subagent_runs"][0]["session_id"]
    pending_tool = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        role="tool_call",
        content=json.dumps(
            {
                "name": "request_confirmation",
                "args": {"force_confirmation": True},
                "status": "pending",
            }
        ),
        conversation_id=passive_child_session_id,
    )
    env.db.add(pending_tool)
    await env.db.commit()
    suspended = await env.client.get(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages"
    )
    assert suspended.status_code == 200
    assert suspended.json()["turn"]["phase"] == "suspended"
    assert suspended.json()["turn"]["generation"] == 1
    replay = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Visible update only",
            "mentions": [],
            "attachments": [],
            "client_message_id": "passive-confirmation-anchor",
        },
    )
    assert replay.status_code == 201, replay.text
    assert replay.json()["idempotent_replay"] is True
    blocked = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Must wait for confirmation",
            "mentions": [],
            "attachments": [],
            "client_message_id": "blocked-while-confirmation-pending",
        },
    )
    assert blocked.status_code == 409, blocked.text
    assert blocked.json()["detail"]["code"] == "project_group_confirmation_pending"
    assert blocked.json()["detail"]["client_message_id"] == "blocked-while-confirmation-pending"
    passive_run.status = "queued"
    pending_tool.content = json.dumps(
        {
            "name": "request_confirmation",
            "args": {"force_confirmation": True},
            "status": "done",
            "result": "confirmed for lifecycle test",
        }
    )
    await env.db.commit()

    mentioned = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Worker build and Reviewer check",
            "llm_content": "[brief.md extracted]\nBuild the evidence and review every acceptance item.",
            "mentions": [str(env.worker_id), str(env.reviewer_id), str(env.worker_id)],
            "attachments": [{"name": "brief.md", "path": "brief.md"}],
            "client_message_id": "mention-1",
        },
    )
    assert mentioned.status_code == 201, mentioned.text
    body = mentioned.json()
    assert body["awakened_agent_ids"] == [
        str(env.worker_id),
        str(env.reviewer_id),
    ]
    assert body["message"]["message_meta"]["wake_policy"] == "structured_mentions_only"
    assert body["message"]["message_meta"]["owner_deferred_until_specialist_result"] is True
    assert len(body["subagent_runs"]) == 3
    assert all(row["project_run_id"] for row in body["subagent_runs"])
    assert body["message"]["display_content"] == "Worker build and Reviewer check"
    assert body["turn"]["phase"] == "active"
    assert body["turn"]["generation"] == 1
    assert body["turn"]["turn_anchor_id"] == cohort_anchor_id
    worker_child = next(row for row in body["subagent_runs"] if row["agent_id"] == str(env.worker_id))
    worker_input = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == worker_child["session_id"],
                ChatMessage.message_meta["project_run_id"].as_string() == worker_child["project_run_id"],
            )
        )
    ).scalar_one()
    assert worker_input.content.startswith("[brief.md extracted]")
    worker_session = await env.db.get(ChatSession, uuid.UUID(worker_child["session_id"]))
    assert worker_session is not None
    assert worker_session.im_config["project_name_snapshot"] == "Project Agent Group"
    assert worker_session.im_config["project_member_name_snapshot"] == "Worker"
    assert worker_session.im_config["project_member_role_snapshot"] == "Build deliverables"
    assert worker_session.im_config["project_role_snapshot"] == "participant"

    replay = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Worker build and Reviewer check",
            "mentions": [str(env.worker_id), str(env.reviewer_id)],
            "client_message_id": "mention-1",
        },
    )
    assert replay.status_code == 201
    assert replay.json()["idempotent_replay"] is True
    assert replay.json()["subagent_runs"] == body["subagent_runs"]

    reused = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Worker follow-up", "mentions": [str(env.worker_id)]},
    )
    assert reused.status_code == 201, reused.text
    reused_worker = next(row for row in reused.json()["subagent_runs"] if row["agent_id"] == str(env.worker_id))
    assert reused_worker["session_id"] == worker_child["session_id"]
    assert reused_worker["project_run_id"] != worker_child["project_run_id"]

    leader_deduped = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Leader and Worker only once",
            "mentions": [str(env.leader_id), str(env.worker_id), str(env.leader_id)],
        },
    )
    assert leader_deduped.status_code == 201, leader_deduped.text
    assert leader_deduped.json()["awakened_agent_ids"] == [
        str(env.leader_id),
        str(env.worker_id),
    ]

    spoofed_agent = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Agent-visible update only",
            "sender_agent_id": str(env.worker_id),
            "mentions": [],
        },
    )
    assert spoofed_agent.status_code == 422

    stored_project = await env.db.get(Project, uuid.UUID(project_id))
    assert stored_project is not None
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "policies": {
            **dict((stored_project.settings or {}).get("policies") or {}),
            "max_a2a_wakes": 12,
        },
    }
    await env.db.commit()
    expanded_group = await env.client.get(f"/api/projects/{project_id}/group-session")
    assert expanded_group.json()["max_mentions"] == 4
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "policies": {
            **dict((stored_project.settings or {}).get("policies") or {}),
            "max_a2a_wakes": 2,
        },
    }
    await env.db.commit()
    limited_group = await env.client.get(f"/api/projects/{project_id}/group-session")
    assert limited_group.json()["max_mentions"] == 1
    over_budget = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Leader plus two participants exceeds total budget",
            "mentions": [str(env.worker_id), str(env.reviewer_id)],
        },
    )
    assert over_budget.status_code == 422

    self_mention = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "No self wake",
            "sender_agent_id": str(env.worker_id),
            "mentions": [str(env.worker_id)],
        },
    )
    assert self_mention.status_code == 422

    history = await env.client.get(f"/api/projects/{project_id}/group-sessions/{group['id']}/messages?limit=500")
    assert history.status_code == 200
    assert len(history.json()["items"]) == 4
    assert history.json()["turn"]["turn_anchor_id"] == cohort_anchor_id
    assert history.json()["turn"]["phase"] == "active"
    attachment_message = next(item for item in history.json()["items"] if item["attachments"])
    assert attachment_message["attachments"][0]["name"] == "brief.md"
    cohort_runs = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.input["group_session_id"].as_string() == str(group["id"])
                )
            )
        ).scalars()
    )
    for cohort_run in cohort_runs:
        cohort_run.status = "succeeded"
        cohort_run.finished_at = datetime.now(UTC)
    await env.db.commit()
    completed_history = await env.client.get(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages?limit=500"
    )
    assert completed_history.status_code == 200
    assert completed_history.json()["turn"]["phase"] == "idle"
    assert completed_history.json()["turn"]["status"] == "completed"
    assert completed_history.json()["turn"]["generation"] == 1


async def test_project_group_partial_completion_versions_cohort_projection(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectRun

    env = project_api
    project = await _create_project(env, name="Versioned group cohort")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    created = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Run two specialists",
            "mentions": [str(env.worker_id), str(env.reviewer_id)],
            "client_message_id": "versioned-cohort-1",
        },
    )
    assert created.status_code == 201, created.text
    initial = created.json()["turn"]
    assert initial["phase"] == "active"
    assert len(initial["active_agent_ids"]) >= 2

    runs = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == uuid.UUID(project_id),
                    ProjectRun.trigger_type.in_(["group_leader_message", "group_mention"]),
                    ProjectRun.input["group_session_id"].as_string() == group["id"],
                )
            )
        ).scalars()
    )
    first = next(run for run in runs if run.agent_id == env.worker_id)
    run_ids = [run.id for run in runs]
    first.status = "failed"
    first.finished_at = datetime.now(UTC)
    await env.db.commit()
    partial = (
        await env.client.get(
            f"/api/projects/{project_id}/group-sessions/{group['id']}/messages"
        )
    ).json()["turn"]
    assert partial["phase"] == "active"
    assert partial["revision"] > initial["revision"]
    assert str(env.worker_id) not in partial["active_agent_ids"]

    env.db.expire_all()
    remaining = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.id.in_(run_ids),
                    ProjectRun.status.not_in(["succeeded", "failed", "cancelled"]),
                )
            )
        ).scalars()
    )
    for run in remaining:
        run.status = "failed"
        run.finished_at = datetime.now(UTC)
    await env.db.commit()
    terminal = (
        await env.client.get(
            f"/api/projects/{project_id}/group-sessions/{group['id']}/messages"
        )
    ).json()["turn"]
    assert terminal["phase"] == "idle"
    assert terminal["revision"] > partial["revision"]
    assert terminal["active_agent_ids"] == []


async def test_explicit_mention_cohort_waits_for_leader_reply_batch_terminal(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectRun
    from app.services.project_group_turn_lifecycle import (
        reconcile_and_publish_project_run_group_turn,
    )

    env = project_api
    project = await _create_project(env, name="Mention owner summary cohort")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, str(project_id))
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    created = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Worker produce the result before the owner summarizes it",
            "mentions": [str(env.worker_id)],
            "client_message_id": "mention-owner-summary-cohort",
        },
    )
    assert created.status_code == 201, created.text
    anchor_id = uuid.UUID(created.json()["message"]["id"])
    runs = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == project_id,
                    ProjectRun.input["group_message_id"].as_string()
                    == str(anchor_id),
                )
            )
        ).scalars()
    )
    owner_run = next(run for run in runs if run.trigger_type == "group_leader_message")
    specialist_run = next(run for run in runs if run.trigger_type == "group_mention")
    assert owner_run.status == "cancelled"
    assert owner_run.output["skip_reason"] == "explicit_mentions_route_to_specialists"

    specialist_run.status = "succeeded"
    specialist_run.finished_at = datetime.now(UTC)
    await env.db.commit()

    reply = ChatMessage(
        agent_id=env.leader_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content="Specialist result ready for owner summary",
        conversation_id=group["id"],
        message_meta={
            "kind": "project_subagent_reply",
            "timeline_anchor_id": str(anchor_id),
            "leader_batch_state": "pending",
            "default_leader_agent_id": str(env.leader_id),
            "source_project_run_ids": [str(specialist_run.id)],
        },
    )
    env.db.add(reply)
    await env.db.commit()
    pending_projection = await reconcile_and_publish_project_run_group_turn(
        specialist_run.id
    )
    assert pending_projection is not None
    assert pending_projection.snapshot.phase == "active"
    assert pending_projection.snapshot.status == "running"
    assert pending_projection.active_agent_ids == (env.leader_id,)

    reply.message_meta = {
        **dict(reply.message_meta or {}),
        "leader_batch_state": "claimed",
    }
    batch_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        execution_user_id=env.owner_id,
        status="queued",
        trigger_type="leader_reply_batch",
        input={
            "group_session_id": group["id"],
            "group_message_id": str(anchor_id),
            "source_group_message_ids": [str(reply.id)],
        },
        output={"group_session_id": group["id"]},
    )
    env.db.add(batch_run)
    await env.db.commit()
    claimed_projection = await reconcile_and_publish_project_run_group_turn(batch_run.id)
    assert claimed_projection is not None
    assert claimed_projection.snapshot.phase == "active"
    assert claimed_projection.run_count == 2
    assert claimed_projection.snapshot.revision > pending_projection.snapshot.revision

    reply.message_meta = {
        **dict(reply.message_meta or {}),
        "leader_batch_state": "delivered",
    }
    batch_run.status = "succeeded"
    batch_run.finished_at = datetime.now(UTC)
    await env.db.commit()
    terminal_projection = await reconcile_and_publish_project_run_group_turn(batch_run.id)
    assert terminal_projection is not None
    assert terminal_projection.snapshot.phase == "idle"
    assert terminal_projection.snapshot.status == "completed"
    assert terminal_projection.snapshot.anchor_id == anchor_id
    assert terminal_projection.run_count == 0


async def test_project_group_reconcile_recomputes_cohort_after_session_mutex(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectRun
    from app.services.project_group_turn_lifecycle import reconcile_project_group_turn

    env = project_api
    project = await _create_project(env, name="Project cohort mutex")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    first = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "First cohort", "mentions": [], "attachments": []},
    )
    assert first.status_code == 201, first.text
    first_turn = first.json()["turn"]
    first_anchor_id = uuid.UUID(first_turn["turn_anchor_id"])

    old_runs = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == project_id,
                    ProjectRun.input["group_session_id"].as_string() == group["id"],
                )
            )
        ).scalars()
    )
    assert old_runs
    for run in old_runs:
        run.status = "succeeded"
        run.finished_at = datetime.now(UTC)
    await env.db.commit()

    async with env.session_factory() as writer_db, env.session_factory() as poll_db:
        locked_session = (
            await writer_db.execute(
                select(ChatSession)
                .where(ChatSession.id == uuid.UUID(group["id"]))
                .with_for_update()
            )
        ).scalar_one()
        second_anchor = ChatMessage(
            agent_id=env.leader_id,
            user_id=env.owner_id,
            sender_user_id=env.owner_id,
            role="user",
            content="Second message joins while poll waits",
            conversation_id=group["id"],
            message_meta={
                "kind": "project_group_message",
                "project_id": str(project_id),
                "visible_to_group": True,
            },
        )
        writer_db.add(second_anchor)
        await writer_db.flush()
        writer_db.add(
            ProjectRun(
                tenant_id=env.tenant_id,
                project_id=project_id,
                agent_id=env.worker_id,
                initiated_by_user_id=env.owner_id,
                execution_user_id=env.owner_id,
                status="queued",
                trigger_type="group_mention",
                input={
                    "group_session_id": group["id"],
                    "group_message_id": str(second_anchor.id),
                },
                output={"group_session_id": group["id"]},
            )
        )
        await writer_db.flush()

        poll_session = await poll_db.get(ChatSession, locked_session.id)
        poll_pid = await poll_db.scalar(text("SELECT pg_backend_pid()"))
        poll_task = asyncio.create_task(
            reconcile_project_group_turn(
                poll_db,
                project_id=project_id,
                session=poll_session,
            )
        )

        async with env.session_factory() as observer_db:
            for _ in range(100):
                wait_type = await observer_db.scalar(
                    text(
                        "SELECT wait_event_type FROM pg_stat_activity "
                        "WHERE pid = :pid"
                    ),
                    {"pid": poll_pid},
                )
                if wait_type == "Lock":
                    break
                await asyncio.sleep(0.01)
            else:
                raise AssertionError("project cohort reconcile did not wait on the session mutex")

        await writer_db.commit()
        projection = await asyncio.wait_for(poll_task, timeout=2)
        await poll_db.commit()

    assert projection.snapshot.phase == "active"
    assert projection.snapshot.generation == 1
    assert projection.snapshot.anchor_id == first_anchor_id
    assert projection.run_count == 1
    assert projection.active_agent_ids == (env.worker_id,)


async def test_project_group_reconcile_refreshes_preloaded_session_pointer(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectRun
    from app.services.conversation_turn_lifecycle import transition_conversation_turn
    from app.services.project_group_turn_lifecycle import reconcile_project_group_turn

    env = project_api
    project = await _create_project(env, name="Project pointer refresh")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    first = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Generation A", "mentions": [], "attachments": []},
    )
    assert first.status_code == 201, first.text
    first_anchor_id = uuid.UUID(first.json()["turn"]["turn_anchor_id"])

    old_runs = list(
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == project_id,
                    ProjectRun.input["group_session_id"].as_string() == group["id"],
                )
            )
        ).scalars()
    )
    for run in old_runs:
        run.status = "succeeded"
        run.finished_at = datetime.now(UTC)
    await env.db.commit()

    async with env.session_factory() as poll_db, env.session_factory() as writer_db:
        preloaded = await poll_db.get(ChatSession, uuid.UUID(group["id"]))
        assert preloaded is not None
        assert preloaded.im_config["conversation_turn"]["turn_anchor_id"] == str(first_anchor_id)

        second_anchor = ChatMessage(
            agent_id=env.leader_id,
            user_id=env.owner_id,
            sender_user_id=env.owner_id,
            role="user",
            content="Generation B",
            conversation_id=group["id"],
            message_meta={
                "kind": "project_group_message",
                "project_id": str(project_id),
                "visible_to_group": True,
            },
        )
        writer_db.add(second_anchor)
        await writer_db.flush()
        writer_db.add(
            ProjectRun(
                tenant_id=env.tenant_id,
                project_id=project_id,
                agent_id=env.reviewer_id,
                initiated_by_user_id=env.owner_id,
                execution_user_id=env.owner_id,
                status="queued",
                trigger_type="group_mention",
                input={
                    "group_session_id": group["id"],
                    "group_message_id": str(second_anchor.id),
                },
                output={"group_session_id": group["id"]},
            )
        )
        await writer_db.flush()
        await transition_conversation_turn(
            writer_db,
            agent_id=env.leader_id,
            conversation_id=group["id"],
            turn_anchor_id=first_anchor_id,
            status="completed",
        )
        await transition_conversation_turn(
            writer_db,
            agent_id=env.leader_id,
            conversation_id=group["id"],
            turn_anchor_id=second_anchor.id,
            status="running",
        )
        await writer_db.commit()

        projection = await reconcile_project_group_turn(
            poll_db,
            project_id=project_id,
            session=preloaded,
        )
        await poll_db.commit()

    assert projection.snapshot.phase == "active"
    assert projection.snapshot.generation == 2
    assert projection.snapshot.anchor_id == second_anchor.id
    assert projection.run_count == 1
    assert projection.active_agent_ids == (env.reviewer_id,)


async def test_project_group_dispatch_outbox_recovers_on_idempotent_replay(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import ProjectMemberSnapshot, ProjectRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Recoverable project dispatch")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    live_events: list[tuple[str, str, dict]] = []

    async def capture_live(agent_id, conversation_id, payload):
        live_events.append((str(agent_id), str(conversation_id), payload))

    monkeypatch.setattr(
        "app.api.websocket.manager.send_to_session",
        capture_live,
    )
    real_dispatch = subagent_runtime.dispatch_project_run

    async def simulated_process_exit(_run_id):
        raise RuntimeError("simulated exit after outbox commit")

    monkeypatch.setattr(subagent_runtime, "dispatch_project_run", simulated_process_exit)
    first = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "This message must survive a dispatcher exit",
            "mentions": [str(env.worker_id)],
            "client_message_id": "recoverable-group-message",
        },
    )
    assert first.status_code == 201, first.text
    assert first.json()["awakened_agent_ids"] == []
    pending = (
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == uuid.UUID(project_id),
                    ProjectRun.trigger_type.in_(["group_leader_message", "group_mention"]),
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(pending) == 2
    assert all(row.status == "queued" and row.input["dispatch"]["task"] for row in pending)
    pending_ids = [row.id for row in pending]
    worker_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == uuid.UUID(project_id),
                ProjectMemberSnapshot.agent_id == env.worker_id,
            )
        )
    ).scalar_one()
    worker_member.is_enabled = False
    await env.db.commit()
    live_events.clear()

    monkeypatch.setattr(subagent_runtime, "dispatch_project_run", real_dispatch)
    for pending_id in pending_ids:
        await real_dispatch(pending_id)
    group_turn_events = [
        payload
        for _agent_id, conversation_id, payload in live_events
        if conversation_id == group["id"] and payload.get("type") == "turn_state"
    ]
    assert group_turn_events
    assert group_turn_events[-1]["turn"]["phase"] == "idle"
    env.db.expire_all()
    recovered = (await env.db.execute(select(ProjectRun).where(ProjectRun.id.in_(pending_ids)))).scalars().all()
    recovered_by_trigger = {row.trigger_type: row for row in recovered}
    assert recovered_by_trigger["group_mention"].status == "failed"
    assert recovered_by_trigger["group_leader_message"].status == "cancelled"
    assert recovered_by_trigger["group_leader_message"].output == {
        "group_session_id": group["id"],
        "status": "skipped",
        "skip_reason": "explicit_mentions_route_to_specialists",
    }


async def test_planning_group_message_only_wakes_owner_without_execution_tools(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import ProjectRun
    from app.services.subagent_runtime import prepare_subagent_tools

    env = project_api
    project = await _create_project(env, name="Human-controlled planning")
    project_id = project["id"]
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()

    async def inherited_tools(_agent_id, *, assignment_snapshot=None):
        return [
            {
                "type": "function",
                "function": {
                    "name": "inherited_write_tool",
                    "description": "Must not be available while planning",
                    "parameters": {"type": "object", "properties": {}},
                },
            }
        ]

    monkeypatch.setattr("app.services.agent_tools.get_agent_tools_for_llm", inherited_tools)
    response = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Help me turn this request into a reviewable delivery plan.",
            "mentions": [],
            "client_message_id": "planning-owner-only",
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["awakened_agent_ids"] == [str(env.leader_id)]
    assert len(body["subagent_runs"]) == 1
    assert body["subagent_runs"][0]["agent_id"] == str(env.leader_id)

    project_run = await env.db.get(ProjectRun, uuid.UUID(body["subagent_runs"][0]["project_run_id"]))
    assert project_run is not None
    assert project_run.trigger_type == "group_leader_message"
    planning_task = project_run.input["dispatch"]["task"]
    assert "project is still in planning" in planning_task
    assert "Do not create or update work items" in planning_task
    assert "do not begin delivery" in planning_task

    child_session_id = uuid.UUID(body["subagent_runs"][0]["session_id"])
    assert (
        await prepare_subagent_tools(
            env.leader_id,
            child_session_id,
            execution_user_id=env.owner_id,
        )
        == []
    )

    mentioned = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Ask the specialist to start now.",
            "mentions": [str(env.worker_id)],
        },
    )
    assert mentioned.status_code == 422
    assert "project owner" in mentioned.json()["detail"]

    stored_project = await env.db.get(Project, uuid.UUID(project_id))
    assert stored_project is not None
    stored_project.status = "paused"
    await env.db.commit()
    paused = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Summarize the current paused state.", "mentions": []},
    )
    assert paused.status_code == 201, paused.text
    paused_body = paused.json()
    assert paused_body["awakened_agent_ids"] == [str(env.leader_id)]
    assert len(paused_body["subagent_runs"]) == 1
    paused_run = await env.db.get(
        ProjectRun,
        uuid.UUID(paused_body["subagent_runs"][0]["project_run_id"]),
    )
    assert paused_run is not None
    assert paused_run.agent_id == env.leader_id
    assert paused_run.input["dispatch"]["execution_tools_enabled"] is False
    assert paused_run.input["dispatch"]["read_only_conversation"] is True
    assert "project is paused" in paused_run.input["dispatch"]["task"]
    assert "Do not restart delivery" in paused_run.input["dispatch"]["task"]
    await env.db.refresh(stored_project)
    assert stored_project.status == "paused"
    paused_mention = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Wake the specialist.", "mentions": [str(env.worker_id)]},
    )
    assert paused_mention.status_code == 422

    # The boundary belongs to the queued message, not the project's later
    # status. Resuming must not retroactively add tools to this advisory turn.
    stored_project.status = "running"
    await env.db.commit()
    assert (
        await prepare_subagent_tools(
            env.leader_id,
            uuid.UUID(paused_body["subagent_runs"][0]["session_id"]),
            execution_user_id=env.owner_id,
        )
        == []
    )

    stored_project.status = "completed"
    await env.db.commit()
    completed = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Explain the completed result.", "mentions": []},
    )
    assert completed.status_code == 201, completed.text
    completed_body = completed.json()
    assert completed_body["awakened_agent_ids"] == [str(env.leader_id)]
    assert len(completed_body["subagent_runs"]) == 1
    completed_run = await env.db.get(
        ProjectRun,
        uuid.UUID(completed_body["subagent_runs"][0]["project_run_id"]),
    )
    assert completed_run is not None
    assert completed_run.agent_id == env.leader_id
    assert completed_run.input["dispatch"]["execution_tools_enabled"] is False
    assert completed_run.input["dispatch"]["read_only_conversation"] is True
    assert "project is completed" in completed_run.input["dispatch"]["task"]
    assert "Do not restart delivery" in completed_run.input["dispatch"]["task"]
    await env.db.refresh(stored_project)
    assert stored_project.status == "completed"
    completed_mention = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Wake the specialist.", "mentions": [str(env.worker_id)]},
    )
    assert completed_mention.status_code == 422


async def test_planning_leader_session_never_gets_project_runtime_tools(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.services.project_runtime_tools import (
        PROJECT_RUNTIME_TOOL_NAMES,
        execute_project_runtime_tool,
    )
    from app.services.subagent_runtime import prepare_subagent_tools

    env = project_api
    project = await _create_project(env, name="Planning scope isolation")
    project_id = project["id"]

    async def no_normal_tools(_agent_id, *, assignment_snapshot=None):
        return []

    monkeypatch.setattr("app.services.agent_tools.get_agent_tools_for_llm", no_normal_tools)
    response = await env.client.get(f"/api/projects/{project_id}/leader-session")
    assert response.status_code == 200, response.text
    assert response.json()["source_channel"] == "web"
    assert response.json()["read_only"] is True
    assert response.json()["planning_transport"] == "project_group"

    planning_session_id = uuid.UUID(response.json()["id"])
    planning_tools = {
        item["function"]["name"]
        for item in await prepare_subagent_tools(
            env.leader_id,
            planning_session_id,
            execution_user_id=env.owner_id,
        )
    }
    assert planning_tools.isdisjoint(PROJECT_RUNTIME_TOOL_NAMES)
    with pytest.raises(ValueError, match="authorized project Subagent runtime"):
        await execute_project_runtime_tool(
            "project_get_context",
            {},
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(planning_session_id),
            tool_call_id="planning-session-scope-bypass",
            turn_anchor_id=None,
        )


async def test_owner_can_resume_waiting_project(project_api: ProjectApiEnv):
    from app.models.project import Project, ProjectEvent

    env = project_api
    project = await _create_project(env, name="Waiting project can resume")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.status = "waiting"
    await env.db.commit()

    resumed = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"status": "running"},
    )

    assert resumed.status_code == 200, resumed.text
    assert resumed.json()["status"] == "running"
    event = await env.db.scalar(
        select(ProjectEvent)
        .where(
            ProjectEvent.project_id == project_id,
            ProjectEvent.event_type == "project.resumed",
        )
        .order_by(ProjectEvent.created_at.desc())
    )
    assert event is not None
    assert event.event_metadata["previous_status"] == "waiting"


async def test_kickoff_requires_leader_discussion_then_freezes_and_starts(project_api: ProjectApiEnv):
    from app.models.project import Project, ProjectRun, ProjectRunMemberSnapshot
    from app.models.subagent_run import SubagentRun
    from app.services.subagent_runtime import prepare_subagent_tools

    env = project_api
    project = await _create_project(env, name="Plan before execution")
    project_id = project["id"]
    assert project["status"] == "planning"

    first = await env.client.get(f"/api/projects/{project_id}/leader-session")
    second = await env.client.get(f"/api/projects/{project_id}/leader-session")
    assert first.status_code == 200, first.text
    assert first.json()["id"] == second.json()["id"]
    assert first.json()["source_channel"] == "web"
    assert first.json()["agent_id"] == str(env.leader_id)
    assert first.json()["read_only"] is True
    assert first.json()["planning_transport"] == "project_group"

    bypass = await env.client.patch(f"/api/projects/{project_id}", json={"status": "running"})
    assert bypass.status_code == 409

    no_discussion = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Start now"},
    )
    assert no_discussion.status_code == 422

    session_id = first.json()["id"]
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    planning_started_at = datetime.now(UTC)
    env.db.add_all(
        [
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Plan the delivery around traceable evidence.",
                conversation_id=group["id"],
                message_meta={"kind": "project_group_message", "visible_to_group": True},
                created_at=planning_started_at,
            ),
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                sender_agent_id=env.leader_id,
                role="assistant",
                content="I will coordinate the team and commit every deliverable.",
                conversation_id=group["id"],
                message_meta={"kind": "project_subagent_reply", "visible_to_group": True},
                created_at=planning_started_at + timedelta(seconds=1),
            ),
        ]
    )
    await env.db.commit()

    confirmed = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Approved. Execute this plan."},
    )
    assert confirmed.status_code == 202, confirmed.text
    body = confirmed.json()
    assert body["status"] == "running"
    assert body["leader_session_id"] == session_id
    assert body["discussion_source"] == "project_group"
    assert body["discussion_session_id"] == group["id"]
    assert body["awakened_agent_ids"] == [str(env.leader_id)]
    assert body["subagent_session_id"] == body["subagent_run_id"]
    assert body["git_start_commit"] != body["transcript_commit"]

    env.db.expire_all()
    stored_project = await env.db.get(Project, uuid.UUID(project_id))
    run = await env.db.get(ProjectRun, uuid.UUID(body["run_id"]))
    child = await env.db.get(SubagentRun, uuid.UUID(body["subagent_run_id"]))
    assert stored_project is not None and stored_project.status == "running"
    assert stored_project.settings["planning"]["state"] == "confirmed"
    assert stored_project.settings["planning"]["launch_confirmed"] is True
    assert stored_project.settings["kickoff"]["project_run_id"] == body["run_id"]
    assert run is not None and run.trigger_type == "leader_kickoff"
    assert run.input["conversation_snapshot"]["message_count"] == 2
    assert run.input["conversation_snapshot"]["transcript_sha256"]
    kickoff_task = run.input["dispatch"]["task"]
    assert "first substantive delivery decision" in kickoff_task
    assert "explicit dependency-aware work-item plan" in kickoff_task
    assert "Do not ask members to acknowledge, wait, or provide routine progress updates" in kickoff_task
    assert "Creating or assigning a work item does not wake its assignee" in kickoff_task
    assert "exactly one project A2A task_delegate" in kickoff_task
    assert "Do not mark a delegated work item in progress without that exact handoff" in kickoff_task
    assert "Do not expose internal narration" in kickoff_task
    assert "report progress to the group" not in kickoff_task
    assert run.output["transcript_commit"] == body["transcript_commit"]
    snapshots = (
        (await env.db.execute(select(ProjectRunMemberSnapshot).where(ProjectRunMemberSnapshot.run_id == run.id)))
        .scalars()
        .all()
    )
    assert len(snapshots) == 3
    assert child is not None and child.project_id == uuid.UUID(project_id)
    assert child.project_member_id == next(item.project_member_id for item in snapshots if item.is_leader)
    kickoff_tool_names = {
        item["function"]["name"]
        for item in await prepare_subagent_tools(
            env.leader_id,
            child.id,
            execution_user_id=env.owner_id,
        )
    }
    assert {"project_get_context", "project_create_work_item", "project_update_plan"} <= kickoff_tool_names

    transcript = subprocess.run(
        [
            "git",
            "-C",
            str(project_repo_path(env.tenant_id, uuid.UUID(project_id))),
            "show",
            f"{body['transcript_commit']}:docs/kickoff-transcript.md",
        ],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    assert "Plan the delivery around traceable evidence." in transcript
    assert "I will coordinate the team" in transcript
    assert "Approved. Execute this plan." in transcript

    event = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.project_id == uuid.UUID(project_id),
                ProjectEvent.event_type == "project.kickoff.confirmed",
            )
        )
    ).scalar_one()
    assert event.run_id == run.id
    assert event.event_metadata["awakened_agent_ids"] == [str(env.leader_id)]
    assert event.event_metadata["transcript_commit"] == body["transcript_commit"]
    kickoff_group_message = (
        await env.db.execute(
            select(ChatMessage).where(ChatMessage.external_event_key.like(f"project-kickoff:{project_id}:%"))
        )
    ).scalar_one()
    assert kickoff_group_message.message_meta["mentions"] == []
    assert kickoff_group_message.message_meta["awakened_agent_ids"] == [str(env.leader_id)]

    repeated = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Do not start twice"},
    )
    assert repeated.status_code == 409


async def test_kickoff_accepts_human_leader_discussion_from_project_group(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectRun

    env = project_api
    project = await _create_project(env, name="Group planning source")
    group = (await env.client.get(f"/api/projects/{project['id']}/group-session")).json()
    planning_started_at = datetime.now(UTC)
    env.db.add_all(
        [
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Keep the plan small and make every result traceable.",
                conversation_id=group["id"],
                message_meta={"kind": "project_group_message", "visible_to_group": True},
                created_at=planning_started_at,
            ),
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                sender_agent_id=env.leader_id,
                role="assistant",
                content="I will deliver in two milestones with Git evidence.",
                conversation_id=group["id"],
                message_meta={"kind": "project_subagent_reply", "visible_to_group": True},
                created_at=planning_started_at + timedelta(seconds=1),
            ),
            # Participant replies remain in the group audit log but are not
            # part of the Human/Leader agreement frozen at kickoff.
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                sender_agent_id=env.worker_id,
                role="assistant",
                content="Worker side note",
                conversation_id=group["id"],
                message_meta={"kind": "project_subagent_reply", "visible_to_group": True},
                created_at=planning_started_at + timedelta(seconds=2),
            ),
        ]
    )
    await env.db.commit()

    confirmed = await env.client.post(
        f"/api/projects/{project['id']}/kickoff/confirm",
        json={"confirmation": "方案确认，开始执行。"},
    )
    assert confirmed.status_code == 202, confirmed.text
    payload = confirmed.json()
    assert payload["status"] == "running"
    assert payload["discussion_source"] == "project_group"
    assert payload["discussion_session_id"] == group["id"]

    run = await env.db.get(ProjectRun, uuid.UUID(payload["run_id"]))
    assert run is not None
    assert run.input["conversation_snapshot"]["source"] == "project_group"
    assert run.input["conversation_snapshot"]["session_id"] == group["id"]
    assert run.input["conversation_snapshot"]["message_count"] == 2

    transcript = subprocess.run(
        [
            "git",
            "-C",
            str(project_repo_path(env.tenant_id, uuid.UUID(project["id"]))),
            "show",
            f"{payload['transcript_commit']}:docs/kickoff-transcript.md",
        ],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    assert "Keep the plan small" in transcript
    assert "I will deliver in two milestones" in transcript
    assert "Worker side note" not in transcript


async def test_kickoff_waits_for_active_project_planning_run(project_api: ProjectApiEnv):
    from app.models.project import ProjectRun

    env = project_api
    project = await _create_project(env, name="Planning run barrier")
    project_id = uuid.UUID(project["id"])
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    env.db.add_all(
        [
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Prepare the final delivery plan.",
                conversation_id=group["id"],
                message_meta={"kind": "project_group_message", "visible_to_group": True},
            ),
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                sender_agent_id=env.leader_id,
                role="assistant",
                content="The plan is ready for approval.",
                conversation_id=group["id"],
                message_meta={"kind": "project_subagent_reply", "visible_to_group": True},
            ),
            ProjectRun(
                tenant_id=env.tenant_id,
                project_id=project_id,
                agent_id=env.leader_id,
                initiated_by_user_id=env.owner_id,
                status="running",
                trigger_type="group_leader_message",
                input={"title": "Finish planning"},
            ),
        ]
    )
    await env.db.commit()

    blocked = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Start only after planning settles"},
    )
    assert blocked.status_code == 409
    assert "still being processed" in blocked.json()["detail"]


async def test_legacy_leader_session_kickoff_waits_for_terminal_reply(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Legacy planning turn barrier")
    project_id = uuid.UUID(project["id"])
    stored_project = await env.db.get(Project, project_id)
    assert stored_project is not None
    stored_project.settings = {
        **dict(stored_project.settings or {}),
        "planning": {
            **dict(dict(stored_project.settings or {}).get("planning") or {}),
            "conversation_mode": "leader_session",
        },
    }
    await env.db.commit()
    leader_session = (await env.client.get(f"/api/projects/{project_id}/leader-session")).json()
    assert leader_session["read_only"] is True
    assert leader_session["planning_transport"] == "leader_session"

    planning_started_at = datetime.now(UTC)
    anchor = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Finish the plan before kickoff.",
        conversation_id=leader_session["id"],
        message_meta={},
        created_at=planning_started_at,
    )
    env.db.add(anchor)
    await env.db.flush()
    anchor_id = anchor.id
    await env.db.commit()

    blocked = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Do not overlap turns"},
    )
    assert blocked.status_code == 409
    assert "still being processed" in blocked.json()["detail"]

    env.db.add(
        ChatMessage(
            agent_id=env.leader_id,
            sender_agent_id=env.leader_id,
            role="assistant",
            content="The plan is now complete.",
            conversation_id=leader_session["id"],
            message_meta={"turn_anchor_id": str(anchor_id), "turn_status": "completed"},
            created_at=planning_started_at + timedelta(seconds=1),
        )
    )
    await env.db.commit()
    confirmed = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "The completed plan is approved"},
    )
    assert confirmed.status_code == 202, confirmed.text
    assert confirmed.json()["discussion_source"] == "leader_session"


async def test_kickoff_outbox_recovers_initializing_project(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import Project, ProjectEvent, ProjectRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Recoverable kickoff")
    project_id = project["id"]
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    planning_started_at = datetime.now(UTC)
    env.db.add_all(
        [
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                user_id=env.owner_id,
                sender_user_id=env.owner_id,
                role="user",
                content="Agree the plan before autonomous execution.",
                conversation_id=group["id"],
                message_meta={"kind": "project_group_message", "visible_to_group": True},
                created_at=planning_started_at,
            ),
            ChatMessage(
                agent_id=uuid.UUID(group["access_agent_id"]),
                sender_agent_id=env.leader_id,
                role="assistant",
                content="The traceable plan is ready for confirmation.",
                conversation_id=group["id"],
                message_meta={"kind": "project_subagent_reply", "visible_to_group": True},
                created_at=planning_started_at + timedelta(seconds=1),
            ),
        ]
    )
    await env.db.commit()
    real_dispatch = subagent_runtime.dispatch_project_run

    async def simulated_process_exit(_run_id):
        raise RuntimeError("simulated exit after kickoff outbox commit")

    monkeypatch.setattr(subagent_runtime, "dispatch_project_run", simulated_process_exit)
    first = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "Confirmed once"},
    )
    assert first.status_code == 202, first.text
    assert first.json()["status"] == "initializing"
    run_id = uuid.UUID(first.json()["run_id"])
    stored = await env.db.get(Project, uuid.UUID(project_id))
    pending = await env.db.get(ProjectRun, run_id)
    assert stored is not None and stored.status == "initializing"
    assert pending is not None and pending.input["dispatch"]["task"]

    monkeypatch.setattr(subagent_runtime, "dispatch_project_run", real_dispatch)
    recovered = await env.client.post(
        f"/api/projects/{project_id}/kickoff/confirm",
        json={"confirmation": "A retry must recover, not start twice"},
    )
    assert recovered.status_code == 202, recovered.text
    assert recovered.json()["status"] == "running"
    assert uuid.UUID(recovered.json()["run_id"]) == run_id
    env.db.expire_all()
    stored = await env.db.get(Project, uuid.UUID(project_id))
    events = (
        (
            await env.db.execute(
                select(ProjectEvent).where(
                    ProjectEvent.project_id == uuid.UUID(project_id),
                    ProjectEvent.event_type == "project.kickoff.confirmed",
                )
            )
        )
        .scalars()
        .all()
    )
    assert stored is not None and stored.status == "running"
    assert len(events) == 1


async def test_project_subagent_reply_materializes_without_resuming_group_root(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.audit import ChatMessage
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectRun
    from app.models.subagent_run import SubagentRun
    from app.services import subagent_runtime

    env = project_api
    live_events: list[tuple[str, str, dict]] = []

    async def capture_live(agent_id, conversation_id, payload):
        live_events.append((str(agent_id), str(conversation_id), payload))

    monkeypatch.setattr(
        "app.api.websocket.manager.send_to_session",
        capture_live,
    )
    project = await _create_project(env, name="Passive child reply")
    await _mark_project_running(env, project["id"])
    group = (await env.client.get(f"/api/projects/{project['id']}/group-session")).json()
    wake = await env.client.post(
        f"/api/projects/{project['id']}/group-sessions/{group['id']}/messages",
        json={"content": "Worker answer once", "mentions": [str(env.worker_id)]},
    )
    assert wake.status_code == 201, wake.text
    group_user_events = [
        payload
        for _agent_id, conversation_id, payload in live_events
        if conversation_id == group["id"]
        and payload.get("type") == "user_message_committed"
    ]
    assert len(group_user_events) == 1
    assert group_user_events[0]["content"] == "Worker answer once"
    assert group_user_events[0]["turn"]["phase"] == "active"
    worker_wake = next(row for row in wake.json()["subagent_runs"] if row["agent_id"] == str(env.worker_id))
    child_id = uuid.UUID(worker_wake["session_id"])
    project_run_id = uuid.UUID(worker_wake["project_run_id"])
    child_input = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["project_run_id"].as_string() == str(project_run_id),
            )
        )
    ).scalar_one()
    child_input_id = child_input.id
    child_input.message_meta = {
        **dict(child_input.message_meta or {}),
        "subagent_input_state": "processing",
        "subagent_turn_anchor_id": str(child_input.id),
        "turn_status": "running",
    }
    durable_run = await env.db.get(SubagentRun, child_id)
    assert durable_run is not None
    durable_run.status = "running"
    durable_run.lease_owner = subagent_runtime.settings.INSTANCE_ID
    await env.db.commit()

    terminal = await subagent_runtime._finish_subagent_turn(
        run_id=child_id,
        anchor_id=child_input_id,
        reply="Worker result",
        failed=False,
    )
    assert terminal is True
    env.db.expire_all()
    project_run = await env.db.get(ProjectRun, project_run_id)
    assert project_run is not None and project_run.status == "succeeded"
    assert project_run.finished_at is not None
    assert project_run.output["subagent_session_id"] == str(child_id)
    assert project_run.output["result"] == "Worker result"
    completion = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["kind"].as_string() == "subagent_completion",
            )
        )
    ).scalar_one()

    async def forbidden_resume(_anchor):
        raise AssertionError("project group completion must not resume root LLM")

    monkeypatch.setattr("app.services.turn_recovery.resume_turn", forbidden_resume)
    assert await subagent_runtime._dispatch_parent_event(completion.id) is True
    materialized = (
        await env.db.execute(
            select(ChatMessage).where(ChatMessage.external_event_key == f"project-subagent:{completion.id}")
        )
    ).scalar_one()
    assert materialized.conversation_id == group["id"]
    assert materialized.sender_agent_id == env.worker_id
    assert materialized.message_meta["visible_to_group"] is True
    assert materialized.message_meta["mentions"] == []
    assert materialized.message_meta["awakened_agent_ids"] == []
    assert materialized.message_meta["leader_batch_state"] == "pending"
    assert materialized.message_meta["default_leader_agent_id"] == str(env.leader_id)
    assert materialized.message_meta["timeline_anchor_id"] == wake.json()["message"]["id"]
    assert materialized.message_meta["producer_scope"] == (
        f"project:{child_id}:{child_input_id}"
    )
    committed_events = [
        payload
        for _agent_id, conversation_id, payload in live_events
        if conversation_id == group["id"]
        and payload.get("type") == "assistant_message_committed"
    ]
    assert len(committed_events) == 1
    committed_event = committed_events[0]
    assert committed_event["message_id"] == str(materialized.id)
    assert committed_event["transient_message_id"] == (
        f"project-stream:{child_id}:{child_input_id}"
    )
    assert committed_event["producer_scope"] == (
        f"project:{child_id}:{child_input_id}"
    )
    assert committed_event["timeline_anchor_id"] == wake.json()["message"]["id"]
    current_group = await env.client.get(
        f"/api/projects/{project['id']}/group-sessions/{group['id']}/messages"
    )
    assert current_group.status_code == 200
    assert committed_event["turn"]["status"] == current_group.json()["turn"]["status"]
    parent = await env.db.get(ChatSession, uuid.UUID(group["id"]))
    assert parent is not None and parent.source_channel == "project"


async def test_project_group_timeline_reuses_standard_child_message_contract(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from datetime import timedelta

    from app.models.audit import ChatMessage
    from app.models.project import ProjectRun
    from app.models.subagent_run import SubagentRun

    env = project_api
    live_events: list[tuple[str, str, dict]] = []

    async def capture_live(agent_id, conversation_id, payload):
        live_events.append((str(agent_id), str(conversation_id), payload))

    monkeypatch.setattr("app.api.websocket.manager.send_to_session", capture_live)
    project = await _create_project(env, name="Standard group timeline")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Show the full execution turn", "mentions": [str(env.worker_id)]},
    )
    assert wake.status_code == 201, wake.text
    worker_wake = next(row for row in wake.json()["subagent_runs"] if row["agent_id"] == str(env.worker_id))
    child_id = uuid.UUID(worker_wake["session_id"])
    project_run_id = uuid.UUID(worker_wake["project_run_id"])
    child_input = (
        await env.db.execute(
            select(ChatMessage).where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.message_meta["project_run_id"].as_string() == str(project_run_id),
            )
        )
    ).scalar_one()
    child_input.message_meta = {
        **dict(child_input.message_meta or {}),
        "subagent_input_state": "processing",
        "subagent_turn_anchor_id": str(child_input.id),
        "turn_status": "running",
    }
    base_time = child_input.created_at + timedelta(seconds=1)
    running_tool = ChatMessage(
        agent_id=env.worker_id,
        role="tool_call",
        content=json.dumps(
            {
                "name": "read_file",
                "call_id": "worker-read-1",
                "args": {"path": "README.md"},
                "status": "running",
                "result": "",
                "reasoning_content": "Inspecting the project evidence",
            }
        ),
        conversation_id=str(child_id),
        message_meta={"turn_anchor_id": str(child_input.id)},
        created_at=base_time,
    )
    done_tool = ChatMessage(
        agent_id=env.worker_id,
        role="tool_call",
        content=json.dumps(
            {
                "name": "read_file",
                "call_id": "worker-read-1",
                "args": {"path": "README.md"},
                "status": "done",
                "result": "# Project evidence",
                "reasoning_content": "Inspecting the project evidence",
            }
        ),
        conversation_id=str(child_id),
        message_meta={"turn_anchor_id": str(child_input.id)},
        created_at=base_time + timedelta(seconds=1),
    )
    confirmation = ChatMessage(
        agent_id=env.worker_id,
        role="tool_call",
        content=json.dumps(
            {
                "name": "request_confirmation",
                "args": {"title": "Approve delivery", "summary": "Publish the evidence"},
                "status": "pending",
                "result": "",
            }
        ),
        conversation_id=str(child_id),
        message_meta={"turn_anchor_id": str(child_input.id), "turn_status": "suspended"},
        created_at=base_time + timedelta(seconds=2),
    )
    fork_context = ChatMessage(
        agent_id=env.worker_id,
        role="assistant",
        content="Fork context must stay in the child session",
        conversation_id=str(child_id),
        message_meta={"kind": "subagent_fork_context"},
        created_at=base_time + timedelta(seconds=3),
    )
    child_final = ChatMessage(
        agent_id=env.worker_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content="Worker final answer",
        thinking="Checked the evidence before answering",
        conversation_id=str(child_id),
        message_meta={
            "kind": "subagent_completion",
            "turn_anchor_id": str(child_input.id),
            "project_run_ids": [str(project_run_id)],
            "attachments": [],
        },
        created_at=base_time + timedelta(seconds=4),
    )
    later_human = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Later Human message C",
        conversation_id=group["id"],
        message_meta={
            "kind": "project_group_message",
            "project_id": str(project_id),
            "visible_to_group": True,
        },
        created_at=base_time + timedelta(seconds=3, milliseconds=500),
    )
    env.db.add_all(
        [running_tool, done_tool, confirmation, fork_context, child_final, later_human]
    )
    await env.db.flush()
    materialized = ChatMessage(
        agent_id=env.leader_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content=child_final.content,
        conversation_id=group["id"],
        external_event_key=f"project-subagent:{child_final.id}",
        message_meta={
            "kind": "project_subagent_reply",
            "visible_to_group": True,
            "child_message_id": str(child_final.id),
            "subagent_id": str(child_id),
            "source_project_run_ids": [str(project_run_id)],
            "attachments": [],
        },
        created_at=base_time + timedelta(seconds=5),
    )
    env.db.add(materialized)
    await env.db.commit()

    paged_items: list[dict] = []
    before = None
    for _page in range(10):
        params = {"limit": 1}
        if before:
            params["before"] = before
        page = await env.client.get(
            f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
            params=params,
        )
        assert page.status_code == 200, page.text
        page_payload = page.json()
        paged_items.extend(page_payload["items"])
        if any(
            item.get("role") == "user"
            and item.get("content") == "Show the full execution turn"
            for item in page_payload["items"]
        ):
            break
        before = page_payload["next_cursor"]
        assert before is not None
    else:
        raise AssertionError("group anchor page was not reached")

    paged_finals = [
        item for item in paged_items if item.get("content") == child_final.content
    ]
    assert [item["id"] for item in paged_finals] == [str(materialized.id)]
    assert paged_finals[0]["turnAnchorId"] == str(
        wake.json()["message"]["id"]
    )
    paged_tools = [item for item in paged_items if item.get("role") == "tool_call"]
    assert paged_tools
    assert {
        item["turnAnchorId"] for item in paged_tools
    } == {str(wake.json()["message"]["id"])}

    response = await env.client.get(f"/api/projects/{project_id}/group-sessions/{group['id']}/messages")
    assert response.status_code == 200, response.text
    items = response.json()["items"]
    assert len([item for item in items if item["role"] == "user"]) == 2
    assert all(item["content"] != fork_context.content for item in items)

    tools = [item for item in items if item["role"] == "tool_call"]
    read_tools = [item for item in tools if item.get("toolName") == "read_file"]
    assert len(read_tools) == 1
    assert read_tools[0]["toolCallId"] == "worker-read-1"
    assert read_tools[0]["toolStatus"] == "done"
    assert read_tools[0]["toolResult"] == "# Project evidence"
    assert read_tools[0]["content"] == ""
    assert read_tools[0]["display_content"] == ""
    assert read_tools[0]["sender_agent_id"] == str(env.worker_id)
    group_anchor_id = next(
        item["id"]
        for item in items
        if item["role"] == "user" and item["content"] == "Show the full execution turn"
    )
    expected_producer_scope = f"project:{child_id}:{child_input.id}"
    assert read_tools[0]["turnAnchorId"] == group_anchor_id
    assert read_tools[0]["producerScope"] == expected_producer_scope
    assert read_tools[0]["message_meta"]["producer_scope"] == expected_producer_scope

    confirmation_item = next(item for item in tools if item.get("toolName") == "request_confirmation")
    assert confirmation_item["toolCallId"] == str(confirmation.id)
    assert confirmation_item["toolStatus"] == "pending"
    assert confirmation_item["sender_agent_id"] == str(env.worker_id)

    confirmation_payload = json.loads(confirmation.content)
    confirmation_payload.update({"status": "done", "result": "approved"})
    confirmation.content = json.dumps(confirmation_payload)
    durable_child = await env.db.get(SubagentRun, child_id)
    assert durable_child is not None
    durable_child.status = "running"
    project_run = await env.db.get(ProjectRun, project_run_id)
    assert project_run is not None
    project_run.status = "queued"
    await env.db.commit()
    from app.services.subagent_runtime import resume_subagent_after_confirmation

    assert await resume_subagent_after_confirmation(
        child_id,
        resolved_tool_payload={
            "type": "tool_call",
            "name": "request_confirmation",
            "call_id": str(confirmation.id),
            "args": confirmation_payload["args"],
            "status": "done",
            "result": "approved",
        },
        child_turn_anchor_id=child_input.id,
    ) is True
    group_resolved = [
        payload
        for _agent_id, conversation_id, payload in live_events
        if conversation_id == group["id"]
        and payload.get("type") == "tool_call"
        and payload.get("call_id") == str(confirmation.id)
    ]
    assert len(group_resolved) == 1
    assert group_resolved[0]["status"] == "done"
    assert group_resolved[0]["timeline_anchor_id"] == group_anchor_id
    assert group_resolved[0]["producer_scope"] == expected_producer_scope
    assert group_resolved[0]["turn"]["phase"] == "active"

    final_items = [item for item in items if item["content"] == child_final.content]
    assert len(final_items) == 1
    assert final_items[0]["id"] == str(materialized.id)
    assert final_items[0]["thinking"] == child_final.thinking
    assert final_items[0]["sender_agent_id"] == str(env.worker_id)
    assert final_items[0]["turnAnchorId"] == group_anchor_id
    assert final_items[0]["producerScope"] == expected_producer_scope
    assert final_items[0]["canonicalDone"] is True
    item_ids = [item["id"] for item in items]
    later_human_index = item_ids.index(str(later_human.id))
    assert item_ids.index(str(materialized.id)) < later_human_index
    assert item_ids.index(str(confirmation.id)) < later_human_index


async def test_project_group_history_uses_shared_tool_projection(
    project_api: ProjectApiEnv,
):
    env = project_api
    project = await _create_project(env, name="Shared tool projection")
    group = (await env.client.get(f"/api/projects/{project['id']}/group-session")).json()
    anchor = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Run the batch",
        conversation_id=group["id"],
        message_meta={"kind": "project_group_message", "visible_to_group": True},
    )
    env.db.add(anchor)
    await env.db.flush()
    tool_base = {
        "name": "toolscall",
        "call_id": "project-batch-1",
        "args": {"table_id": 18, "rows": [{"name": "A"}]},
        "result": "",
    }
    running = ChatMessage(
        agent_id=env.leader_id,
        sender_agent_id=env.leader_id,
        role="tool_call",
        content=json.dumps({**tool_base, "status": "running"}),
        conversation_id=group["id"],
        message_meta={"turn_anchor_id": str(anchor.id)},
    )
    done = ChatMessage(
        agent_id=env.leader_id,
        sender_agent_id=env.leader_id,
        role="tool_call",
        content=json.dumps({**tool_base, "status": "done", "result": "created"}),
        conversation_id=group["id"],
        message_meta={"turn_anchor_id": str(anchor.id)},
    )
    env.db.add_all([running, done])
    await env.db.commit()

    response = await env.client.get(
        f"/api/projects/{project['id']}/group-sessions/{group['id']}/messages"
    )
    assert response.status_code == 200, response.text
    tools = [item for item in response.json()["items"] if item["role"] == "tool_call"]
    assert len(tools) == 1
    assert tools[0]["toolCallId"] == "project-batch-1"
    assert tools[0]["toolStatus"] == "done"
    assert tools[0]["toolResult"] == "created"
    assert tools[0]["toolArgs"] == {"table_id": 18, "rows": [{"name": "A"}]}
    assert tools[0]["content"] == ""
    assert tools[0]["display_content"] == ""


async def test_project_group_timeline_page_does_not_materialize_unrelated_history(
    project_api: ProjectApiEnv,
):
    from sqlalchemy import event as sqlalchemy_event

    from app.models.audit import ChatMessage
    from app.models.project import ProjectRun
    from app.services.project_group_timeline import build_project_group_timeline

    env = project_api
    project = await _create_project(env, name="Bounded group timeline")
    project_id = uuid.UUID(project["id"])
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    child_id = uuid.uuid4()
    page_message = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Current page anchor",
        conversation_id=group["id"],
        message_meta={"kind": "project_group_message", "visible_to_group": True},
    )
    old_message = ChatMessage(
        agent_id=env.leader_id,
        user_id=env.owner_id,
        sender_user_id=env.owner_id,
        role="user",
        content="Unrelated old anchor",
        conversation_id=group["id"],
        message_meta={"kind": "project_group_message", "visible_to_group": True},
    )
    env.db.add_all([page_message, old_message])
    await env.db.flush()
    relevant_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.worker_id,
        initiated_by_user_id=env.owner_id,
        execution_user_id=env.owner_id,
        status="running",
        trigger_type="group_mention",
        input={"group_message_id": str(page_message.id)},
        output={"subagent_session_id": str(child_id)},
    )
    unrelated_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=project_id,
        agent_id=env.worker_id,
        initiated_by_user_id=env.owner_id,
        execution_user_id=env.owner_id,
        status="succeeded",
        trigger_type="group_mention",
        input={"group_message_id": str(old_message.id)},
        output={"subagent_session_id": str(child_id)},
    )
    env.db.add_all([relevant_run, unrelated_run])
    await env.db.flush()
    relevant_input_id = uuid.uuid4()
    unrelated_input_id = uuid.uuid4()
    relevant_input = ChatMessage(
        id=relevant_input_id,
        agent_id=env.worker_id,
        user_id=env.owner_id,
        role="user",
        content="Relevant child input",
        conversation_id=str(child_id),
        message_meta={
            "kind": "subagent_input",
            "project_run_id": str(relevant_run.id),
            "subagent_turn_anchor_id": str(relevant_input_id),
        },
    )
    unrelated_input = ChatMessage(
        id=unrelated_input_id,
        agent_id=env.worker_id,
        user_id=env.owner_id,
        role="user",
        content="Unrelated child input",
        conversation_id=str(child_id),
        message_meta={
            "kind": "subagent_input",
            "project_run_id": str(unrelated_run.id),
            "subagent_turn_anchor_id": str(unrelated_input_id),
        },
    )
    relevant_reply = ChatMessage(
        agent_id=env.worker_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content="Relevant child reply",
        conversation_id=str(child_id),
        message_meta={"turn_anchor_id": str(relevant_input_id)},
    )
    unrelated_reply = ChatMessage(
        agent_id=env.worker_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content="Unrelated child reply",
        conversation_id=str(child_id),
        message_meta={"turn_anchor_id": str(unrelated_input_id)},
    )
    env.db.add_all(
        [relevant_input, unrelated_input, relevant_reply, unrelated_reply]
    )
    await env.db.commit()

    page_message_id = page_message.id
    relevant_run_id = relevant_run.id
    unrelated_run_id = unrelated_run.id
    relevant_reply_id = relevant_reply.id
    unrelated_reply_id = unrelated_reply.id
    env.db.expunge_all()
    page_row = await env.db.get(ChatMessage, page_message_id)
    loaded_run_ids: set[uuid.UUID] = set()
    loaded_message_ids: set[uuid.UUID] = set()

    def capture_loaded(_session, instance):
        if isinstance(instance, ProjectRun):
            loaded_run_ids.add(instance.id)
        elif isinstance(instance, ChatMessage):
            loaded_message_ids.add(instance.id)

    sqlalchemy_event.listen(env.db.sync_session, "loaded_as_persistent", capture_loaded)
    try:
        timeline = await build_project_group_timeline(
            env.db,
            project_id=project_id,
            group_messages=[page_row],
        )
    finally:
        sqlalchemy_event.remove(
            env.db.sync_session,
            "loaded_as_persistent",
            capture_loaded,
        )

    assert relevant_run_id in loaded_run_ids
    assert unrelated_run_id not in loaded_run_ids
    assert relevant_input_id in loaded_message_ids
    assert unrelated_input_id not in loaded_message_ids
    assert relevant_reply_id in loaded_message_ids
    assert unrelated_reply_id not in loaded_message_ids
    assert any(item["content"] == "Relevant child reply" for item in timeline)
    assert all(item["content"] != "Unrelated child reply" for item in timeline)


async def test_project_participant_replies_coalesce_into_one_durable_leader_turn(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.audit import ChatMessage
    from app.models.project import ProjectEvent, ProjectRun, ProjectWorkItem
    from app.models.subagent_run import SubagentRun
    from app.services import subagent_runtime

    env = project_api
    project = await _create_project(env, name="Coalesced Leader inbox")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    dependency = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.reviewer_id,
        created_by_agent_id=env.leader_id,
        title="Confirm the upstream evidence source",
        description="The source contract must be settled first",
        status="done",
        priority="medium",
        acceptance_criteria=["Source contract is recorded"],
        dependency_ids=[],
    )
    env.db.add(dependency)
    await env.db.flush()
    dependency_id = dependency.id
    work_item = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.worker_id,
        created_by_agent_id=env.leader_id,
        title="Collect exact participant evidence",
        description="All replies in this batch belong to one explicit work item",
        status="todo",
        priority="medium",
        acceptance_criteria=["Both participant replies are consolidated with attributed evidence"],
        dependency_ids=[str(dependency_id)],
    )
    env.db.add(work_item)
    await env.db.flush()
    env.db.add(
        ProjectEvent(
            tenant_id=env.tenant_id,
            project_id=project_id,
            work_item_id=work_item.id,
            actor_agent_id=env.worker_id,
            event_type="work_item.updated",
            summary="Participant evidence recorded",
            event_metadata={"evidence": ["docs/participant-evidence.md", "commit-a1b2c3d4"]},
        )
    )
    await env.db.commit()
    work_item_id = work_item.id
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={
            "content": "Ask two participants and let Leader coordinate results",
            "mentions": [str(env.worker_id), str(env.reviewer_id)],
            "work_item_id": str(work_item_id),
        },
    )
    assert wake.status_code == 201, wake.text
    run_by_agent = {row["agent_id"]: row for row in wake.json()["subagent_runs"]}
    created_runs = (
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.id.in_([uuid.UUID(row["project_run_id"]) for row in run_by_agent.values()])
                )
            )
        )
        .scalars()
        .all()
    )
    assert created_runs
    assert {row.work_item_id for row in created_runs} == {work_item_id}
    assert run_by_agent[str(env.leader_id)]["run_id"] is None
    assert run_by_agent[str(env.leader_id)]["status"] == "skipped"
    worker_child_id = uuid.UUID(run_by_agent[str(env.worker_id)]["session_id"])
    reviewer_child_id = uuid.UUID(run_by_agent[str(env.reviewer_id)]["session_id"])

    completion_rows = [
        ChatMessage(
            agent_id=agent_id,
            sender_agent_id=agent_id,
            role="assistant",
            content=content,
            conversation_id=str(child_id),
            message_meta={
                "kind": "subagent_completion",
                "subagent_wake": True,
                "project_run_ids": [project_run_id],
                "attachments": attachments,
            },
        )
        for agent_id, child_id, content, project_run_id, attachments in [
            (
                env.worker_id,
                worker_child_id,
                "Worker evidence A",
                run_by_agent[str(env.worker_id)]["project_run_id"],
                [{"name": "evidence-a.md", "path": "evidence-a.md"}],
            ),
            (
                env.reviewer_id,
                reviewer_child_id,
                "Reviewer finding B",
                run_by_agent[str(env.reviewer_id)]["project_run_id"],
                [],
            ),
            (
                env.worker_id,
                worker_child_id,
                "Worker follow-up C",
                run_by_agent[str(env.worker_id)]["project_run_id"],
                [],
            ),
        ]
    ]
    env.db.add_all(completion_rows)
    await env.db.commit()

    async def forbidden_resume(_anchor):
        raise AssertionError("participant replies must not resume the project group root")

    monkeypatch.setattr("app.services.turn_recovery.resume_turn", forbidden_resume)
    for completion in completion_rows:
        assert await subagent_runtime._dispatch_parent_event(completion.id) is True

    assert (
        await subagent_runtime._dispatch_project_leader_batch(
            uuid.UUID(group["id"]),
            debounce_seconds=0,
        )
        is True
    )
    env.db.expire_all()
    leader_child_id = (
        await env.db.execute(
            select(ChatSession.id)
            .where(
                ChatSession.project_id == project_id,
                ChatSession.agent_id == env.leader_id,
                ChatSession.source_channel == "subagent",
            )
            .order_by(ChatSession.created_at, ChatSession.id)
        )
    ).scalar_one()
    leader_inputs = (
        (
            await env.db.execute(
                select(ChatMessage).where(
                    ChatMessage.conversation_id == str(leader_child_id),
                    ChatMessage.message_meta["project_leader_batch"].as_boolean().is_(True),
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(leader_inputs) == 1
    batch_input = leader_inputs[0]
    assert len(batch_input.message_meta["source_group_message_ids"]) == 3
    assert len(batch_input.message_meta["source_replies"]) == 3
    assert batch_input.message_meta["batch_limits"] == {
        "max_replies": subagent_runtime.PROJECT_LEADER_BATCH_MAX_REPLIES,
        "max_bytes": subagent_runtime.PROJECT_LEADER_BATCH_MAX_BYTES,
    }
    assert batch_input.message_meta["batch_input_bytes"] <= subagent_runtime.PROJECT_LEADER_BATCH_MAX_BYTES
    assert batch_input.message_meta["original_human_request"]["content"] == (
        "Ask two participants and let Leader coordinate results"
    )
    [work_item_snapshot] = batch_input.message_meta["work_item_snapshots"]
    assert work_item_snapshot["title"] == "Collect exact participant evidence"
    assert work_item_snapshot["acceptance_criteria"] == [
        "Both participant replies are consolidated with attributed evidence"
    ]
    assert work_item_snapshot["dependencies"] == [
        {
            "id": str(dependency_id),
            "title": "Confirm the upstream evidence source",
            "status": "done",
        }
    ]
    assert work_item_snapshot["evidence"] == ["docs/participant-evidence.md", "commit-a1b2c3d4"]
    assert "Ask two participants and let Leader coordinate results" in batch_input.content
    assert "Both participant replies are consolidated with attributed evidence" in batch_input.content
    assert "docs/participant-evidence.md" in batch_input.content
    assert {
        (reply["source_agent_name"], reply["source_role_snapshot"])
        for reply in batch_input.message_meta["source_replies"]
    } == {
        ("Worker", "Build deliverables"),
        ("Reviewer", "Review evidence"),
    }
    batch_attachments = [
        attachment for reply in batch_input.message_meta["source_replies"] for attachment in reply["attachments"]
    ]
    assert batch_attachments[0]["name"] == "evidence-a.md"

    materialized = (
        (
            await env.db.execute(
                select(ChatMessage).where(
                    ChatMessage.conversation_id == group["id"],
                    ChatMessage.message_meta["kind"].as_string() == "project_subagent_reply",
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(materialized) == 3
    assert {row.message_meta["leader_batch_state"] for row in materialized} == {"delivered"}
    assert len({row.message_meta["leader_batch_id"] for row in materialized}) == 1

    batch_runs = (
        (
            await env.db.execute(
                select(ProjectRun).where(
                    ProjectRun.project_id == project_id,
                    ProjectRun.trigger_type == "leader_reply_batch",
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(batch_runs) == 1
    assert batch_runs[0].work_item_id == work_item_id
    assert batch_runs[0].input["group_message_id"] == wake.json()["message"]["id"]
    assert batch_runs[0].input["related_work_item_ids"] == [str(work_item_id)]
    assert batch_runs[0].input["original_human_request"] == batch_input.message_meta["original_human_request"]
    assert batch_runs[0].input["work_item_snapshots"] == batch_input.message_meta["work_item_snapshots"]
    assert batch_runs[0].output["subagent_session_id"] == str(leader_child_id)
    assert batch_runs[0].output["source_count"] == 3
    leader_run = await env.db.get(SubagentRun, leader_child_id)
    assert leader_run is not None and leader_run.parent_session_id == uuid.UUID(group["id"])

    assert (
        await subagent_runtime._dispatch_project_leader_batch(
            uuid.UUID(group["id"]),
            debounce_seconds=0,
        )
        is False
    )
    env.db.expire_all()
    replay_inputs = (
        (
            await env.db.execute(
                select(ChatMessage).where(
                    ChatMessage.conversation_id == str(leader_child_id),
                    ChatMessage.message_meta["project_leader_batch"].as_boolean().is_(True),
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(replay_inputs) == 1


async def test_leader_reply_batch_lineage_requires_exact_source_run_consensus(
    project_api: ProjectApiEnv,
):
    """Concurrent work-item replies are never guessed into one batch lineage."""
    from app.models.audit import ChatMessage
    from app.models.project import ProjectRun, ProjectWorkItem
    from app.services.subagent_runtime import _resolve_batch_work_item_lineage

    env = project_api
    project = await _create_project(env, name="Exact batch lineage")
    project_id = uuid.UUID(project["id"])
    first_item = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.worker_id,
        created_by_agent_id=env.leader_id,
        title="Concurrent item A",
        status="in_progress",
        priority="medium",
        acceptance_criteria=[],
        dependency_ids=[],
    )
    second_item = ProjectWorkItem(
        tenant_id=env.tenant_id,
        project_id=project_id,
        assignee_agent_id=env.worker_id,
        created_by_agent_id=env.leader_id,
        title="Concurrent item B",
        status="in_progress",
        priority="medium",
        acceptance_criteria=[],
        dependency_ids=[],
    )
    env.db.add_all([first_item, second_item])
    await env.db.flush()
    runs = [
        ProjectRun(
            tenant_id=env.tenant_id,
            project_id=project_id,
            work_item_id=work_item_id,
            agent_id=env.worker_id,
            initiated_by_user_id=env.owner_id,
            status="succeeded",
            trigger_type="a2a",
            input={},
            output={},
        )
        for work_item_id in [first_item.id, first_item.id, second_item.id]
    ]
    env.db.add_all(runs)
    await env.db.flush()
    rows = [
        ChatMessage(
            agent_id=env.worker_id,
            sender_agent_id=env.worker_id,
            role="assistant",
            content=f"reply-{index}",
            conversation_id=str(uuid.uuid4()),
            message_meta={"source_project_run_ids": [str(run.id)]},
        )
        for index, run in enumerate(runs)
    ]
    missing = ChatMessage(
        agent_id=env.worker_id,
        sender_agent_id=env.worker_id,
        role="assistant",
        content="reply-without-source",
        conversation_id=str(uuid.uuid4()),
        message_meta={},
    )

    exact, related = await _resolve_batch_work_item_lineage(env.db, project_id, rows[:2])
    assert exact == first_item.id
    assert related == [first_item.id]

    exact, related = await _resolve_batch_work_item_lineage(env.db, project_id, [rows[0], rows[2]])
    assert exact is None
    assert set(related) == {first_item.id, second_item.id}

    exact, related = await _resolve_batch_work_item_lineage(env.db, project_id, [rows[0], missing])
    assert exact is None
    assert related == [first_item.id]


async def test_peer_a2a_completion_only_escalates_decisions_failures_and_owner_requests() -> None:
    from app.services.subagent_runtime import _a2a_completion_leader_policy

    leader_id = uuid.uuid4()
    peer_id = uuid.uuid4()
    informational = SimpleNamespace(
        input={"dispatch": {"a2a": {"from_agent_id": str(peer_id), "mode": "task_delegate"}}},
        work_item_id=None,
    )
    assert await _a2a_completion_leader_policy(
        None,
        project_run=informational,
        leader_agent_id=leader_id,
        failed=False,
    ) == (False, "peer_completion_recorded")

    consult = SimpleNamespace(
        input={"dispatch": {"a2a": {"from_agent_id": str(peer_id), "mode": "consult"}}},
        work_item_id=None,
    )
    assert await _a2a_completion_leader_policy(
        None,
        project_run=consult,
        leader_agent_id=leader_id,
        failed=False,
    ) == (True, "decision_consult")

    owner_requested = SimpleNamespace(
        input={"dispatch": {"a2a": {"from_agent_id": str(leader_id), "mode": "task_delegate"}}},
        work_item_id=None,
    )
    assert (
        await _a2a_completion_leader_policy(
            None,
            project_run=owner_requested,
            leader_agent_id=leader_id,
            failed=False,
        )
    )[0] is True
    assert (
        await _a2a_completion_leader_policy(
            None,
            project_run=informational,
            leader_agent_id=leader_id,
            failed=True,
        )
    )[0] is True


async def test_leader_reply_batch_is_bounded_and_causally_isolated() -> None:
    from app.services.subagent_runtime import (
        PROJECT_LEADER_BATCH_MAX_BYTES,
        PROJECT_LEADER_BATCH_MAX_REPLIES,
        _select_leader_batch_rows,
        _truncate_batch_content,
    )

    first_cause = [
        ChatMessage(id=uuid.uuid4(), content="x" * 100, role="assistant", conversation_id="group")
        for _ in range(PROJECT_LEADER_BATCH_MAX_REPLIES + 2)
    ]
    unrelated = ChatMessage(
        id=uuid.uuid4(),
        content="other",
        role="assistant",
        conversation_id="group",
    )
    rows = [first_cause[0], unrelated, *first_cause[1:]]
    causal_keys = {row.id: "work-item:a" for row in first_cause}
    causal_keys[unrelated.id] = "work-item:b"
    selected = _select_leader_batch_rows(rows, causal_keys)
    assert len(selected) == PROJECT_LEADER_BATCH_MAX_REPLIES
    assert unrelated not in selected

    large_rows = [
        ChatMessage(id=uuid.uuid4(), content="字" * 4_000, role="assistant", conversation_id="group") for _ in range(3)
    ]
    selected_large = _select_leader_batch_rows(
        large_rows,
        {row.id: "cause:one" for row in large_rows},
    )
    assert len(selected_large) == 1
    truncated = _truncate_batch_content("字" * PROJECT_LEADER_BATCH_MAX_BYTES)
    assert len(truncated.encode("utf-8")) <= PROJECT_LEADER_BATCH_MAX_BYTES


async def test_project_standard_file_tools_use_shared_git_without_project_specific_size_limits(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.services.agent_runtime_workspace import (
        bind_agent_runtime_workspace,
        project_agent_runtime_workspace,
    )
    from app.services.agent_tools import execute_tool

    env = project_api
    project = await _create_project(env, name="Standard project files")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Prepare project files", "mentions": [str(env.worker_id)]},
    )
    assert wake.status_code == 201, wake.text
    child_id = next(
        row["session_id"]
        for row in wake.json()["subagent_runs"]
        if row["agent_id"] == str(env.worker_id)
    )
    runtime_workspace = project_agent_runtime_workspace(
        agent_id=env.worker_id,
        tenant_id=env.tenant_id,
        project_id=uuid.UUID(project_id),
    )

    with bind_agent_runtime_workspace(runtime_workspace):
        missing_workspace = await execute_tool(
            "list_files",
            {"path": ""},
            env.worker_id,
            env.owner_id,
            session_id=child_id,
            tool_call_id="project-file-missing-workspace",
            skip_autonomy=True,
        )
    assert "require workspace='agent' or workspace='project'" in missing_workspace

    async def run(tool_name: str, arguments: dict[str, Any], *, workspace: str = "project") -> str:
        with bind_agent_runtime_workspace(runtime_workspace):
            return await execute_tool(
                tool_name,
                {"workspace": workspace, **arguments},
                env.worker_id,
                env.owner_id,
                session_id=child_id,
                tool_call_id=f"project-file-{tool_name}-{uuid.uuid4().hex[:8]}",
                skip_autonomy=True,
            )

    repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    head_before_workspace_reads = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()

    async def deny_project_write(*_args, **_kwargs):
        return {"allowed": False, "level": "L2", "message": "project write denied"}

    from app.services.autonomy_service import autonomy_service

    monkeypatch.setattr(autonomy_service, "check_and_enforce", deny_project_write)
    with bind_agent_runtime_workspace(runtime_workspace):
        denied_by_autonomy = await execute_tool(
            "write_file",
            {"workspace": "project", "path": "blocked.txt", "content": "must not exist"},
            env.worker_id,
            env.owner_id,
            session_id=child_id,
            tool_call_id="project-file-autonomy-denied",
        )
    assert "project write denied" in denied_by_autonomy
    assert not (repo / "blocked.txt").exists()

    private_listing = await run("list_files", {"path": ""}, workspace="agent")
    assert "PROJECT.json" not in private_listing
    project_root = await run("list_files", {"path": ""})
    assert "PROJECT.json" in project_root
    assert subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip() == head_before_workspace_reads

    large_content = "project file content\n" + ("x" * (1024 * 1024 + 1))
    written = json.loads(
        await run("write_file", {"path": "docs/result.txt", "content": large_content})
    )
    assert written["operation"] == "write_file"
    assert written["path"] == "docs/result.txt"

    listed = await run("list_files", {"path": "docs"})
    assert "docs/result.txt" in listed
    read = await run("read_file", {"path": "docs/result.txt", "offset": 0, "limit": 1})
    assert "project file content" in read
    searched = await run(
        "search_files",
        {"pattern": "project file content", "path": "docs", "file_pattern": "*.txt"},
    )
    assert "docs/result.txt:1" in searched
    found = await run("find_files", {"pattern": "**/*.txt", "path": "."})
    assert "docs/result.txt" in found

    edited = json.loads(
        await run(
            "edit_file",
            {
                "path": "docs/result.txt",
                "old_string": "project file content",
                "new_string": "project file updated",
            },
        )
    )
    assert edited["replacements"] == 1
    moved = json.loads(
        await run(
            "move_file",
            {"source_path": "docs/result.txt", "destination_path": "deliverables/result.txt"},
        )
    )
    assert moved["destination_path"] == "deliverables/result.txt"
    deleted = json.loads(await run("delete_file", {"path": "deliverables"}))
    assert deleted["operation"] == "delete_file"

    denied = await run("write_file", {"path": ".agents/forbidden.txt", "content": "blocked"})
    assert "project deliverables" in denied

    stored = await env.db.get(Project, uuid.UUID(project_id))
    assert stored.settings["git"]["head"] == subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    committed_events = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.project_id == uuid.UUID(project_id),
                ProjectEvent.event_type == "project.file.committed",
            )
        )
    ).scalars().all()
    assert len(committed_events) == 4

    stored.status = "paused"
    await env.db.commit()
    paused_write = await run("write_file", {"path": "paused.txt", "content": "blocked"})
    assert "only while the project is running" in paused_write
    assert not (repo / "paused.txt").exists()


async def test_approved_project_tool_resumes_with_the_original_project_workspace(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    import app.database as database
    from app.services import agent_tools
    from app.services.agent_runtime_workspace import current_agent_runtime_workspace
    from app.services.autonomy_service import autonomy_service

    env = project_api
    project = await _create_project(env, name="Approved project workspace")
    project_id = uuid.UUID(project["id"])
    session = ChatSession(
        project_id=project_id,
        agent_id=env.worker_id,
        user_id=env.owner_id,
        title="Approved project action",
        source_channel="project",
        external_conv_id=f"approval-{uuid.uuid4()}",
    )
    env.db.add(session)
    await env.db.commit()

    observed: dict[str, Any] = {}

    async def observe_execute_tool(
        _tool_name: str,
        _arguments: dict[str, Any],
        agent_id: uuid.UUID,
        **_kwargs: Any,
    ) -> str:
        workspace = current_agent_runtime_workspace(agent_id)
        observed["agent_id"] = workspace.agent_id
        observed["project_id"] = workspace.project_id
        observed["is_project"] = workspace.is_project
        return "approved project tool executed"

    monkeypatch.setattr(database, "async_session", env.session_factory)
    monkeypatch.setattr(agent_tools, "execute_tool", observe_execute_tool)
    result = await autonomy_service._execute_approved_action(
        env.worker_id,
        "write_workspace_files",
        {
            "tool": "write_file",
            "args": {"workspace": "project", "path": "approved.txt", "content": "ok"},
            "requested_by": str(env.owner_id),
            "session_id": str(session.id),
            "tool_call_id": "approved-project-write",
        },
        env.owner_id,
    )

    assert result == "approved project tool executed"
    assert observed == {
        "agent_id": env.worker_id,
        "project_id": project_id,
        "is_project": True,
    }
    assert await autonomy_service._execute_approved_action(
        env.worker_id,
        "write_workspace_files",
        {
            "tool": "write_file",
            "args": {"workspace": "project", "path": "invalid.txt", "content": "blocked"},
            "session_id": str(session.id),
        },
        env.owner_id,
    ) == "Execution failed: approval has no valid original requester"


async def test_project_structured_readers_and_isolated_sandbox_share_the_project_repository(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from docx import Document

    from app.services import agent_tools
    from app.services.agent_runtime_workspace import (
        bind_agent_runtime_workspace,
        project_agent_runtime_workspace,
    )
    from app.services.tools import read_image as read_image_tool

    env = project_api
    project = await _create_project(env, name="Structured project workspace")
    project_id = uuid.UUID(project["id"])
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()
    wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Inspect and update project files", "mentions": [str(env.worker_id)]},
    )
    assert wake.status_code == 201, wake.text
    child_id = next(
        row["session_id"]
        for row in wake.json()["subagent_runs"]
        if row["agent_id"] == str(env.worker_id)
    )
    runtime_workspace = project_agent_runtime_workspace(
        agent_id=env.worker_id,
        tenant_id=env.tenant_id,
        project_id=project_id,
    )
    repo = project_repo_path(env.tenant_id, project_id)
    (repo / "docs").mkdir(exist_ok=True)
    document = Document()
    document.add_paragraph("Project document content")
    document.save(repo / "docs" / "brief.docx")
    (repo / "assets").mkdir(exist_ok=True)
    (repo / "assets" / "diagram.png").write_bytes(b"project-image-marker")
    (repo / "sandbox-delete.txt").write_text("remove me\n", encoding="utf-8")
    subprocess.run(
        [
            "git",
            "add",
            "--",
            "docs/brief.docx",
            "assets/diagram.png",
            "sandbox-delete.txt",
        ],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    )
    subprocess.run(
        [
            "git",
            "-c",
            "user.name=Project test",
            "-c",
            "user.email=project@test.invalid",
            "commit",
            "-m",
            "Add structured project inputs",
        ],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    )

    async def run(tool_name: str, arguments: dict[str, Any]) -> str:
        with bind_agent_runtime_workspace(runtime_workspace):
            return await agent_tools.execute_tool(
                tool_name,
                {"workspace": "project", **arguments},
                env.worker_id,
                env.owner_id,
                session_id=child_id,
                tool_call_id=f"project-{tool_name}-{uuid.uuid4().hex[:8]}",
                skip_autonomy=True,
            )

    document_result = await run("read_document", {"path": "docs/brief.docx"})
    assert "Project document content" in document_result

    async def fake_read_image(agent_id, arguments, *, workspace_root=None):
        assert agent_id == env.worker_id
        assert workspace_root is not None
        assert (workspace_root / "assets" / "diagram.png").read_bytes() == b"project-image-marker"
        assert not (workspace_root / ".agents").exists()
        assert not (workspace_root / ".git").exists()
        return "project image read"

    monkeypatch.setattr(read_image_tool, "handle_read_image", fake_read_image)
    assert await run("read_image", {"image_paths": ["assets/diagram.png"]}) == "project image read"

    async def fake_execute_code(_agent_id, _ws, _arguments, **kwargs):
        sandbox_root = kwargs["work_dir_override"]
        runtime_temp_root = kwargs["runtime_temp_path_override"]
        assert runtime_temp_root.parent == sandbox_root.parent
        assert runtime_temp_root != sandbox_root / ".tmp"
        runtime_temp_root.mkdir(parents=True, exist_ok=True)
        (runtime_temp_root / "runtime-only.txt").write_text("not a project file")
        assert (sandbox_root / "PROJECT.json").is_file()
        assert not (sandbox_root / ".agents").exists()
        assert not (sandbox_root / ".git").exists()
        if _arguments.get("code") == "create oversized output":
            with (sandbox_root / "oversized.bin").open("wb") as oversized:
                oversized.truncate(10 * 1024 * 1024 + 1)
            return "oversized output created"
        (sandbox_root / "sandbox").mkdir(exist_ok=True)
        (sandbox_root / "sandbox" / "result.txt").write_text(
            "sandbox project output\n",
            encoding="utf-8",
        )
        (sandbox_root / "sandbox-delete.txt").unlink()
        return "sandbox ok"

    monkeypatch.setattr(agent_tools, "_execute_code", fake_execute_code)
    sandbox_result = await run(
        "execute_code",
        {
            "action": "execute",
            "language": "python",
            "code": "print('ok')",
            "execution_mode": "foreground",
        },
    )
    assert "sandbox ok" in sandbox_result
    assert "sandbox_commit" in sandbox_result
    assert (repo / "sandbox" / "result.txt").read_text(encoding="utf-8") == "sandbox project output\n"
    assert not (repo / "sandbox-delete.txt").exists()
    assert not (repo / ".tmp").exists()
    assert not (repo / ".runtime-tmp").exists()
    assert subprocess.run(
        ["git", "status", "--porcelain"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout == ""
    sandbox_parent = Path(tempfile.gettempdir()) / "clawith-project-sandboxes"
    assert not list(sandbox_parent.glob("project-code-*"))
    assert not list(sandbox_parent.glob("project-read-*"))

    oversized_result = await run(
        "execute_code",
        {
            "action": "execute",
            "language": "python",
            "code": "create oversized output",
            "execution_mode": "foreground",
        },
    )
    assert "项目文件提交未完成" in oversized_result
    assert not (repo / "oversized.bin").exists()

    stored = await env.db.get(Project, project_id)
    assert stored.settings["git"]["head"] == subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=repo,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    sandbox_events = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.project_id == project_id,
                ProjectEvent.event_type == "project.file.committed",
                ProjectEvent.event_metadata["operation"].astext == "sandbox_commit",
            )
        )
    ).scalars().all()
    assert len(sandbox_events) == 1


async def test_project_uncertain_audit_resolution_preserves_only_reachable_commits(
    project_api: ProjectApiEnv,
):
    from app.services.project_git_service import (
        repository_state,
        write_project_workspace_file,
    )
    from app.services.project_runtime_tools import _resolve_uncertain_project_file_audit
    from app.services.project_service import add_event

    env = project_api
    created = await _create_project(env, name="Uncertain file audit")
    project = await env.db.get(Project, uuid.UUID(created["id"]))
    member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(ProjectMemberSnapshot.project_id == project.id)
        )
    ).scalars().first()
    assert project is not None and member is not None

    initial_head = (await repository_state(project))["head"]
    exact = await write_project_workspace_file(project, "exact.txt", "exact\n")
    assert await _resolve_uncertain_project_file_audit(
        project,
        previous_head=initial_head,
        result=exact,
        tool_call_id="exact-compensation",
        member=member,
        agent_id=member.agent_id,
        project_run=None,
        summary="exact compensation",
        metadata={**exact, "tool_call_id": "exact-compensation"},
    ) is False
    assert (await repository_state(project))["head"] == initial_head

    committed = await write_project_workspace_file(project, "committed.txt", "committed\n")
    add_event(
        env.db,
        project,
        "project.file.committed",
        "already committed",
        actor_agent_id=member.agent_id,
        metadata={**committed, "tool_call_id": "ambiguous-success"},
    )
    await env.db.commit()
    assert await _resolve_uncertain_project_file_audit(
        project,
        previous_head=initial_head,
        result=committed,
        tool_call_id="ambiguous-success",
        member=member,
        agent_id=member.agent_id,
        project_run=None,
        summary="already committed",
        metadata={**committed, "tool_call_id": "ambiguous-success"},
    ) is True
    assert (await repository_state(project))["head"] == committed["commit"]

    descendant_base = committed["commit"]
    ancestor = await write_project_workspace_file(project, "ancestor.txt", "ancestor\n")
    descendant = await write_project_workspace_file(project, "descendant.txt", "descendant\n")
    assert await _resolve_uncertain_project_file_audit(
        project,
        previous_head=descendant_base,
        result=ancestor,
        tool_call_id="descendant-recovery",
        member=member,
        agent_id=member.agent_id,
        project_run=None,
        summary="descendant recovery",
        metadata={**ancestor, "tool_call_id": "descendant-recovery"},
    ) is True
    assert (await repository_state(project))["head"] == descendant["commit"]
    recovered = (
        await env.db.execute(
            select(ProjectEvent).where(
                ProjectEvent.project_id == project.id,
                ProjectEvent.event_metadata["tool_call_id"].astext == "descendant-recovery",
            )
        )
    ).scalar_one()
    assert recovered.event_metadata["audit_recovered"] is True
    assert recovered.event_metadata["current_head"] == descendant["commit"]


async def test_project_runtime_tools_are_role_projected_and_double_enforced(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.chat_session import ChatSession
    from app.services.agent_tools import execute_tool
    from app.services.project_runtime_tools import execute_project_runtime_tool
    from app.services.subagent_runtime import prepare_subagent_tools

    env = project_api

    collaboration_bypass_names = {
        "send_message_to_agent",
        "send_file_to_agent",
        "send_message_to_parent",
        "send_session_message",
    }
    standard_file_names = {
        "delete_file",
        "edit_file",
        "find_files",
        "list_files",
        "move_file",
        "read_file",
        "search_files",
        "write_file",
    }
    structured_read_names = {"read_document", "read_image"}
    sandbox_names = {"execute_code", "execute_code_e2b", "execute_code_aio"}
    project_sandbox_names = {"execute_code"}

    async def normal_tools_with_collaboration_bypasses(_agent_id, *, assignment_snapshot=None):
        return [
            {
                "type": "function",
                "function": {
                    "name": name,
                    "description": "Generic collaboration path",
                    "parameters": {"type": "object", "properties": {}},
                },
            }
            for name in (
                collaboration_bypass_names
                | structured_read_names
                | standard_file_names
                | sandbox_names
            )
        ]

    monkeypatch.setattr(
        "app.services.agent_tools.get_agent_tools_for_llm",
        normal_tools_with_collaboration_bypasses,
    )
    project = await _create_project(env, name="Role projected tools")
    project_id = project["id"]
    await _mark_project_running(env, project_id)
    group = (await env.client.get(f"/api/projects/{project_id}/group-session")).json()

    assigned = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={
            "title": "Worker item",
            "description": "Worker-owned delivery",
            "assignee_agent_id": str(env.worker_id),
            "acceptance_criteria": ["Evidence attached"],
        },
    )
    other = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={
            "title": "Reviewer item",
            "description": "Reviewer-owned delivery",
            "assignee_agent_id": str(env.reviewer_id),
            "acceptance_criteria": ["Review complete"],
        },
    )
    assert assigned.status_code == other.status_code == 201

    worker_wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Worker runtime", "mentions": [str(env.worker_id)]},
    )
    leader_wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Leader runtime", "mentions": [str(env.leader_id)]},
    )
    assert worker_wake.status_code == leader_wake.status_code == 201
    worker_child_id = uuid.UUID(
        next(row for row in worker_wake.json()["subagent_runs"] if row["agent_id"] == str(env.worker_id))["session_id"]
    )
    leader_child_id = uuid.UUID(
        next(row for row in leader_wake.json()["subagent_runs"] if row["agent_id"] == str(env.leader_id))["session_id"]
    )

    worker_tools = await prepare_subagent_tools(
        env.worker_id,
        worker_child_id,
        execution_user_id=env.owner_id,
    )
    worker_names = {item["function"]["name"] for item in worker_tools}
    leader_names = {
        item["function"]["name"]
        for item in await prepare_subagent_tools(
            env.leader_id,
            leader_child_id,
            execution_user_id=env.owner_id,
        )
    }
    assert {"project_get_context", "project_list_work_items", "project_update_work_item", "project_message_agent"} <= worker_names
    assert {"project_list_files", "project_read_file", "project_write_file"}.isdisjoint(worker_names)
    assert standard_file_names <= worker_names
    for tool in worker_tools:
        if tool["function"]["name"] not in (
            standard_file_names | structured_read_names | project_sandbox_names
        ):
            continue
        parameters = tool["function"]["parameters"]
        assert "workspace" in parameters["required"]
        assert parameters["properties"]["workspace"]["enum"] == ["agent", "project"]
    assert "project_create_work_item" not in worker_names
    assert "project_update_plan" not in worker_names
    assert "project_set_status" not in worker_names
    assert worker_names.isdisjoint(collaboration_bypass_names)
    assert leader_names.isdisjoint(collaboration_bypass_names)
    assert structured_read_names | sandbox_names <= worker_names
    assert structured_read_names | sandbox_names <= leader_names
    e2b_tool = next(item for item in worker_tools if item["function"]["name"] == "execute_code_e2b")
    assert "workspace" not in e2b_tool["function"]["parameters"].get("properties", {})
    aio_tool = next(item for item in worker_tools if item["function"]["name"] == "execute_code_aio")
    assert "workspace" not in aio_tool["function"]["parameters"].get("properties", {})
    assert standard_file_names <= leader_names
    assert {
        "project_create_work_item",
        "project_update_plan",
        "project_restore_commit",
        "project_set_status",
    } <= leader_names

    with pytest.raises(ValueError, match="authorized project Subagent runtime"):
        await prepare_subagent_tools(
            env.worker_id,
            worker_child_id,
            execution_user_id=env.viewer_id,
        )
    with pytest.raises(ValueError, match="authorized project Subagent runtime"):
        await prepare_subagent_tools(
            env.reviewer_id,
            worker_child_id,
            execution_user_id=env.owner_id,
        )
    with pytest.raises(ValueError, match="authorized project Subagent runtime"):
        await execute_project_runtime_tool(
            "project_get_context",
            {},
            agent_id=env.worker_id,
            execution_user_id=env.viewer_id,
            session_id=str(worker_child_id),
            tool_call_id="wrong-execution-user",
            turn_anchor_id=None,
        )

    settings_update = await env.client.patch(
        f"/api/projects/{project_id}/settings",
        json={"policies": {"project_tools": {"participant_disabled": ["project_message_agent"]}}},
    )
    members = (await env.client.get(f"/api/projects/{project_id}/members")).json()
    worker_member = next(item for item in members if item["agent_id"] == str(env.worker_id))
    member_update = await env.client.patch(
        f"/api/projects/{project_id}/members/{worker_member['id']}",
        json={
            "config_snapshot": {
                **worker_member["config_snapshot"],
                "disabled_project_tools": ["project_update_work_item"],
            }
        },
    )
    assert settings_update.status_code == member_update.status_code == 200
    projected_names = {
        item["function"]["name"]
        for item in await prepare_subagent_tools(
            env.worker_id,
            worker_child_id,
            execution_user_id=env.owner_id,
        )
    }
    assert "project_message_agent" not in projected_names
    assert "project_update_work_item" in projected_names
    refreshed_wake = await env.client.post(
        f"/api/projects/{project_id}/group-sessions/{group['id']}/messages",
        json={"content": "Worker runtime after member policy change", "mentions": [str(env.worker_id)]},
    )
    assert refreshed_wake.status_code == 201, refreshed_wake.text
    refreshed_child_id = uuid.UUID(
        next(
            row
            for row in refreshed_wake.json()["subagent_runs"]
            if row["agent_id"] == str(env.worker_id)
        )["session_id"]
    )
    refreshed_names = {
        item["function"]["name"]
        for item in await prepare_subagent_tools(
            env.worker_id,
            refreshed_child_id,
            execution_user_id=env.owner_id,
        )
    }
    assert "project_message_agent" not in refreshed_names
    assert "project_update_work_item" not in refreshed_names
    with pytest.raises(ValueError, match="not allowed"):
        await execute_project_runtime_tool(
            "project_message_agent",
            {},
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(worker_child_id),
            tool_call_id="worker-policy-bypass",
            turn_anchor_id=None,
        )
    restore_member_tools = await env.client.patch(
        f"/api/projects/{project_id}/members/{worker_member['id']}",
        json={
            "config_snapshot": {
                **worker_member["config_snapshot"],
                "disabled_project_tools": [],
            }
        },
    )
    assert restore_member_tools.status_code == 200, restore_member_tools.text

    long_progress = "Implementation started: " + ("progress " * 100)
    long_evidence = [f"docs/progress-{index}.md:" + ("e" * 400) for index in range(8)]
    updated = await execute_project_runtime_tool(
        "project_update_work_item",
        {
            "work_item_id": assigned.json()["id"],
            "status": "in_progress",
            "progress_note": long_progress,
            "evidence": long_evidence,
        },
        agent_id=env.worker_id,
        execution_user_id=env.owner_id,
        session_id=str(worker_child_id),
        tool_call_id="worker-update",
        turn_anchor_id=None,
    )
    assert json.loads(updated)["status"] == "in_progress"

    with pytest.raises(ValueError, match="only status, progress_note and evidence"):
        await execute_project_runtime_tool(
            "project_update_work_item",
            {"work_item_id": assigned.json()["id"], "title": "Privilege escalation"},
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(worker_child_id),
            tool_call_id="worker-title",
            turn_anchor_id=None,
        )
    with pytest.raises(ValueError, match="assigned to themselves"):
        await execute_project_runtime_tool(
            "project_update_work_item",
            {"work_item_id": other.json()["id"], "status": "done"},
            agent_id=env.worker_id,
            execution_user_id=env.owner_id,
            session_id=str(worker_child_id),
            tool_call_id="worker-other",
            turn_anchor_id=None,
        )

    denied = await execute_tool(
        "project_update_plan",
        {"goal": "Participant must not change this"},
        env.worker_id,
        env.owner_id,
        session_id=str(worker_child_id),
        tool_call_id="worker-plan",
    )
    assert denied.startswith("❌")
    worker_member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == uuid.UUID(project_id),
                ProjectMemberSnapshot.agent_id == env.worker_id,
            )
        )
    ).scalar_one()
    stored_professional_role = ("Build traceable deliverables. " + ("role " * 100))[:500]
    worker_member.role_snapshot = stored_professional_role
    await env.db.commit()
    context = await execute_tool(
        "project_get_context",
        {},
        env.worker_id,
        env.owner_id,
        session_id=str(worker_child_id),
        tool_call_id="worker-context",
    )
    context_payload = json.loads(context)
    assert context_payload["id"] == project_id
    context_worker = next(row for row in context_payload["members"] if row["agent_id"] == str(env.worker_id))
    context_leader = next(row for row in context_payload["members"] if row["agent_id"] == str(env.leader_id))
    assert context_worker["project_role"] == "participant"
    assert context_leader["project_role"] == "owner"
    assert context_worker["professional_role"] == stored_professional_role.rstrip()

    work_items = json.loads(
        await execute_tool(
            "project_list_work_items",
            {"mine_only": True},
            env.worker_id,
            env.owner_id,
            session_id=str(worker_child_id),
            tool_call_id="worker-items",
        )
    )
    assert len(work_items) == 1
    assert work_items[0]["assignee_name"] == worker_member.name_snapshot
    assert work_items[0]["professional_role"] == context_worker["professional_role"]
    assert len(work_items[0]["progress"]) == 600
    assert work_items[0]["progress"].endswith("…")
    assert len(work_items[0]["evidence"]) == 6
    assert all(len(value) == 300 and value.endswith("…") for value in work_items[0]["evidence"])
    assert await env.db.get(ChatSession, worker_child_id) is not None


async def test_work_item_mutations_update_dashboard_and_audit(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Work item audit")
    project_id = project["id"]

    parent = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={"title": "Prepare evidence", "priority": "high"},
    )
    assert parent.status_code == 201, parent.text
    child = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={
            "title": "Review evidence",
            "parent_id": parent.json()["id"],
            "dependency_ids": [parent.json()["id"]],
            "assignee_agent_id": str(env.reviewer_id),
            "status": "todo",
        },
    )
    assert child.status_code == 201, child.text
    completed = await env.client.patch(
        f"/api/projects/{project_id}/work-items/{parent.json()['id']}",
        json={"status": "done", "priority": "urgent"},
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()["status"] == "done"

    dashboard = (await env.client.get(f"/api/projects/{project_id}/dashboard")).json()
    assert dashboard["progress"] == 50
    assert dashboard["work_item_counts"] == {"done": 1, "todo": 1}
    assert len(dashboard["work_items"]) == 2

    events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
    created = [event for event in events if event["event_type"] == "work_item.created"]
    updated = next(event for event in events if event["event_type"] == "work_item.updated")
    assert len(created) == 2
    assert updated["work_item_id"] == parent.json()["id"]
    assert updated["event_metadata"]["after"]["status"] == "done"
    assert {"status", "priority"} <= set(updated["event_metadata"]["changed_fields"])


async def test_run_work_item_and_milestone_contracts_are_explicit(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.project import ProjectRun
    from app.services.project_runtime_tools import (
        execute_project_runtime_tool,
        execute_project_workspace_tool,
    )

    env = project_api
    project = await _create_project(env, name="Explicit trace contract")
    project_id = project["id"]
    stored_project = await env.db.get(Project, uuid.UUID(project_id))
    assert stored_project is not None
    stored_project.status = "running"
    await env.db.commit()

    item_response = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={
            "title": "Ship traced artifact",
            "assignee_agent_id": str(env.leader_id),
            "acceptance_criteria": ["Artifact and milestone are linked"],
        },
    )
    assert item_response.status_code == 201, item_response.text
    item_id = item_response.json()["id"]
    run_response = await env.client.post(
        f"/api/projects/{project_id}/runs",
        json={"work_item_id": item_id, "agent_id": str(env.leader_id)},
    )
    assert run_response.status_code == 201, run_response.text
    run = run_response.json()
    assert run["work_item_id"] == item_id
    assert run["agent_name"] == "Leader"
    assert run["project_member_id"]
    assert run["member_snapshot"]["name"] == "Leader"
    assert run["session_id"] == run["subagent_session_id"]
    child_id = uuid.UUID(run["subagent_session_id"])
    child_input = (
        await env.db.execute(
            select(ChatMessage)
            .where(
                ChatMessage.conversation_id == str(child_id),
                ChatMessage.role == "user",
            )
            .order_by(ChatMessage.created_at.desc())
            .limit(1)
        )
    ).scalar_one()
    assert child_input.message_meta["project_run_id"] == run["id"]
    anchor_id = child_input.id

    await execute_project_runtime_tool(
        "project_update_work_item",
        {
            "work_item_id": item_id,
            "status": "in_progress",
            "progress_note": "Implementation is traceable",
            "evidence": ["deliverables/traced.md"],
        },
        agent_id=env.leader_id,
        execution_user_id=env.owner_id,
        session_id=str(child_id),
        tool_call_id="trace-update",
        turn_anchor_id=anchor_id,
    )
    write_result = json.loads(
        await execute_project_workspace_tool(
            "write_file",
            {
                "workspace": "project",
                "path": "deliverables/traced.md",
                "content": "# Durable evidence\n",
            },
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="trace-file",
            turn_anchor_id=anchor_id,
        )
    )
    stored_run = await env.db.get(ProjectRun, uuid.UUID(run["id"]))
    assert stored_run is not None
    stored_run.status = "succeeded"
    await env.db.commit()
    long_milestone_message = "里程碑完整说明：" + "六个Agent的交付证据、评审结论与回滚锚点均已核验。" * 40
    assert len(long_milestone_message) > 500
    milestone_result = json.loads(
        await execute_project_runtime_tool(
            "project_create_milestone",
            {
                "message": long_milestone_message,
                # Requested milestone paths are operation intent, not proof of
                # an actual change. Both files are clean here, so the empty
                # milestone commit must not make PROJECT.json appear in the
                # work item's code/file change list.
                "paths": ["deliverables/traced.md", "PROJECT.json"],
                "related_work_item_ids": [item_id],
            },
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="trace-milestone",
            turn_anchor_id=anchor_id,
        )
    )

    detail_response = await env.client.get(f"/api/projects/{project_id}/work-items/{item_id}")
    assert detail_response.status_code == 200, detail_response.text
    detail = detail_response.json()
    assert detail["work_item"]["id"] == item_id
    assert [row["id"] for row in detail["runs"]] == [run["id"]]
    assert detail["sessions"][0] == {
        "run_id": run["id"],
        "project_run_id": run["id"],
        "work_item_id": item_id,
        "agent_id": str(env.leader_id),
        "agent_name": "Leader",
        "status": detail["runs"][0]["status"],
        "source_channel": "subagent",
        "session_intent": "execution",
        "session_id": str(child_id),
        "subagent_session_id": str(child_id),
        "anchor_message_id": str(anchor_id),
    }
    assert {row["path"] for row in detail["files"]} == {"deliverables/traced.md"}
    assert detail["files"][0]["commit"] == write_result["commit"]
    assert {row["commit"] for row in detail["commits"]} >= {
        write_result["commit"],
        milestone_result["commit"],
    }
    milestone_trace = next(row for row in detail["commits"] if row["commit"] == milestone_result["commit"])
    assert milestone_trace["paths"] == []
    assert milestone_trace["diff_available"] is True
    assert any(row["kind"] == "progress_note" for row in detail["evidence"])
    assert any(row["kind"] == "evidence" for row in detail["evidence"])
    traced_events = [
        row
        for row in detail["events"]
        if row["event_type"] in {"work_item.updated", "project.file.committed", "git.milestone.created"}
    ]
    assert all(row["run_id"] == run["id"] for row in traced_events)
    assert all(row["work_item_id"] == item_id for row in traced_events)

    # A durable Leader child is reused across work items. Session equality is
    # only a conversation entry point and must never pull another work item's
    # Run, events, commits, files, or evidence into this detail contract.
    other_item_response = await env.client.post(
        f"/api/projects/{project_id}/work-items",
        json={"title": "Unrelated work", "assignee_agent_id": str(env.leader_id)},
    )
    assert other_item_response.status_code == 201
    other_item_id = other_item_response.json()["id"]
    unrelated_run = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=uuid.UUID(project_id),
        work_item_id=uuid.UUID(other_item_id),
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="succeeded",
        trigger_type="manual",
        output={"subagent_session_id": str(child_id)},
    )
    env.db.add(unrelated_run)
    await env.db.flush()
    foreign_commit = "f" * 40
    env.db.add(
        ProjectEvent(
            tenant_id=env.tenant_id,
            project_id=uuid.UUID(project_id),
            work_item_id=uuid.UUID(other_item_id),
            run_id=unrelated_run.id,
            actor_agent_id=env.leader_id,
            event_type="foreign.work.evidence",
            summary="Must stay with the other work item",
            event_metadata={
                "project_run_id": str(unrelated_run.id),
                "session_id": str(child_id),
                "subagent_session_id": str(child_id),
                "commit": foreign_commit,
                "path": "foreign/only.md",
                "evidence": ["foreign-evidence"],
            },
        )
    )
    await env.db.commit()
    detail_again = (await env.client.get(f"/api/projects/{project_id}/work-items/{item_id}")).json()
    assert [row["id"] for row in detail_again["runs"]] == [run["id"]]
    assert all(row["event_type"] != "foreign.work.evidence" for row in detail_again["events"])
    assert all(row["commit"] != foreign_commit for row in detail_again["commits"])
    assert all(row["path"] != "foreign/only.md" for row in detail_again["files"])
    assert all(row["value"] != "foreign-evidence" for row in detail_again["evidence"])
    assert len(detail_again["sessions"]) == 1

    milestones_response = await env.client.get(f"/api/projects/{project_id}/milestones")
    assert milestones_response.status_code == 200, milestones_response.text
    milestone = milestones_response.json()[0]
    assert milestone["commit"] == milestone_result["commit"]
    assert milestone["run_id"] == run["id"]
    assert milestone["work_item_id"] == item_id
    assert milestone["session_id"] == str(child_id)
    assert milestone["subagent_session_id"] == str(child_id)
    assert milestone["agent_name"] == "Leader"
    assert milestone["related_run_ids"] == [run["id"]]
    assert milestone["related_work_item_ids"] == [item_id]
    assert milestone["message"] == long_milestone_message
    milestone_event = await env.db.get(ProjectEvent, uuid.UUID(milestone_result["event_id"]))
    assert milestone_event is not None
    assert milestone_event.event_type == "git.milestone.created"
    assert len(milestone_event.summary) <= 500
    assert milestone_event.event_metadata["milestone_message"] == long_milestone_message
    assert milestone_event.event_metadata["description"] == long_milestone_message

    # A repository commit can win immediately before the metadata transaction
    # fails.  The durable prepared event and Git operation trailer let a retry
    # finalize that same semantic operation without another empty commit.
    from app.services import project_runtime_tools

    commit_attempts = {"value": 0}

    class FailSecondCommitSession:
        def __init__(self):
            self._session = env.session_factory()

        async def __aenter__(self):
            await self._session.__aenter__()
            return self

        async def __aexit__(self, exc_type, exc, traceback):
            return await self._session.__aexit__(exc_type, exc, traceback)

        def __getattr__(self, name):
            return getattr(self._session, name)

        async def commit(self):
            commit_attempts["value"] += 1
            if commit_attempts["value"] == 2:
                await self._session.rollback()
                raise RuntimeError("injected metadata commit failure")
            await self._session.commit()

    recovery_arguments = {
        "message": "Recover this semantic milestone after DB failure",
        "related_work_item_ids": [item_id],
    }
    monkeypatch.setattr(project_runtime_tools, "async_session", FailSecondCommitSession)
    with pytest.raises(RuntimeError, match="injected metadata commit failure"):
        await execute_project_runtime_tool(
            "project_create_milestone",
            recovery_arguments,
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="recovery-first-call",
            turn_anchor_id=anchor_id,
        )

    repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    operation_commits_after_failure = subprocess.check_output(
        ["git", "-C", str(repo), "log", "--format=%H", "--fixed-strings", "--grep=Project-Milestone-Operation:"],
        text=True,
    ).splitlines()
    assert len(operation_commits_after_failure) == 2

    late_success = ProjectRun(
        tenant_id=env.tenant_id,
        project_id=uuid.UUID(project_id),
        work_item_id=uuid.UUID(item_id),
        agent_id=env.leader_id,
        initiated_by_user_id=env.owner_id,
        status="succeeded",
        trigger_type="retry",
        output={"subagent_session_id": str(child_id)},
    )
    env.db.add(late_success)
    await env.db.commit()

    monkeypatch.setattr(project_runtime_tools, "async_session", env.session_factory)
    recovered_result = json.loads(
        await execute_project_runtime_tool(
            "project_create_milestone",
            recovery_arguments,
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            # A regenerated model tool call has another tool_call_id; semantic
            # milestone identity must still recover the original operation.
            tool_call_id="recovery-regenerated-call",
            turn_anchor_id=anchor_id,
        )
    )
    assert recovered_result["idempotent_replay"] is True
    assert recovered_result["commit"] == operation_commits_after_failure[0]
    operation_commits_after_retry = subprocess.check_output(
        ["git", "-C", str(repo), "log", "--format=%H", "--fixed-strings", "--grep=Project-Milestone-Operation:"],
        text=True,
    ).splitlines()
    assert operation_commits_after_retry == operation_commits_after_failure
    recovered_events = (
        (
            await env.db.execute(
                select(ProjectEvent).where(
                    ProjectEvent.project_id == uuid.UUID(project_id),
                    ProjectEvent.event_type.in_(["git.milestone.prepared", "git.milestone.created"]),
                )
            )
        )
        .scalars()
        .all()
    )
    operation_keys = [
        event.event_metadata.get("milestone_operation_key")
        for event in recovered_events
        if event.event_metadata.get("milestone_message") == recovery_arguments["message"]
    ]
    assert len(operation_keys) == 1
    assert (
        next(
            event
            for event in recovered_events
            if event.event_metadata.get("milestone_message") == recovery_arguments["message"]
        ).event_type
        == "git.milestone.created"
    )
    recovered_event = next(
        event
        for event in recovered_events
        if event.event_metadata.get("milestone_message") == recovery_arguments["message"]
    )
    assert set(recovered_event.event_metadata["related_run_ids"]) == {run["id"], str(late_success.id)}
    assert not subprocess.check_output(["git", "-C", str(repo), "status", "--porcelain"], text=True).strip()

    completed = json.loads(
        await execute_project_runtime_tool(
            "project_set_status",
            {"status": "completed", "reason": "Acceptance evidence is committed"},
            agent_id=env.leader_id,
            execution_user_id=env.owner_id,
            session_id=str(child_id),
            tool_call_id="trace-complete",
            turn_anchor_id=anchor_id,
        )
    )
    assert completed["status"] == "completed"
    env.db.expire(stored_project)
    refreshed_project = (await env.client.get(f"/api/projects/{project_id}")).json()
    assert refreshed_project["status"] == "completed"
    project_events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
    status_event = next(row for row in project_events if row["event_type"] == "project.status.updated")
    assert status_event["run_id"] == run["id"]
    assert status_event["event_metadata"]["after"] == "completed"


async def test_file_commits_restore_branch_and_milestone_preserve_history(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Git evidence")
    project_id = project["id"]

    initial_git = await env.client.get(f"/api/projects/{project_id}/git")
    assert initial_git.status_code == 200, initial_git.text
    initial_head = initial_git.json()["head"]
    assert {"README.md", "PROJECT.json"} <= set(initial_git.json()["files"])

    write_response = await env.client.put(
        f"/api/projects/{project_id}/files",
        json={"path": "deliverables/report.md", "content": "# Evidence\n\nversion one\n"},
    )
    assert write_response.status_code == 200, write_response.text
    write_commit = write_response.json()["commit"]
    assert write_commit != initial_head

    diff_response = await env.client.get(
        f"/api/projects/{project_id}/git/diff",
        params={"commit": write_commit, "path": "deliverables/report.md"},
    )
    assert diff_response.status_code == 200, diff_response.text
    diff_payload = diff_response.json()
    assert diff_payload["parent_commit"] == initial_head
    assert diff_payload["files"][0]["path"] == "deliverables/report.md"
    assert diff_payload["files"][0]["original_content"] == ""
    assert diff_payload["files"][0]["modified_content"].startswith("# Evidence")

    unsafe_diff = await env.client.get(
        f"/api/projects/{project_id}/git/diff",
        params={"commit": write_commit, "path": "../outside.md"},
    )
    assert unsafe_diff.status_code == 422
    env.authenticate_as(env.viewer_id)
    hidden_diff = await env.client.get(
        f"/api/projects/{project_id}/git/diff",
        params={"commit": write_commit},
    )
    assert hidden_diff.status_code == 404
    env.authenticate_as(env.owner_id)

    for unsafe_path in ("../outside.md", ".git/config", ".GIT/config", "deliverables//hidden.md"):
        rejected = await env.client.put(
            f"/api/projects/{project_id}/files",
            json={"path": unsafe_path, "content": "must not be written"},
        )
        assert rejected.status_code == 422, rejected.text

    unchanged = await env.client.put(
        f"/api/projects/{project_id}/files",
        json={"path": "deliverables/report.md", "content": "# Evidence\n\nversion one\n"},
    )
    assert unchanged.status_code == 409

    files_response = await env.client.get(f"/api/projects/{project_id}/files")
    assert files_response.status_code == 200, files_response.text
    files_payload = files_response.json()
    files = files_payload["files"] if isinstance(files_payload, dict) else files_payload
    assert any(
        (item == "deliverables/report.md") or (isinstance(item, dict) and item.get("path") == "deliverables/report.md")
        for item in files
    )

    milestone = await env.client.post(
        f"/api/projects/{project_id}/git/commit",
        json={"message": "Acceptance milestone", "milestone": True},
    )
    assert milestone.status_code == 200, milestone.text
    milestone_commit = milestone.json()["commit"]

    restore = await env.client.post(
        f"/api/projects/{project_id}/git/restore",
        json={"commit": initial_head, "message": "Restore initial project tree"},
    )
    assert restore.status_code == 200, restore.text
    restore_commit = restore.json()["commit"]
    assert restore_commit not in {initial_head, write_commit, milestone_commit}

    branch = await env.client.post(
        f"/api/projects/{project_id}/git/branches",
        json={"name": "review/version-one", "from_commit": write_commit},
    )
    assert branch.status_code == 200, branch.text
    assert branch.json()["from_commit"] == write_commit

    final_git = (await env.client.get(f"/api/projects/{project_id}/git?limit=20")).json()
    commit_ids = {item["commit"] for item in final_git["commits"]}
    assert {initial_head, write_commit, milestone_commit, restore_commit} <= commit_ids
    assert "review/version-one" in final_git["branches"]

    # Prove the old commits still resolve in the real repository: restore is a
    # new commit, never reset/force-push history rewriting.
    repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    for commit in (initial_head, write_commit, milestone_commit, restore_commit):
        resolved = subprocess.run(
            ["git", "-C", str(repo), "cat-file", "-e", f"{commit}^{{commit}}"],
            check=False,
            capture_output=True,
        )
        assert resolved.returncode == 0

    events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
    event_types = {event["event_type"] for event in events}
    assert {
        "project.file.committed",
        "git.milestone.created",
        "git.restore_commit.created",
        "git.branch.created",
    } <= event_types

    db_events = (
        (await env.db.execute(select(ProjectEvent).where(ProjectEvent.project_id == uuid.UUID(project_id))))
        .scalars()
        .all()
    )
    assert len(db_events) == len(events)
    assert any(event.event_metadata.get("commit") == restore_commit for event in db_events)


async def test_project_head_file_preview_media_range_and_acl(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="HEAD file previews")
    project_id = project["id"]
    repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    (repo / "assets").mkdir()
    binary = b"\x00\x01\x02\x03\x04\x05\xff\x10"
    (repo / "assets" / "sample.mp4").write_bytes(binary)
    (repo / "large.txt").write_text("x" * (1024 * 1024 + 1), encoding="utf-8")
    subprocess.run(
        ["git", "-C", str(repo), "add", "--", "assets/sample.mp4", "large.txt"],
        check=True,
    )
    subprocess.run(
        ["git", "-C", str(repo), "commit", "-m", "Add preview fixtures"],
        check=True,
        capture_output=True,
    )

    env.authenticate_as(env.viewer_id)
    hidden = await env.client.get(
        f"/api/projects/{project_id}/files/content",
        params={"path": "assets/sample.mp4"},
    )
    assert hidden.status_code == 404

    env.authenticate_as(env.owner_id)
    for unsafe_path in ("../README.md", ".git/config", "nested/.GIT/config", "assets//sample.mp4"):
        rejected = await env.client.get(
            f"/api/projects/{project_id}/files/content",
            params={"path": unsafe_path},
        )
        assert rejected.status_code == 422, rejected.text

    listing = (await env.client.get(f"/api/projects/{project_id}/files")).json()
    media_item = next(item for item in listing if item["path"] == "assets/sample.mp4")
    assert media_item["kind"] == "video"
    assert media_item["mime_type"] == "video/mp4"
    assert media_item["preview"] == ""
    assert media_item["is_editable"] is False
    large_item = next(item for item in listing if item["path"] == "large.txt")
    assert large_item["kind"] == "text"
    assert large_item["is_editable"] is False

    media = await env.client.get(
        f"/api/projects/{project_id}/files/content",
        params={"path": "assets/sample.mp4"},
    )
    assert media.status_code == 200, media.text
    media_payload = media.json()
    assert media_payload["content"] is None
    assert media_payload["raw_url"].startswith(f"/api/projects/{project_id}/files/raw?")

    partial = await env.client.get(media_payload["raw_url"], headers={"Range": "bytes=2-5"})
    assert partial.status_code == 206, partial.text
    assert partial.content == binary[2:6]
    assert partial.headers["content-type"] == "video/mp4"
    assert partial.headers["content-range"] == f"bytes 2-5/{len(binary)}"
    assert partial.headers["accept-ranges"] == "bytes"
    assert partial.headers["content-disposition"].startswith("inline;")

    downloaded = await env.client.get(media_payload["download_url"])
    assert downloaded.status_code == 200
    assert downloaded.content == binary
    assert downloaded.headers["content-disposition"].startswith("attachment;")
    raw_head = await env.client.head(media_payload["raw_url"])
    assert raw_head.status_code == 200
    assert raw_head.content == b""
    assert raw_head.headers["content-length"] == str(len(binary))
    invalid_range = await env.client.get(media_payload["raw_url"], headers={"Range": "bytes=99-100"})
    assert invalid_range.status_code == 416
    assert invalid_range.headers["content-range"] == f"bytes */{len(binary)}"

    large = await env.client.get(
        f"/api/projects/{project_id}/files/content",
        params={"path": "large.txt", "max_chars": 64},
    )
    assert large.status_code == 200
    assert large.json()["content"] == "x" * 64
    assert large.json()["truncated"] is True
    assert large.json()["is_editable"] is False

    shared = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert shared.status_code == 200
    env.authenticate_as(env.viewer_id)
    viewer_content = await env.client.get(
        f"/api/projects/{project_id}/files/content",
        params={"path": "assets/sample.mp4"},
    )
    assert viewer_content.status_code == 200

    env.authenticate_as(env.owner_id)
    revoked = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "private", "shared_with_user_ids": []},
    )
    assert revoked.status_code == 200
    # Raw tickets are re-authorized on every request, not bearer URLs that
    # outlive revoked project access.
    old_viewer_raw = await env.client.get(viewer_content.json()["raw_url"])
    assert old_viewer_raw.status_code == 404


async def test_project_directory_archive_and_html_preview_use_one_immutable_head(project_api: ProjectApiEnv):
    env = project_api
    project = await _create_project(env, name="Immutable preview")
    project_id = project["id"]
    repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    (repo / "site" / "assets").mkdir(parents=True)
    (repo / "site" / "index.html").write_text(
        '<!doctype html><link rel="stylesheet" href="./style.css"><script src="./app.js"></script>'
        '<img src="./assets/logo.svg">',
        encoding="utf-8",
    )
    (repo / "site" / "style.css").write_text("body { color: rgb(1, 2, 3); }\n", encoding="utf-8")
    (repo / "site" / "app.js").write_text("document.body.dataset.loaded = 'yes';\n", encoding="utf-8")
    (repo / "site" / "assets" / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><circle r="4" cx="4" cy="4"/></svg>',
        encoding="utf-8",
    )
    subprocess.run(["git", "-C", str(repo), "add", "--", "site"], check=True)
    subprocess.run(["git", "-C", str(repo), "commit", "-m", "Add preview site"], check=True, capture_output=True)
    snapshot_head = subprocess.check_output(["git", "-C", str(repo), "rev-parse", "HEAD"], text=True).strip()
    (repo / "site" / "untracked-secret.txt").write_text("must-not-download", encoding="utf-8")

    archive_ticket = await env.client.get(
        f"/api/projects/{project_id}/files/archive",
        params={"path": "site"},
    )
    assert archive_ticket.status_code == 200, archive_ticket.text
    archive_payload = archive_ticket.json()
    assert archive_payload["head"] == snapshot_head
    assert archive_payload["file_count"] == 4

    html_content = await env.client.get(
        f"/api/projects/{project_id}/files/content",
        params={"path": "site/index.html"},
    )
    assert html_content.status_code == 200, html_content.text
    preview_url = html_content.json()["html_preview_url"]
    assert preview_url.startswith(f"/api/projects/{project_id}/files/preview/")

    # Change HEAD after both tickets were issued. Their resources must remain
    # an internally consistent snapshot rather than mixing new HEAD content.
    (repo / "site" / "style.css").write_text("body { color: red; }\n", encoding="utf-8")
    subprocess.run(["git", "-C", str(repo), "add", "--", "site/style.css"], check=True)
    subprocess.run(["git", "-C", str(repo), "commit", "-m", "Change preview style"], check=True, capture_output=True)

    archive = await env.client.get(archive_payload["download_url"])
    assert archive.status_code == 200, archive.text
    assert archive.headers["content-type"].startswith("application/zip")
    assert archive.headers["x-project-git-head"] == snapshot_head
    with zipfile.ZipFile(io.BytesIO(archive.content)) as bundle:
        assert set(bundle.namelist()) == {
            "site/",
            "site/app.js",
            "site/assets/",
            "site/assets/logo.svg",
            "site/index.html",
            "site/style.css",
        }
        assert bundle.read("site/style.css") == b"body { color: rgb(1, 2, 3); }\n"
        assert "site/untracked-secret.txt" not in bundle.namelist()

    preview = await env.client.get(preview_url)
    assert preview.status_code == 200, preview.text
    assert preview.headers["x-project-git-head"] == snapshot_head
    preview_csp = preview.headers["content-security-policy"]
    assert "sandbox allow-scripts" in preview_csp
    assert "connect-src 'none'" in preview_csp
    assert "allow-same-origin" not in preview_csp
    assert preview_url.split("/files/preview/", 1)[1].split("/", 1)[0] not in preview_csp
    assert len(preview_csp) < 2_048
    preview_base = preview_url.rsplit("/", 1)[0]
    css = await env.client.get(f"{preview_base}/style.css")
    script = await env.client.get(f"{preview_base}/app.js")
    image = await env.client.get(f"{preview_base}/assets/logo.svg")
    assert css.text == "body { color: rgb(1, 2, 3); }\n"
    assert css.headers["x-project-git-head"] == snapshot_head
    assert script.status_code == 200 and script.headers["content-type"].startswith("text/javascript")
    assert image.status_code == 200 and image.headers["content-type"].startswith("image/svg+xml")

    tampered_archive = await env.client.get(archive_payload["download_url"].replace("path=site", "path=site%2Fassets"))
    assert tampered_archive.status_code == 401

    shared = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert shared.status_code == 200
    env.authenticate_as(env.viewer_id)
    viewer_archive = (
        await env.client.get(
            f"/api/projects/{project_id}/files/archive",
            params={"path": "site"},
        )
    ).json()["download_url"]
    viewer_preview = (
        await env.client.get(
            f"/api/projects/{project_id}/files/content",
            params={"path": "site/index.html"},
        )
    ).json()["html_preview_url"]
    env.authenticate_as(env.owner_id)
    revoked = await env.client.patch(
        f"/api/projects/{project_id}",
        json={"visibility": "private", "shared_with_user_ids": []},
    )
    assert revoked.status_code == 200
    assert (await env.client.get(viewer_archive)).status_code == 404
    assert (await env.client.get(viewer_preview)).status_code == 404


async def test_owner_manages_provider_neutral_remotes_and_atomically_clones(
    project_api: ProjectApiEnv,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.services import project_git_service

    async def allow_local_git_fixture(raw_url: str) -> str:
        return project_git_service._safe_remote_url(raw_url)

    monkeypatch.setattr(project_git_service, "validate_project_remote_url", allow_local_git_fixture)
    env = project_api
    project = await _create_project(env, name="Clone target")
    project_id = project["id"]
    initial = (await env.client.get(f"/api/projects/{project_id}/git")).json()
    initial_repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
    initial_files = set(
        subprocess.run(
            ["git", "-C", str(initial_repo), "ls-tree", "-r", "--name-only", "HEAD"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.splitlines()
    )

    empty = await env.client.get(f"/api/projects/{project_id}/git/remotes")
    assert empty.status_code == 200
    assert empty.json()["items"] == []
    configured = await env.client.put(
        f"/api/projects/{project_id}/git/remotes/upstream",
        json={"url": "https://example.com/vendor/project.git"},
    )
    assert configured.status_code == 200, configured.text
    assert configured.json()["created"] is True
    assert (await env.client.get(f"/api/projects/{project_id}/git")).json()["head"] == initial["head"]
    assert (await env.client.get(f"/api/projects/{project_id}/git/remotes")).json()["items"] == [
        {"name": "upstream", "url": "https://example.com/vendor/project.git"}
    ]
    ssh_remote = await env.client.put(
        f"/api/projects/{project_id}/git/remotes/ssh-upstream",
        json={"url": "ssh://git@example.com/vendor/project.git"},
    )
    assert ssh_remote.status_code == 200, ssh_remote.text
    assert ssh_remote.json()["url"] == "ssh://git@example.com/vendor/project.git"
    scp_remote = await env.client.put(
        f"/api/projects/{project_id}/git/remotes/scp-upstream",
        json={"url": "git@example.com:vendor/project.git"},
    )
    assert scp_remote.status_code == 200, scp_remote.text

    for name, url in [
        ("origin", "file:///tmp/repository.git"),
        ("origin", "/tmp/repository.git"),
        ("origin", "https://token@example.com/vendor/project.git"),
        ("origin", "ssh://git:secret@example.com/vendor/project.git"),
        ("origin", "ssh://git@example.com/vendor/$(touch-pwned).git"),
        ("origin", "ssh://git@example.com/vendor/repo.git;touch-pwned"),
        ("origin", "ssh://git@example.com/vendor/%24%28touch-pwned%29.git"),
        ("origin", "ssh://git@example.com/vendor/../private.git"),
        ("origin", "git@example.com:vendor/project.git;touch-pwned"),
        ("origin", "git@example.com:vendor/$(touch-pwned).git"),
        ("origin", "git@example.com:vendor/../private.git"),
        ("--upload-pack", "https://example.com/vendor/project.git"),
    ]:
        rejected = await env.client.put(
            f"/api/projects/{project_id}/git/remotes/{name}",
            json={"url": url},
        )
        assert rejected.status_code == 422, rejected.text

    env.authenticate_as(env.viewer_id)
    hidden = await env.client.get(f"/api/projects/{project_id}/git/remotes")
    assert hidden.status_code == 404
    env.authenticate_as(env.owner_id)
    deleted = await env.client.delete(f"/api/projects/{project_id}/git/remotes/upstream")
    assert deleted.status_code == 200, deleted.text
    deleted_ssh = await env.client.delete(f"/api/projects/{project_id}/git/remotes/ssh-upstream")
    assert deleted_ssh.status_code == 200, deleted_ssh.text
    deleted_scp = await env.client.delete(f"/api/projects/{project_id}/git/remotes/scp-upstream")
    assert deleted_scp.status_code == 200, deleted_scp.text

    source = tmp_path / "source"
    bare = tmp_path / "served" / "source.git"
    source.mkdir()
    subprocess.run(["git", "-C", str(source), "init", "-b", "main"], check=True, capture_output=True)
    subprocess.run(["git", "-C", str(source), "config", "user.name", "Source Author"], check=True)
    subprocess.run(["git", "-C", str(source), "config", "user.email", "source@example.test"], check=True)
    (source / "SOURCE.md").write_text("# Imported source\n", encoding="utf-8")
    subprocess.run(["git", "-C", str(source), "add", "--", "SOURCE.md"], check=True)
    subprocess.run(["git", "-C", str(source), "commit", "-m", "Imported baseline"], check=True)
    bare.parent.mkdir()
    subprocess.run(["git", "clone", "--bare", "--", str(source), str(bare)], check=True, capture_output=True)
    subprocess.run(["git", "--git-dir", str(bare), "update-server-info"], check=True)
    with socket.socket() as reservation:
        reservation.bind(("127.0.0.1", 0))
        port = reservation.getsockname()[1]
    server = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "http.server",
            str(port),
            "--bind",
            "127.0.0.1",
            "--directory",
            str(bare.parent),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        for _attempt in range(50):
            try:
                with socket.create_connection(("127.0.0.1", port), timeout=0.1):
                    break
            except OSError:
                time.sleep(0.02)
        else:
            raise AssertionError("local Git HTTP fixture did not start")
        clone_url = f"http://127.0.0.1:{port}/source.git"

        # A prepared journal survives the exact process-exit window after the
        # filesystem swap and restores the initialization repository on access.
        from app.models.project import ProjectRepositoryOperation

        model = await env.db.get(Project, uuid.UUID(project_id))
        assert model is not None
        staged_only = await project_git_service.begin_project_repository_clone(model, clone_url, "main")
        env.db.add(
            ProjectRepositoryOperation(
                id=staged_only.id,
                tenant_id=model.tenant_id,
                project_id=model.id,
                operation_type="clone",
                state="prepared",
                old_head=staged_only.old_head,
                new_head=staged_only.new_head,
                backup_name=staged_only.backup.name,
                staging_name=staged_only.staging_root.name,
            )
        )
        await env.db.commit()
        assert staged_only.staging_root.exists()
        assert staged_only.lock_handle is not None
        staged_only.lock_handle.close()  # Simulate exit before the filesystem swap.
        staged_only.lock_handle = None
        assert (
            await project_git_service.reconcile_project_repository_operations(
                model.id,
                db=env.db,
            )
            == 1
        )
        await env.db.commit()
        assert not staged_only.staging_root.exists()
        assert not (staged_only.repo / "SOURCE.md").exists()

        abandoned = await project_git_service.begin_project_repository_clone(model, clone_url, "main")
        env.db.add(
            ProjectRepositoryOperation(
                id=abandoned.id,
                tenant_id=model.tenant_id,
                project_id=model.id,
                operation_type="clone",
                state="prepared",
                old_head=abandoned.old_head,
                new_head=abandoned.new_head,
                backup_name=abandoned.backup.name,
                staging_name=abandoned.staging_root.name,
            )
        )
        await env.db.commit()
        await project_git_service.apply_project_repository_clone(abandoned)
        assert (abandoned.repo / "SOURCE.md").exists()
        assert abandoned.lock_handle is not None
        abandoned.lock_handle.close()  # Simulate OS releasing flock on process death.
        abandoned.lock_handle = None
        assert (
            await project_git_service.reconcile_project_repository_operations(
                model.id,
                db=env.db,
            )
            == 1
        )
        await env.db.commit()
        assert not (abandoned.repo / "SOURCE.md").exists()
        assert list(abandoned.repo.parent.glob(".repo-backup-*")) == []
        assert (await env.db.execute(select(ProjectRepositoryOperation))).scalars().all() == []

        # The repository swap stays compensatable until settings and the
        # audit row are durable. A database failure restores the generated
        # initialization baseline and leaves the endpoint safely retryable.
        original_commit = AsyncSession.commit
        clone_commit_count = {"value": 0}

        async def fail_clone_commit_once(self: AsyncSession):
            if self is env.db:
                clone_commit_count["value"] += 1
            if self is env.db and clone_commit_count["value"] == 2:
                raise RuntimeError("forced clone metadata commit failure")
            return await original_commit(self)

        monkeypatch.setattr(AsyncSession, "commit", fail_clone_commit_once)
        with pytest.raises(RuntimeError, match="forced clone metadata commit failure"):
            await env.client.post(
                f"/api/projects/{project_id}/git/clone",
                json={"url": clone_url, "branch": "main"},
            )
        monkeypatch.setattr(AsyncSession, "commit", original_commit)

        restored_repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
        restored_files = subprocess.run(
            ["git", "-C", str(restored_repo), "ls-tree", "-r", "--name-only", "HEAD"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.splitlines()
        assert set(restored_files) == initial_files
        assert not (restored_repo / "SOURCE.md").exists()
        assert (
            subprocess.run(
                ["git", "-C", str(restored_repo), "rev-parse", "HEAD"],
                check=True,
                capture_output=True,
                text=True,
            ).stdout.strip()
            == initial["head"]
        )
        assert list(restored_repo.parent.glob(".repo-backup-*")) == []
        failed_settings = (await env.client.get(f"/api/projects/{project_id}/settings")).json()
        assert failed_settings["git"].get("source") != "cloned"
        failed_events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
        assert "git.repository.cloned" not in {event["event_type"] for event in failed_events}

        # A transport error after the database accepted COMMIT must not roll
        # the repository back to a state that contradicts durable metadata.
        tenant = await env.db.get(Tenant, env.tenant_id)
        assert tenant is not None
        ambiguous_leader = await _agent(env.db, tenant, env.owner, "Clone Leader", "Drive clone recovery")
        await env.db.commit()
        ambiguous_project = await _create_project(
            env,
            name="Ambiguous clone commit",
            leader_id=ambiguous_leader.id,
        )
        ambiguous_project_id = ambiguous_project["id"]
        ambiguous_commit_count = {"value": 0}

        async def commit_then_report_disconnect(self: AsyncSession):
            if self is env.db:
                ambiguous_commit_count["value"] += 1
            if self is env.db and ambiguous_commit_count["value"] == 2:
                await original_commit(self)
                raise RuntimeError("simulated disconnect after metadata commit")
            return await original_commit(self)

        monkeypatch.setattr(AsyncSession, "commit", commit_then_report_disconnect)
        with pytest.raises(RuntimeError, match="simulated disconnect after metadata commit"):
            await env.client.post(
                f"/api/projects/{ambiguous_project_id}/git/clone",
                json={"url": clone_url, "branch": "main"},
            )
        monkeypatch.setattr(AsyncSession, "commit", original_commit)
        ambiguous_repo = project_repo_path(env.tenant_id, uuid.UUID(ambiguous_project_id))
        assert (ambiguous_repo / "SOURCE.md").exists()
        ambiguous_settings = (await env.client.get(f"/api/projects/{ambiguous_project_id}/settings")).json()
        assert ambiguous_settings["git"]["source"] == "cloned"
        ambiguous_events = (await env.client.get(f"/api/projects/{ambiguous_project_id}/events?limit=200")).json()
        assert "git.repository.cloned" in {event["event_type"] for event in ambiguous_events}
        ambiguous_journal = (
            await env.db.execute(
                select(ProjectRepositoryOperation).where(
                    ProjectRepositoryOperation.project_id == uuid.UUID(ambiguous_project_id)
                )
            )
        ).scalar_one()
        assert ambiguous_journal.state == "committed"
        recovered_ambiguous_git = await env.client.get(f"/api/projects/{ambiguous_project_id}/git")
        assert recovered_ambiguous_git.status_code == 200
        assert (
            await env.db.execute(
                select(ProjectRepositoryOperation).where(
                    ProjectRepositoryOperation.project_id == uuid.UUID(ambiguous_project_id)
                )
            )
        ).scalars().all() == []

        original_finalize = projects_api.finalize_project_repository_clone

        async def simulate_exit_after_metadata_commit(_operation):
            assert _operation.lock_handle is not None
            _operation.lock_handle.close()  # Process exit releases the kernel flock.
            _operation.lock_handle = None
            raise RuntimeError("simulated exit before clone backup cleanup")

        monkeypatch.setattr(projects_api, "finalize_project_repository_clone", simulate_exit_after_metadata_commit)
        cloned = await env.client.post(
            f"/api/projects/{project_id}/git/clone",
            json={"url": clone_url, "branch": "main"},
        )
        assert cloned.status_code == 200, cloned.text
        assert cloned.json()["operation"] == "clone"
        assert cloned.json()["default_branch"] == "main"
        imported_repo = project_repo_path(env.tenant_id, uuid.UUID(project_id))
        assert (imported_repo / "SOURCE.md").read_text(encoding="utf-8") == "# Imported source\n"
        committed_journal = (await env.db.execute(select(ProjectRepositoryOperation))).scalar_one()
        assert committed_journal.state == "committed"
        assert list(imported_repo.parent.glob(".repo-backup-*"))
        monkeypatch.setattr(projects_api, "finalize_project_repository_clone", original_finalize)
        recovered_git = await env.client.get(f"/api/projects/{project_id}/git")
        assert recovered_git.status_code == 200, recovered_git.text
        assert recovered_git.json()["head"] == cloned.json()["head"]
        assert (await env.db.execute(select(ProjectRepositoryOperation))).scalars().all() == []
        assert list(imported_repo.parent.glob(".repo-backup-*")) == []
        settings = (await env.client.get(f"/api/projects/{project_id}/settings")).json()
        assert settings["git"]["mode"] == "managed"
        assert settings["git"]["repository_mode"] == "managed"
        assert settings["git"]["source"] == "cloned"
        assert settings["git"]["remotes"] == [
            {
                "name": "origin",
                "url_sha256": hashlib.sha256(clone_url.encode("utf-8")).hexdigest(),
            }
        ]
        assert settings["git"]["remote_count"] == 1
        second_clone = await env.client.post(
            f"/api/projects/{project_id}/git/clone",
            json={"url": clone_url, "branch": "main"},
        )
        assert second_clone.status_code == 409

        running_model = await env.db.get(Project, uuid.UUID(project_id))
        assert running_model is not None
        running_model.status = "running"
        await env.db.commit()
        running_clone = await env.client.post(
            f"/api/projects/{project_id}/git/clone",
            json={"url": clone_url, "branch": "main"},
        )
        assert running_clone.status_code == 409
        assert "planning" in running_clone.json()["detail"]
    finally:
        server.terminate()
        server.wait(timeout=5)

    events = (await env.client.get(f"/api/projects/{project_id}/events?limit=200")).json()
    assert {
        "git.remote.configured",
        "git.remote.deleted",
        "git.repository.cloned",
    } <= {event["event_type"] for event in events}

    configured_event = next(
        event
        for event in events
        if event["event_type"] == "git.remote.configured" and event["event_metadata"]["remote_name"] == "upstream"
    )
    assert configured_event["event_metadata"]["remote_name"] == "upstream"
    assert (
        configured_event["event_metadata"]["remote_url_sha256"]
        == hashlib.sha256(b"https://example.com/vendor/project.git").hexdigest()
    )
    cloned_event = next(event for event in events if event["event_type"] == "git.repository.cloned")
    assert cloned_event["event_metadata"]["remote_name"] == "origin"
    assert cloned_event["event_metadata"]["remote_url_sha256"] == hashlib.sha256(clone_url.encode("utf-8")).hexdigest()

    # Audit events are visible to explicitly shared viewers, so no Git event
    # may disclose a provider URL or the remote response object.
    shared = await env.client.patch(
        f"/api/projects/{project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert shared.status_code == 200, shared.text
    env.authenticate_as(env.viewer_id)
    viewer_events_response = await env.client.get(f"/api/projects/{project_id}/events?limit=200")
    assert viewer_events_response.status_code == 200, viewer_events_response.text
    viewer_git_events = [event for event in viewer_events_response.json() if event["event_type"].startswith("git.")]
    serialized_events = json.dumps(viewer_git_events, sort_keys=True)
    for private_url in (
        "https://example.com/vendor/project.git",
        "ssh://git@example.com/vendor/project.git",
        "git@example.com:vendor/project.git",
        clone_url,
    ):
        assert private_url not in serialized_events
    assert viewer_git_events
    assert all("remote" not in event["event_metadata"] for event in viewer_git_events)
    assert all("url" not in event["event_metadata"] for event in viewer_git_events)


async def test_concurrent_member_file_deliveries_keep_their_exact_parent_sessions(
    project_api: ProjectApiEnv,
):
    """Concurrent project members must not collapse onto another A2A thread."""
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectMemberSnapshot
    from app.models.subagent_run import SubagentRun
    from app.services.a2a_file_delivery import (
        append_a2a_file_delivery_message,
        resolve_a2a_file_origin_scope,
    )

    env = project_api
    project = await _create_project(env, name="Concurrent exact file routing")
    project_id = uuid.UUID(project["id"])
    senders = [(env.worker_id, "Worker"), (env.reviewer_id, "Reviewer")]
    routes: list[tuple[uuid.UUID, str, ChatSession, ChatSession]] = []

    for index, (sender_id, sender_name) in enumerate(senders):
        member = (
            await env.db.execute(
                select(ProjectMemberSnapshot).where(
                    ProjectMemberSnapshot.project_id == project_id,
                    ProjectMemberSnapshot.agent_id == sender_id,
                )
            )
        ).scalar_one()
        access_agent_id = min(sender_id, env.leader_id, key=str)
        peer_agent_id = max(sender_id, env.leader_id, key=str)
        decoy = ChatSession(
            project_id=project_id,
            agent_id=access_agent_id,
            peer_agent_id=peer_agent_id,
            source_channel="agent",
            title=f"{sender_name} old thread",
            external_conv_id=f"a2a-decoy-{index}",
        )
        parent = ChatSession(
            project_id=project_id,
            agent_id=access_agent_id,
            peer_agent_id=peer_agent_id,
            source_channel="agent",
            title=f"{sender_name} current thread",
            external_conv_id=f"a2a-current-{index}",
        )
        child = ChatSession(
            project_id=project_id,
            agent_id=sender_id,
            source_channel="subagent",
            title=f"{sender_name} project child",
            external_conv_id=f"subagent-file-{index}",
        )
        env.db.add_all([decoy, parent, child])
        await env.db.flush()
        env.db.add(
            SubagentRun(
                id=child.id,
                parent_session_id=parent.id,
                project_id=project_id,
                project_member_id=member.id,
                execution_user_id=env.owner_id,
                origin_tool_call_id=f"file-parent-{index}",
                mode="async",
                status="running",
            )
        )
        routes.append((sender_id, sender_name, parent, child))
    await env.db.commit()

    async def deliver(index: int, sender_id: uuid.UUID, sender_name: str, child: ChatSession) -> uuid.UUID:
        async with env.session_factory() as db:
            scope = await resolve_a2a_file_origin_scope(
                db,
                origin_session_id=child.id,
                sender_agent_id=sender_id,
            )
            session_id = await append_a2a_file_delivery_message(
                db,
                sender_agent_id=sender_id,
                target_agent_id=env.leader_id,
                sender_creator_id=env.owner_id,
                sender_name=sender_name,
                target_name="Leader",
                project_id=scope.project_id,
                preferred_session_id=scope.preferred_session_id,
                source_path=f"workspace/report-{index}.md",
                delivered_path=f"workspace/inbox/files/report-{index}.md",
                delivered_name=f"report-{index}.md",
                delivery_note="Review this exact project artifact",
                file_size=128 + index,
                created_at=datetime.now(UTC),
                external_event_key=f"test-project-file-{project_id}-{index}",
                origin_session_id=str(child.id),
                tool_call_id=f"send-file-{index}",
            )
            await db.commit()
            return session_id

    delivered_session_ids = await asyncio.gather(
        *(deliver(index, sender_id, sender_name, child) for index, (sender_id, sender_name, _parent, child) in enumerate(routes))
    )
    assert delivered_session_ids == [route[2].id for route in routes]

    messages = (
        (
            await env.db.execute(
                select(ChatMessage).where(
                    ChatMessage.external_event_key.in_(
                        [f"test-project-file-{project_id}-0", f"test-project-file-{project_id}-1"]
                    )
                )
            )
        )
        .scalars()
        .all()
    )
    assert {uuid.UUID(row.conversation_id) for row in messages} == {route[2].id for route in routes}
    assert {row.sender_agent_id for row in messages} == {env.worker_id, env.reviewer_id}
    assert all(len(row.message_meta["attachments"]) == 1 for row in messages)


async def test_activity_enum_includes_agent_file_delivery_actions():
    from app.models.activity_log import AgentActivityLog

    enum_values = set(AgentActivityLog.__table__.c.action_type.type.enums)
    assert {"agent_file_sent", "agent_file_received"} <= enum_values


async def test_project_skill_assets_are_owner_managed_and_template_portable(
    project_api: ProjectApiEnv,
    monkeypatch: pytest.MonkeyPatch,
):
    from app.models.mcp_server import MCPServer
    from app.models.project import ProjectCapabilityBinding, ProjectTemplate
    from app.models.skill import Skill, SkillFile
    from app.models.tool import AgentTool
    from app.services.project_agent_workspace import project_agent_workspace
    from app.services.storage import get_storage_backend, normalize_storage_key

    env = project_api
    source = await _create_project(env, name="Project Skill source")
    source_project_id = uuid.UUID(source["id"])
    worker_role = env.worker.role_description
    env.worker.autonomy_policy = {"write": "L2"}
    source_tool = (await env.db.execute(select(Tool).where(Tool.name == "send_message_to_parent"))).scalar_one()
    source_tool_id = source_tool.id
    source_tool_description = source_tool.description
    env.db.add(
        AgentTool(
            agent_id=env.source_worker_id,
            tool_id=source_tool_id,
            enabled=True,
            config={"api_key": "must-not-cross-project-boundary"},
            source="user_installed",
        )
    )
    mcp_server = MCPServer(
        tenant_id=env.tenant_id,
        name=f"private-runtime-{uuid.uuid4().hex[:8]}",
        display_name="Private Runtime MCP",
        base_url_template="https://example.invalid/mcp",
        headers_template={},
        instructions="Internal protocol instructions must not be product copy.",
    )
    shared_mcp_server = MCPServer(
        tenant_id=env.tenant_id,
        name=f"shared-runtime-{uuid.uuid4().hex[:8]}",
        display_name="Shared Runtime MCP",
        base_url_template="https://shared.example.invalid/mcp",
        headers_template={},
    )
    env.db.add_all([mcp_server, shared_mcp_server])
    await env.db.flush()
    source_mcp_tool = Tool(
        name=f"mcp_release_evidence_{uuid.uuid4().hex[:8]}",
        display_name="Release evidence query",
        description="Read release evidence from an approved connection.",
        type="mcp",
        category="engineering",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="agent",
        tenant_id=env.tenant_id,
        mcp_server_id=mcp_server.id,
    )
    shared_mcp_tool = Tool(
        name=f"shared_mcp_release_evidence_{uuid.uuid4().hex[:8]}",
        display_name="Shared release evidence query",
        description="Read release evidence from a shared connection.",
        type="mcp",
        category="engineering",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="admin",
        tenant_id=env.tenant_id,
        mcp_server_id=shared_mcp_server.id,
    )
    env.db.add_all([source_mcp_tool, shared_mcp_tool])
    await env.db.flush()
    source_mcp_assignment = AgentTool(
        agent_id=env.source_worker_id,
        tool_id=source_mcp_tool.id,
        enabled=True,
        config={"token": "must-not-cross-project-boundary"},
        source="user_installed",
    )
    env.db.add(source_mcp_assignment)
    await env.db.commit()

    bootstrap_response = await env.client.get("/api/projects/bootstrap-options")
    assert bootstrap_response.status_code == 200, bootstrap_response.text
    bootstrap_payload = bootstrap_response.json()
    bootstrap_capabilities = bootstrap_payload["capabilities"]
    bootstrap_tool = next(
        item for item in bootstrap_payload["tools"] if item["id"] == str(source_tool_id)
    )
    assert bootstrap_tool["name"] == source_tool.name
    assert bootstrap_tool["agent_config"] == {}
    bootstrap_mcp_tool = next(
        item
        for item in bootstrap_payload["tools"]
        if item["id"] == str(source_mcp_tool.id)
        and item["installed_by_agent_id"] == str(env.source_worker_id)
    )
    assert bootstrap_mcp_tool["mcp_server_name"] == mcp_server.display_name
    assert bootstrap_mcp_tool["agent_tool_source"] == "user_installed"
    assert bootstrap_mcp_tool["agent_config"] == {}
    assert bootstrap_mcp_tool["enabled"] is False
    assert all(item["type"] != "mcp" for item in bootstrap_capabilities)
    folder = f"release-check-{uuid.uuid4().hex[:8]}"
    source_prefix = normalize_storage_key(f"{env.source_worker_id}/skills/{folder}")
    storage = get_storage_backend()
    manifest_content = (
        "---\n"
        "name: Release Check\n"
        "version: 3\n"
        "description: Verify release evidence\n"
        "---\n\n"
        "# Release Check\n"
    )
    await storage.write_text(f"{source_prefix}/SKILL.md", manifest_content)
    await storage.write_text(f"{source_prefix}/references/checklist.md", "# Checklist\n")

    created_response = await env.client.post(
        f"/api/projects/{source_project_id}/agents",
        json={"source_agent_id": str(env.source_worker_id), "name": "Project release owner"},
    )
    assert created_response.status_code == 201, created_response.text
    project_agent = created_response.json()
    project_agent_id = uuid.UUID(project_agent["id"])

    explicit_tool_binding = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "tool",
            "capability_id": str(source_tool_id),
            "capability_name": source_tool.display_name,
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    explicit_mcp_binding = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "mcp",
            "capability_id": str(mcp_server.id),
            "capability_name": mcp_server.display_name,
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    shared_mcp_binding = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "mcp",
            "capability_id": str(shared_mcp_server.id),
            "capability_name": shared_mcp_server.display_name,
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    assert (
        explicit_tool_binding.status_code
        == explicit_mcp_binding.status_code
        == shared_mcp_binding.status_code
        == 201
    )

    capabilities_response = await env.client.get(f"/api/projects/{source_project_id}/capabilities")
    assert capabilities_response.status_code == 200, capabilities_response.text
    skill_binding = next(
        item
        for item in capabilities_response.json()
        if item["capability_type"] == "skill" and item["inherited_from_agent_id"] == str(project_agent_id)
    )
    tool_binding = next(
        item
        for item in capabilities_response.json()
        if item["capability_type"] == "tool" and item["inherited_from_agent_id"] == str(project_agent_id)
    )
    mcp_binding = next(
        item
        for item in capabilities_response.json()
        if item["capability_type"] == "mcp" and item["inherited_from_agent_id"] == str(project_agent_id)
    )
    assert tool_binding["capability_id"] == str(source_tool_id)
    assert tool_binding["key"] == source_tool.name
    assert tool_binding["availability"] == "available"
    assert tool_binding["description"] == source_tool_description
    assert tool_binding["config"] == {}
    copied_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == project_agent_id,
                AgentTool.tool_id == source_tool_id,
            )
        )
    ).scalar_one()
    assert copied_assignment.enabled is True
    assert copied_assignment.config == {}
    copied_mcp_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == project_agent_id,
                AgentTool.tool_id == source_mcp_tool.id,
            )
        )
    ).scalar_one()
    assert copied_mcp_assignment.enabled is True
    assert copied_mcp_assignment.config == {}
    disabled_tool_binding = await env.client.patch(
        f"/api/projects/{source_project_id}/capabilities/{tool_binding['id']}",
        json={"is_enabled": False},
    )
    disabled_mcp_binding = await env.client.patch(
        f"/api/projects/{source_project_id}/capabilities/{mcp_binding['id']}",
        json={"is_enabled": False},
    )
    assert disabled_tool_binding.status_code == disabled_mcp_binding.status_code == 200
    await env.db.refresh(copied_assignment)
    await env.db.refresh(copied_mcp_assignment)
    await env.db.refresh(source_mcp_assignment)
    assert copied_assignment.enabled is False
    assert copied_mcp_assignment.enabled is False
    assert source_mcp_assignment.enabled is True
    source_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == env.source_worker_id,
                AgentTool.tool_id == source_tool_id,
            )
        )
    ).scalar_one()
    assert source_assignment.enabled is True
    assert source_tool.enabled is True
    assert source_mcp_tool.enabled is True
    assert await env.db.get(MCPServer, mcp_server.id) is not None
    assert (
        await env.client.patch(
            f"/api/projects/{source_project_id}/capabilities/{tool_binding['id']}",
            json={"is_enabled": True},
        )
    ).status_code == 200
    assert (
        await env.client.patch(
            f"/api/projects/{source_project_id}/capabilities/{mcp_binding['id']}",
            json={"is_enabled": True},
        )
    ).status_code == 200
    source_tool_row = await env.db.get(Tool, source_tool_id)
    assert source_tool_row is not None
    source_tool_row.enabled = False
    await env.db.commit()
    restricted_capabilities = (
        await env.client.get(f"/api/projects/{source_project_id}/capabilities")
    ).json()
    assert next(item for item in restricted_capabilities if item["id"] == tool_binding["id"])[
        "availability"
    ] == "restricted"
    source_tool_row = await env.db.get(Tool, source_tool_id)
    assert source_tool_row is not None
    source_tool_row.enabled = True
    await env.db.commit()
    stored_source_binding = await env.db.get(ProjectCapabilityBinding, uuid.UUID(skill_binding["id"]))
    assert stored_source_binding is not None
    metadata = stored_source_binding.config["skill_asset"]
    assert metadata == {
        "schema_version": 2,
        "asset_id": metadata["asset_id"],
        "version": "3",
        "sha256": metadata["sha256"],
        "path": f"skills/{folder}",
        "source": "agent",
        "source_agent_id": str(env.source_worker_id),
        "file_count": 2,
        "size_bytes": len(manifest_content.encode()) + len("# Checklist\n".encode()),
    }
    assert skill_binding["config"] == {}
    assert skill_binding["availability"] == "available"
    assert skill_binding["description"] == "Verify release evidence"
    assert skill_binding["version"] == "3"
    assert skill_binding["file_count"] == 2
    assert skill_binding["size_bytes"] == metadata["size_bytes"]
    source_project = await env.db.get(Project, source_project_id)
    assert source_project is not None
    source_layout = project_agent_workspace(
        project_repo_path(source_project.tenant_id, source_project.id),
        project_agent_id,
    )
    copied_manifest = source_layout.root / "skills" / folder / "SKILL.md"
    assert copied_manifest.read_text(encoding="utf-8") == manifest_content

    library_folder = f"library-check-{uuid.uuid4().hex[:8]}"
    library_skill = Skill(
        tenant_id=env.tenant_id,
        name="Library Check",
        description="Validate a library-backed release check",
        category="engineering",
        folder_name=library_folder,
        version=5,
        visibility="tenant",
        status="published",
        publisher_agent_id=env.source_worker_id,
    )
    env.db.add(library_skill)
    await env.db.flush()
    library_skill_id = library_skill.id
    env.db.add_all(
        [
            SkillFile(
                skill_id=library_skill_id,
                path="SKILL.md",
                content="---\nname: Library Check\ndescription: Validate a library check\n---\n",
            ),
            SkillFile(skill_id=library_skill_id, path="references/guide.md", content="# Guide\n"),
        ]
    )
    await env.db.commit()
    rejected_shared_skill = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={"capability_type": "skill", "capability_id": str(library_skill_id), "source": "shared"},
    )
    assert rejected_shared_skill.status_code == 422
    library_binding_response = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "skill",
            "capability_id": str(library_skill_id),
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    assert library_binding_response.status_code == 201, library_binding_response.text
    library_binding = library_binding_response.json()
    assert library_binding["availability"] == "available"
    assert library_binding["version"] == "5"
    assert library_binding["file_count"] == 2
    assert library_binding["description"] == "Validate a library-backed release check"
    assert library_binding["config"] == {}
    library_root = source_layout.root / "skills" / library_folder
    assert (library_root / "SKILL.md").is_file()
    hidden_library_root = library_root.with_name(f".{library_root.name}.missing")
    os.replace(library_root, hidden_library_root)
    missing_capabilities = (await env.client.get(f"/api/projects/{source_project_id}/capabilities")).json()
    assert next(item for item in missing_capabilities if item["id"] == library_binding["id"])[
        "availability"
    ] == "missing"
    os.replace(hidden_library_root, library_root)
    duplicate_library_binding = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "skill",
            "capability_id": str(library_skill_id),
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    assert duplicate_library_binding.status_code == 409
    secret_folder = f"secret-skill-{uuid.uuid4().hex[:8]}"
    secret_skill = Skill(
        tenant_id=env.tenant_id,
        name="Unsafe Skill",
        description="Must not cross the project boundary",
        category="engineering",
        folder_name=secret_folder,
        version=1,
        visibility="tenant",
        status="published",
    )
    env.db.add(secret_skill)
    await env.db.flush()
    secret_skill_id = secret_skill.id
    env.db.add(
        SkillFile(
            skill_id=secret_skill_id,
            path="SKILL.md",
            content="---\nname: Unsafe Skill\n---\napi_key = 'abcdefghijklmnop123456'\n",
        )
    )
    await env.db.commit()
    rejected_secret = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "skill",
            "capability_id": str(secret_skill_id),
            "source": "inherited",
            "inherited_from_agent_id": str(project_agent_id),
        },
    )
    assert rejected_secret.status_code == 422
    assert not (source_layout.root / "skills" / secret_folder).exists()

    manifest_response = await env.client.get(f"/api/projects/{source_project_id}/template-manifest")
    assert manifest_response.status_code == 200, manifest_response.text
    manifest_skill = next(
        item for item in manifest_response.json()["skills"] if item["binding_id"] == skill_binding["id"]
    )
    assert manifest_skill["selected"] is False
    assert manifest_skill["selection_state"] == "unselected"
    assert manifest_skill["affected_member_count"] == 1
    assert {"path", "sha256", "asset_id"}.isdisjoint(manifest_skill)
    manifest_tool = next(
        item
        for item in manifest_response.json()["capabilities"]
        if item["type"] == "tool" and item["key"] == source_tool.name
    )
    assert manifest_tool["selected"] is True
    assert manifest_tool["affected_member_count"] == 1
    assert {
        key: manifest_skill[key]
        for key in (
            "binding_id",
            "member_id",
            "member_agent_id",
            "member_name",
            "member_role",
            "name",
            "version",
            "file_count",
            "size_bytes",
        )
    } == {
        "binding_id": skill_binding["id"],
        "member_id": project_agent["member_id"],
        "member_agent_id": str(project_agent_id),
        "member_name": "Project release owner",
        "member_role": worker_role,
        "name": "Release Check",
        "version": "3",
        "file_count": 2,
        "size_bytes": metadata["size_bytes"],
    }

    shared = await env.client.patch(
        f"/api/projects/{source_project_id}",
        json={
            "visibility": "shared",
            "shared_with_user_ids": [str(env.viewer_id)],
            "execution_user_id": str(env.viewer_id),
        },
    )
    assert shared.status_code == 200, shared.text
    grant = (
        await env.db.execute(
            select(ProjectAccessGrant).where(
                ProjectAccessGrant.project_id == source_project_id,
                ProjectAccessGrant.user_id == env.viewer_id,
            )
        )
    ).scalar_one()
    grant.role = "edit"
    await env.db.commit()
    env.authenticate_as(env.viewer_id)
    assert (await env.client.get(f"/api/projects/{source_project_id}/capabilities")).status_code == 200
    denied_create = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={"capability_type": "tool", "capability_name": "editor-tool"},
    )
    assert denied_create.status_code == 404
    denied_patch = await env.client.patch(
        f"/api/projects/{source_project_id}/capabilities/{skill_binding['id']}",
        json={"is_enabled": False},
    )
    assert denied_patch.status_code == 404
    env.authenticate_as(env.owner_id)

    default_template_response = await env.client.post(
        f"/api/projects/{source_project_id}/templates",
        json={"name": "No Skills by default", "version": "1.0.0"},
    )
    assert default_template_response.status_code == 201, default_template_response.text
    default_template = await env.db.get(ProjectTemplate, uuid.UUID(default_template_response.json()["id"]))
    assert default_template is not None
    assert default_template.definition["skill_assets"] == []
    assert default_template_response.json()["skills"] == []

    selected_template_response = await env.client.post(
        f"/api/projects/{source_project_id}/templates",
        json={
            "name": "Release Skill template",
            "version": "1.0.0",
            "is_published": True,
            "included_skill_binding_ids": [skill_binding["id"]],
        },
    )
    assert selected_template_response.status_code == 201, selected_template_response.text
    selected_template_body = selected_template_response.json()
    assert selected_template_body["skills"] == [
        {
            "name": "Release Check",
            "version": "3",
            "member_name": "Project release owner",
            "member_role": worker_role,
            "file_count": 2,
            "size_bytes": metadata["size_bytes"],
        }
    ]
    selected_template = await env.db.get(ProjectTemplate, uuid.UUID(selected_template_body["id"]))
    assert selected_template is not None
    packaged_skill = selected_template.definition["skill_assets"][0]
    assert packaged_skill["sha256"] == metadata["sha256"]
    assert packaged_skill["digital_employee_index"] == 3
    assert skill_binding["id"] not in str(packaged_skill)
    assert str(env.source_worker_id) not in str(packaged_skill)
    assert "capability_id" not in packaged_skill
    packaged_tool = next(
        item for item in selected_template.definition["capabilities"] if item["capability_type"] == "tool"
    )
    packaged_mcp = [
        item for item in selected_template.definition["capabilities"] if item["capability_type"] == "mcp"
    ]
    assert packaged_tool["capability_id"] == str(source_tool_id)
    assert [item["capability_id"] for item in packaged_mcp] == [str(shared_mcp_server.id)]
    assert packaged_mcp[0]["source"] == "inherited"
    assert packaged_mcp[0]["digital_employee_index"] == 3
    assert str(mcp_server.id) not in str(selected_template.definition["capabilities"])
    assert "config" not in packaged_tool
    assert "must-not-cross-project-boundary" not in str(selected_template.definition)

    # Older published snapshots may still contain registry identifiers that
    # were portable before the canonical authorization boundary existed.
    # They must not become authority when another user instantiates them.
    selected_template.definition = {
        **selected_template.definition,
        "capabilities": [
            *selected_template.definition["capabilities"],
            {
                "schema_version": 1,
                "capability_type": "mcp",
                "capability_id": str(mcp_server.id),
                "capability_name": mcp_server.display_name,
                "source": "inherited",
                "digital_employee_index": 3,
                "is_enabled": True,
                "scope": {},
            },
                {
                    "schema_version": 1,
                    "capability_type": "skill",
                    "capability_id": str(library_skill_id),
                    "capability_name": "Library Check",
                "source": "inherited",
                "digital_employee_index": 3,
                "is_enabled": True,
                "scope": {},
            },
        ],
    }
    await env.db.commit()

    env.authenticate_as(env.viewer_id)
    restored_response = await env.client.post(
        "/api/projects/from-template",
        json={"template_id": selected_template_body["id"], "name": "Restored Skill project"},
    )
    assert restored_response.status_code == 201, restored_response.text
    restored = restored_response.json()
    restored_project_id = uuid.UUID(restored["id"])
    assert restored["template_setup_summary"]["restored_skill_count"] == 1
    restored_agents = (await env.client.get(f"/api/projects/{restored_project_id}/agents")).json()
    assert len(restored_agents) == 4
    restored_agent_id = uuid.UUID(
        next(item for item in restored_agents if item["name"] == "Project release owner")["id"]
    )
    assert restored_agent_id != project_agent_id
    restored_binding = (
        await env.db.execute(
            select(ProjectCapabilityBinding).where(
                ProjectCapabilityBinding.project_id == restored_project_id,
                ProjectCapabilityBinding.capability_type == "skill",
            )
        )
    ).scalar_one()
    assert restored_binding.capability_id is None
    assert restored_binding.inherited_from_agent_id == restored_agent_id
    assert restored_binding.config["skill_asset"]["source"] == "template"
    assert restored_binding.config["skill_asset"]["source_agent_id"] is None
    restored_tool_bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == restored_project_id,
                    ProjectCapabilityBinding.capability_type == "tool",
                )
            )
        ).scalars()
    )
    restored_tool_binding = next(
        item for item in restored_tool_bindings if item.inherited_from_agent_id == restored_agent_id
    )
    assert restored_tool_binding.inherited_from_agent_id == restored_agent_id
    restored_assignment = (
        await env.db.execute(
            select(AgentTool).where(
                AgentTool.agent_id == restored_agent_id,
                AgentTool.tool_id == source_tool_id,
            )
        )
    ).scalar_one()
    assert restored_assignment.enabled is True
    assert restored_assignment.config == {}
    restored_mcp_bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == restored_project_id,
                    ProjectCapabilityBinding.capability_type == "mcp",
                )
            )
        ).scalars()
    )
    assert [(item.capability_id, item.source) for item in restored_mcp_bindings] == [
        (shared_mcp_server.id, "inherited")
    ]
    assert restored_mcp_bindings[0].inherited_from_agent_id == restored_agent_id
    restored_project = await env.db.get(Project, restored_project_id)
    assert restored_project is not None
    restored_layout = project_agent_workspace(
        project_repo_path(restored_project.tenant_id, restored_project.id),
        restored_agent_id,
    )
    assert (restored_layout.root / "skills" / folder / "SKILL.md").read_text(encoding="utf-8") == manifest_content

    env.authenticate_as(env.owner_id)
    deactivated = await env.client.post(f"/api/projects/{source_project_id}/agents/{project_agent_id}/deactivate")
    assert deactivated.status_code == 200, deactivated.text
    assert copied_manifest.read_text(encoding="utf-8") == manifest_content
    retained_binding = await env.db.get(ProjectCapabilityBinding, uuid.UUID(skill_binding["id"]))
    assert retained_binding is not None
    retained_capabilities = (await env.client.get(f"/api/projects/{source_project_id}/capabilities")).json()
    assert next(item for item in retained_capabilities if item["id"] == skill_binding["id"])[
        "availability"
    ] == "available"

    restored_member = await env.client.post(
        f"/api/projects/{source_project_id}/agents/{project_agent_id}/restore"
    )
    assert restored_member.status_code == 200, restored_member.text
    second_agent_response = await env.client.post(
        f"/api/projects/{source_project_id}/agents",
        json={"source_agent_id": str(env.source_worker_id), "name": "Second release owner"},
    )
    assert second_agent_response.status_code == 201, second_agent_response.text
    second_agent_id = uuid.UUID(second_agent_response.json()["id"])
    second_layout = project_agent_workspace(
        project_repo_path(source_project.tenant_id, source_project.id),
        second_agent_id,
    )
    second_source_binding = (
        await env.db.execute(
            select(ProjectCapabilityBinding).where(
                ProjectCapabilityBinding.project_id == source_project_id,
                ProjectCapabilityBinding.capability_type == "skill",
                ProjectCapabilityBinding.inherited_from_agent_id == second_agent_id,
                ProjectCapabilityBinding.capability_name == "Release Check",
            )
        )
    ).scalar_one()
    assert second_source_binding.config["skill_asset"]["asset_id"] == metadata["asset_id"]
    assert os.stat(copied_manifest).st_ino == os.stat(
        second_layout.root / "skills" / folder / "SKILL.md"
    ).st_ino

    second_library_response = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "skill",
            "capability_id": str(library_skill_id),
            "source": "inherited",
            "inherited_from_agent_id": str(second_agent_id),
        },
    )
    assert second_library_response.status_code == 201, second_library_response.text
    second_library_binding = second_library_response.json()
    first_library_row = await env.db.get(ProjectCapabilityBinding, uuid.UUID(library_binding["id"]))
    second_library_row = await env.db.get(ProjectCapabilityBinding, uuid.UUID(second_library_binding["id"]))
    assert first_library_row is not None and second_library_row is not None
    assert (
        first_library_row.config["skill_asset"]["asset_id"]
        == second_library_row.config["skill_asset"]["asset_id"]
    )
    second_library_root = second_layout.root / "skills" / library_folder
    assert os.stat(library_root / "SKILL.md").st_ino == os.stat(second_library_root / "SKILL.md").st_ino

    disabled_response = await env.client.patch(
        f"/api/projects/{source_project_id}/capabilities/{second_library_binding['id']}",
        json={"is_enabled": False},
    )
    assert disabled_response.status_code == 200, disabled_response.text
    assert disabled_response.json()["is_enabled"] is False
    assert disabled_response.json()["availability"] == "available"
    assert not second_library_root.exists()
    source_library_after_disable = await env.db.get(Skill, library_skill_id)
    assert source_library_after_disable is not None
    assert source_library_after_disable.status == "published"
    assert len(
        list(
            (
                await env.db.execute(select(SkillFile).where(SkillFile.skill_id == library_skill_id))
            ).scalars()
        )
    ) == 2
    enabled_response = await env.client.patch(
        f"/api/projects/{source_project_id}/capabilities/{second_library_binding['id']}",
        json={"is_enabled": True},
    )
    assert enabled_response.status_code == 200, enabled_response.text
    assert second_library_root.is_dir()

    library_skill = await env.db.get(Skill, library_skill_id)
    assert library_skill is not None
    library_skill.version = 6
    library_manifest_row = (
        await env.db.execute(
            select(SkillFile).where(
                SkillFile.skill_id == library_skill_id,
                SkillFile.path == "SKILL.md",
            )
        )
    ).scalar_one()
    refreshed_library_content = (
        "---\nname: Library Check\nversion: 6\ndescription: Updated library check\n---\n"
    )
    library_manifest_row.content = refreshed_library_content
    await env.db.commit()
    refresh_response = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities/{library_binding['id']}/refresh"
    )
    assert refresh_response.status_code == 200, refresh_response.text
    assert refresh_response.json()["changed"] is True
    assert refresh_response.json()["affected_member_count"] == 2
    assert refresh_response.json()["capability"]["version"] == "6"
    assert (library_root / "SKILL.md").read_text(encoding="utf-8") == refreshed_library_content
    assert (second_library_root / "SKILL.md").read_text(encoding="utf-8") == refreshed_library_content
    assert os.stat(library_root / "SKILL.md").st_ino == os.stat(second_library_root / "SKILL.md").st_ino

    impact_response = await env.client.get(
        f"/api/projects/{source_project_id}/capabilities/{library_binding['id']}/delete-impact"
    )
    assert impact_response.status_code == 200, impact_response.text
    assert impact_response.json()["affected_member_count"] == 2
    rejected_delete = await env.client.delete(
        f"/api/projects/{source_project_id}/capabilities/{library_binding['id']}"
    )
    assert rejected_delete.status_code == 409
    assert rejected_delete.json()["detail"]["affected_member_count"] == 2
    confirmed_delete = await env.client.delete(
        f"/api/projects/{source_project_id}/capabilities/{library_binding['id']}?confirm=true"
    )
    assert confirmed_delete.status_code == 200, confirmed_delete.text
    assert confirmed_delete.json()["affected_member_count"] == 2
    remaining_library_bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == source_project_id,
                    ProjectCapabilityBinding.capability_id == library_skill_id,
                )
            )
        ).scalars()
    )
    assert remaining_library_bindings == []
    assert not library_root.exists()
    assert not second_library_root.exists()

    workspace_manifest_content = (
        "---\nname: Release Check\nversion: 4\ndescription: Edited in the project Agent Skill page\n---\n"
    )
    workspace_write = await env.client.put(
        f"/api/agents/{project_agent_id}/files/content",
        params={"path": f"skills/{folder}/SKILL.md"},
        json={"content": workspace_manifest_content},
    )
    assert workspace_write.status_code == 200, workspace_write.text
    assert copied_manifest.read_text(encoding="utf-8") == workspace_manifest_content
    workspace_capabilities = (await env.client.get(f"/api/projects/{source_project_id}/capabilities")).json()
    edited_skill = next(item for item in workspace_capabilities if item["id"] == skill_binding["id"])
    assert edited_skill["version"] == "4"
    assert edited_skill["description"] == "Edited in the project Agent Skill page"
    edited_manifest = (await env.client.get(f"/api/projects/{source_project_id}/template-manifest")).json()
    edited_manifest_skill = next(
        item for item in edited_manifest["skills"] if item["binding_id"] == skill_binding["id"]
    )
    assert edited_manifest_skill["version"] == "4"
    assert {"path", "sha256", "asset_id"}.isdisjoint(edited_manifest_skill)

    workspace_delete = await env.client.delete(
        f"/api/agents/{project_agent_id}/files/content",
        params={"path": f"skills/{folder}/SKILL.md"},
    )
    assert workspace_delete.status_code == 200, workspace_delete.text
    assert workspace_delete.json()["project_skill_deleted"] is True
    assert workspace_delete.json()["affected_member_count"] == 2
    remaining_source_skill_bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == source_project_id,
                    ProjectCapabilityBinding.capability_type == "skill",
                    ProjectCapabilityBinding.capability_name == "Release Check",
                )
            )
        ).scalars()
    )
    assert remaining_source_skill_bindings == []
    assert not copied_manifest.parent.exists()
    assert not (second_layout.root / "skills" / folder).exists()
    deletion_events = (await env.client.get(f"/api/projects/{source_project_id}/events?limit=200")).json()
    file_delete_event = next(
        event
        for event in reversed(deletion_events)
        if event["event_type"] == "capability.deleted"
        and event["event_metadata"].get("source") == "agent_files"
    )
    assert file_delete_event["event_metadata"]["affected_member_count"] == 2

    def fail_template_skill_copy(*_args, **_kwargs):
        raise OSError("simulated template Skill copy failure")

    monkeypatch.setattr(
        "app.services.project_skill_assets._write_files",
        fail_template_skill_copy,
    )
    failed_project_name = f"Failed Skill restore {uuid.uuid4().hex[:8]}"
    with pytest.raises(OSError, match="simulated template Skill copy failure"):
        await env.client.post(
            "/api/projects/from-template",
            json={"template_id": selected_template_body["id"], "name": failed_project_name},
        )
    failed_project = (
        await env.db.execute(select(Project).where(Project.name == failed_project_name))
    ).scalar_one_or_none()
    assert failed_project is None


async def test_project_agent_template_api_round_trip_preserves_assets_with_fresh_identity(
    project_api: ProjectApiEnv,
):
    from app.models.chat_session import ChatSession
    from app.models.project import ProjectCapabilityBinding, ProjectMemberSnapshot, ProjectTemplate

    env = project_api
    source = await _create_project(env, name="Template source")
    source_project_id = uuid.UUID(source["id"])

    created_agent_response = await env.client.post(
        f"/api/projects/{source_project_id}/agents",
        json={
            "name": "Project release specialist",
            "role_description": "Own release readiness inside this project",
            "soul": "# Release specialist\nKeep delivery evidence concise.\n",
            "core_memory": "# Durable context\nThe acceptance gate requires a signed release checklist.\n",
        },
    )
    assert created_agent_response.status_code == 201, created_agent_response.text
    source_agent = created_agent_response.json()

    portable_tool = (
        await env.db.execute(select(Tool).where(Tool.name == "send_message_to_parent"))
    ).scalar_one()
    inherited_tool_response = await env.client.post(
        f"/api/projects/{source_project_id}/capabilities",
        json={
            "capability_type": "tool",
            "capability_id": str(portable_tool.id),
            "source": "inherited",
            "inherited_from_agent_id": source_agent["id"],
        },
    )
    assert inherited_tool_response.status_code == 201, inherited_tool_response.text
    deactivated_source_agent = await env.client.post(
        f"/api/projects/{source_project_id}/agents/{source_agent['id']}/deactivate"
    )
    assert deactivated_source_agent.status_code == 200, deactivated_source_agent.text

    source_project = await env.db.get(Project, source_project_id)
    assert source_project is not None
    source_repo = project_repo_path(source_project.tenant_id, source_project.id)
    (source_repo / ".gitignore").write_text("frontend/dist/\n", encoding="utf-8")
    ignored_asset = source_repo / "frontend" / "dist" / "app.js"
    ignored_asset.parent.mkdir(parents=True, exist_ok=True)
    ignored_asset.write_text("console.log('portable build');\n", encoding="utf-8")
    subprocess.run(
        ["git", "-C", str(source_repo), "add", ".gitignore"],
        check=True,
        capture_output=True,
        text=True,
    )
    subprocess.run(
        ["git", "-C", str(source_repo), "add", "-f", "frontend/dist/app.js"],
        check=True,
        capture_output=True,
        text=True,
    )
    subprocess.run(
        [
            "git",
            "-C",
            str(source_repo),
            "-c",
            "user.name=Project test",
            "-c",
            "user.email=project@test.invalid",
            "commit",
            "-m",
            "Add portable ignored build asset",
        ],
        check=True,
        capture_output=True,
        text=True,
    )

    env.authenticate_as(env.viewer_id)
    forbidden_template_response = await env.client.post(
        f"/api/projects/{source_project_id}/templates",
        json={"name": "Unauthorized export"},
    )
    assert forbidden_template_response.status_code == 404
    env.authenticate_as(env.owner_id)

    template_response = await env.client.post(
        f"/api/projects/{source_project_id}/templates",
        json={
            "name": "Release readiness template",
            "description": "Reusable release workflow",
            "category": "engineering",
            "version": "1.0.0",
        },
    )
    assert template_response.status_code == 201, template_response.text
    template = template_response.json()
    assert "agents" not in template["definition"]
    assert len(template["definition"]["roles"]) == 4
    release_role = next(
        item for item in template["definition"]["roles"] if item["name"] == "Project release specialist"
    )
    assert release_role["description"] == "Own release readiness inside this project"
    assert template["definition"]["asset_summary"]["digital_employee_count"] == 4

    stored_template = await env.db.get(ProjectTemplate, uuid.UUID(template["id"]))
    assert stored_template is not None
    assert "frontend/dist/app.js" in {
        item["path"] for item in stored_template.definition["project_snapshot"]["files"]
    }
    exported_agents = stored_template.definition["agents"]
    assert len(exported_agents) == 4
    exported_agent = next(item for item in exported_agents if item["name"] == "Project release specialist")
    assert {
        key: exported_agent[key]
        for key in (
            "name",
            "role_description",
            "is_leader",
            "is_enabled",
            "soul",
            "core_memory",
            "workspace_files",
        )
    } == {
        "name": "Project release specialist",
        "role_description": "Own release readiness inside this project",
        "is_leader": False,
        "is_enabled": False,
        "soul": "# Release specialist\nKeep delivery evidence concise.\n",
        "core_memory": "# Durable context\nThe acceptance gate requires a signed release checklist.\n",
        "workspace_files": [],
    }
    assert exported_agent["runtime"]["max_tool_rounds"] > 0
    assert exported_agent["member_config"]["enabled_project_tools"] == []
    assert source_agent["id"] not in str(stored_template.definition)
    assert str(source_project_id) not in str(stored_template.definition)

    target_response = await env.client.post(
        "/api/projects/from-template",
        json={
            "template_id": template["id"],
            "name": "Template target",
            "visibility": "private",
        },
    )
    assert target_response.status_code == 201, target_response.text
    target = target_response.json()
    target_project_id = uuid.UUID(target["id"])
    assert target_project_id != source_project_id
    assert target["goal"] == source["goal"]
    assert target["success_criteria"] == source["success_criteria"]
    assert target["template_setup_summary"] == {
        "restored_file_count": len(stored_template.definition["project_snapshot"]["files"]),
        "restored_digital_employee_count": 4,
        "restored_skill_count": 0,
        "restored_connection_count": 0,
        "restored_tool_count": 1,
    }
    target_project = await env.db.get(Project, target_project_id)
    assert target_project is not None
    target_repo = project_repo_path(target_project.tenant_id, target_project.id)
    assert (target_repo / "frontend" / "dist" / "app.js").read_text(
        encoding="utf-8"
    ) == "console.log('portable build');\n"
    assert (
        subprocess.run(
            ["git", "-C", str(target_repo), "ls-files", "--error-unmatch", "frontend/dist/app.js"],
            check=False,
            capture_output=True,
            text=True,
        ).returncode
        == 0
    )

    target_agents_response = await env.client.get(f"/api/projects/{target_project_id}/agents")
    assert target_agents_response.status_code == 200, target_agents_response.text
    target_agents = target_agents_response.json()
    assert len(target_agents) == 4
    target_agent = next(item for item in target_agents if item["name"] == source_agent["name"])
    assert target_agent["id"] != source_agent["id"]
    assert target_agent["name"] == source_agent["name"]
    assert target_agent["soul"] == source_agent["soul"]
    assert target_agent["core_memory"] == source_agent["core_memory"]
    assert target_agent["is_leader"] is False

    member = (
        await env.db.execute(
            select(ProjectMemberSnapshot).where(
                ProjectMemberSnapshot.project_id == target_project_id,
                ProjectMemberSnapshot.agent_id == uuid.UUID(target_agent["id"]),
            )
        )
    ).scalar_one()
    assert member.is_enabled is False
    assert member.is_leader is False
    restored_departed_capability = (
        await env.db.execute(
            select(ProjectCapabilityBinding).where(
                ProjectCapabilityBinding.project_id == target_project_id,
                ProjectCapabilityBinding.inherited_from_agent_id == uuid.UUID(target_agent["id"]),
                ProjectCapabilityBinding.capability_id == portable_tool.id,
            )
        )
    ).scalar_one()
    assert restored_departed_capability.is_enabled is False

    sessions = (
        (await env.db.execute(select(ChatSession).where(ChatSession.project_id == target_project_id))).scalars().all()
    )
    assert {(session.source_channel, session.is_group) for session in sessions} == {
        ("project", True),
        ("web", False),
    }
    target_leader = next(item for item in target_agents if item["is_leader"])
    assert {session.agent_id for session in sessions} == {uuid.UUID(target_leader["id"])}


async def test_project_skill_legacy_backfill_dry_run_apply_and_rollback_are_auditable(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectCapabilityBinding, ProjectEvent
    from app.services.project_agent_workspace import project_agent_workspace

    env = project_api
    source = await _create_project(env, name="Legacy Skill normalization")
    project_id = uuid.UUID(source["id"])
    created_agent_response = await env.client.post(
        f"/api/projects/{project_id}/agents",
        json={"name": "Legacy Skill owner", "role_description": "Maintain project knowledge"},
    )
    assert created_agent_response.status_code == 201, created_agent_response.text
    project_agent_id = uuid.UUID(created_agent_response.json()["id"])
    manifests = {
        "Legacy Zero": "---\nname: Legacy Zero\nversion: 7\n---\n\n# Zero\n",
        "Legacy One": "---\nname: Legacy One\nversion: 8\n---\n\n# One\n",
    }
    folders = {"Legacy Zero": "legacy-zero", "Legacy One": "legacy-one"}
    for name, content in manifests.items():
        response = await env.client.put(
            f"/api/agents/{project_agent_id}/files/content",
            params={"path": f"skills/{folders[name]}/SKILL.md"},
            json={"content": content},
        )
        assert response.status_code == 200, response.text

    bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == project_id,
                    ProjectCapabilityBinding.capability_type == "skill",
                )
            )
        ).scalars()
    )
    by_name = {binding.capability_name: binding for binding in bindings}
    zero = by_name["Legacy Zero"]
    one = by_name["Legacy One"]
    disabled_one = await env.client.patch(
        f"/api/projects/{project_id}/capabilities/{one.id}",
        json={"is_enabled": False},
    )
    assert disabled_one.status_code == 200, disabled_one.text
    await env.db.refresh(one)
    current_one_metadata = dict(one.config["skill_asset"])
    v1_metadata = {key: value for key, value in current_one_metadata.items() if key != "asset_id"}
    v1_metadata["schema_version"] = 1
    project = await env.db.get(Project, project_id)
    assert project is not None
    layout = project_agent_workspace(project_repo_path(project.tenant_id, project.id), project_agent_id)
    current_disabled_root = (
        layout.root
        / ".disabled-skills"
        / current_one_metadata["asset_id"]
        / folders["Legacy One"]
    )
    legacy_disabled_root = (
        layout.root
        / ".disabled-skills"
        / current_one_metadata["sha256"]
        / folders["Legacy One"]
    )
    legacy_disabled_root.parent.mkdir(parents=True, exist_ok=True)
    os.replace(current_disabled_root, legacy_disabled_root)
    zero.config = {"legacy_note": "keep-zero"}
    one.config = {"legacy_note": "keep-one", "skill_asset": v1_metadata}
    await env.db.commit()

    env.authenticate_as(env.viewer_id)
    hidden = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "dry_run"},
    )
    assert hidden.status_code == 404
    env.authenticate_as(env.owner_id)

    preview = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "dry_run"},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["operation_id"] is None
    assert preview.json()["ready_count"] == 2
    assert preview.json()["blocked_count"] == 0
    assert {item["previous_schema_version"] for item in preview.json()["items"]} == {0, 1}
    await env.db.refresh(zero)
    await env.db.refresh(one)
    assert zero.config == {"legacy_note": "keep-zero"}
    assert one.config == {"legacy_note": "keep-one", "skill_asset": v1_metadata}
    dry_run_events = list(
        (
            await env.db.execute(
                select(ProjectEvent).where(
                    ProjectEvent.project_id == project_id,
                    ProjectEvent.event_type.like("capability.skill_backfill.%"),
                )
            )
        ).scalars()
    )
    assert dry_run_events == []

    applied = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "apply"},
    )
    assert applied.status_code == 200, applied.text
    applied_body = applied.json()
    assert applied_body["applied_count"] == 2
    operation_id = uuid.UUID(applied_body["operation_id"])
    await env.db.refresh(zero)
    await env.db.refresh(one)
    for binding, note in ((zero, "keep-zero"), (one, "keep-one")):
        assert binding.config["legacy_note"] == note
        assert binding.config["skill_asset"]["schema_version"] == 2
        uuid.UUID(binding.config["skill_asset"]["asset_id"])

    assert (layout.root / "skills" / folders["Legacy Zero"] / "SKILL.md").read_text(
        encoding="utf-8"
    ) == manifests["Legacy Zero"]
    applied_one_metadata = one.config["skill_asset"]
    assert (
        layout.root
        / ".disabled-skills"
        / applied_one_metadata["asset_id"]
        / folders["Legacy One"]
        / "SKILL.md"
    ).read_text(encoding="utf-8") == manifests["Legacy One"]
    apply_event = await env.db.get(ProjectEvent, operation_id)
    assert apply_event is not None
    assert apply_event.event_type == "capability.skill_backfill.applied"
    assert apply_event.actor_user_id == env.owner_id
    assert apply_event.event_metadata["applied_count"] == 2
    assert len(apply_event.event_metadata["rollback_entries"]) == 2
    assert all(
        {"path", "sha256", "config"}.isdisjoint(entry)
        for entry in apply_event.event_metadata["rollback_entries"]
    )

    no_op = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "apply"},
    )
    assert no_op.status_code == 200, no_op.text
    assert no_op.json()["applied_count"] == 0
    assert no_op.json()["operation_id"] is None

    rolled_back = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "rollback", "operation_id": str(operation_id)},
    )
    assert rolled_back.status_code == 200, rolled_back.text
    assert rolled_back.json()["rolled_back_count"] == 2
    await env.db.refresh(zero)
    await env.db.refresh(one)
    assert zero.config == {"legacy_note": "keep-zero"}
    assert one.config == {"legacy_note": "keep-one", "skill_asset": v1_metadata}
    assert (legacy_disabled_root / "SKILL.md").read_text(encoding="utf-8") == manifests["Legacy One"]
    rollback_event = await env.db.get(ProjectEvent, uuid.UUID(rolled_back.json()["rollback_event_id"]))
    assert rollback_event is not None
    assert rollback_event.event_metadata["operation_id"] == str(operation_id)
    assert rollback_event.actor_user_id == env.owner_id
    duplicate_rollback = await env.client.post(
        f"/api/projects/{project_id}/capabilities/skill-backfill",
        json={"action": "rollback", "operation_id": str(operation_id)},
    )
    assert duplicate_rollback.status_code == 409

    missing_root = layout.root / "skills" / folders["Legacy Zero"]
    hidden_root = layout.root / ".legacy-zero-missing"
    os.replace(missing_root, hidden_root)
    try:
        blocked_preview = await env.client.post(
            f"/api/projects/{project_id}/capabilities/skill-backfill",
            json={"action": "dry_run"},
        )
        assert blocked_preview.status_code == 200, blocked_preview.text
        assert blocked_preview.json()["blocked_count"] == 1
        rejected_apply = await env.client.post(
            f"/api/projects/{project_id}/capabilities/skill-backfill",
            json={"action": "apply"},
        )
        assert rejected_apply.status_code == 409
        await env.db.refresh(one)
        assert one.config == {"legacy_note": "keep-one", "skill_asset": v1_metadata}
    finally:
        os.replace(hidden_root, missing_root)


async def test_template_manifest_and_restore_include_legacy_members_effective_platform_dependencies(
    project_api: ProjectApiEnv,
):
    from app.models.mcp_server import MCPServer
    from app.models.project import ProjectCapabilityBinding, ProjectTemplate
    from app.models.tool import AgentTool

    env = project_api
    for agent in (env.leader, env.worker, env.reviewer):
        agent.autonomy_policy = {"write": "L2"}
    common_tool = Tool(
        name=f"template-common-{uuid.uuid4().hex[:8]}",
        display_name="Shared delivery checklist",
        description="Review the delivery checklist assigned to a project member.",
        type="builtin",
        category="project",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="builtin",
    )
    server = MCPServer(
        tenant_id=env.tenant_id,
        name=f"template-evidence-{uuid.uuid4().hex[:8]}",
        display_name="Evidence catalog",
        base_url_template="https://evidence.example.test/mcp",
    )
    env.db.add_all([common_tool, server])
    await env.db.flush()
    mcp_tool = Tool(
        name=f"template-mcp-{uuid.uuid4().hex[:8]}",
        display_name="Read evidence catalog",
        description="Read evidence available to the organization.",
        type="mcp",
        category="project",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="admin",
        tenant_id=env.tenant_id,
        mcp_server_id=server.id,
        mcp_server_name=server.display_name,
    )
    env.db.add(mcp_tool)
    await env.db.flush()
    env.db.add_all(
        [
            AgentTool(
                agent_id=agent_id,
                tool_id=common_tool.id,
                enabled=True,
                source="user_installed",
                config={"access_token": "must-not-enter-template"},
            )
            for agent_id in (env.leader_id, env.worker_id, env.reviewer_id)
        ]
        + [
            AgentTool(
                agent_id=env.leader_id,
                tool_id=mcp_tool.id,
                enabled=True,
                source="user_installed",
                config={"credential": "must-not-enter-template"},
            )
        ]
    )
    await env.db.commit()

    source_response = await env.client.post(
        "/api/projects",
        json={
            "name": "Legacy member template source",
            "goal": "Preserve the actual member setup",
            "members": [
                {
                    "agent_id": str(env.leader_id),
                    "is_leader": True,
                    "enabled_inherited_capability_ids": [str(common_tool.id), str(mcp_tool.id)],
                },
                {
                    "agent_id": str(env.worker_id),
                    "enabled_inherited_capability_ids": [str(common_tool.id)],
                },
                {
                    "agent_id": str(env.reviewer_id),
                    "enabled_inherited_capability_ids": [str(common_tool.id)],
                },
            ],
            "capabilities": [],
        },
    )
    assert source_response.status_code == 201, source_response.text
    source_project_id = source_response.json()["id"]
    members_response = await env.client.get(f"/api/projects/{source_project_id}/members")
    assert members_response.status_code == 200, members_response.text
    member_agent_ids = {
        item["name_snapshot"]: item["agent_id"] for item in members_response.json()
    }
    reviewer_member = next(
        item for item in members_response.json() if item["name_snapshot"] == env.reviewer.name
    )
    disable_response = await env.client.put(
        f"/api/projects/{source_project_id}/members/{reviewer_member['id']}/tools",
        json=[{"tool_id": str(common_tool.id), "enabled": False}],
    )
    assert disable_response.status_code == 200, disable_response.text
    disabled_common = next(item for item in disable_response.json() if item["id"] == str(common_tool.id))
    assert disabled_common["enabled"] is False

    manifest_response = await env.client.get(f"/api/projects/{source_project_id}/template-manifest")
    assert manifest_response.status_code == 200, manifest_response.text
    manifest = manifest_response.json()
    assert manifest["asset_summary"]["digital_employee_count"] == 3
    assert len(manifest["roles"]) == 3
    assert manifest["asset_summary"]["capability_count"] == 2
    common_manifest = next(
        item for item in manifest["capabilities"] if item["capability_id"] == str(common_tool.id)
    )
    assert common_manifest["key"] == common_tool.name
    assert common_manifest["description"] == common_tool.description
    assert common_manifest["availability"] == "available"
    assert common_manifest["affected_member_count"] == 2
    assert {item["agent_id"] for item in common_manifest["affected_members"]} == {
        member_agent_ids["Leader"],
        member_agent_ids["Worker"],
    }
    mcp_manifest = next(
        item for item in manifest["capabilities"] if item["capability_id"] == str(server.id)
    )
    assert mcp_manifest["key"] == server.name
    assert mcp_manifest["affected_member_count"] == 1
    assert mcp_manifest["affected_members"][0]["agent_id"] == member_agent_ids["Leader"]
    assert "must-not-enter-template" not in json.dumps(manifest)

    template_response = await env.client.post(
        f"/api/projects/{source_project_id}/templates",
        json={"name": "Legacy member setup", "version": "1.0.0"},
    )
    assert template_response.status_code == 201, template_response.text
    stored_template = await env.db.get(ProjectTemplate, uuid.UUID(template_response.json()["id"]))
    assert stored_template is not None
    assert len(stored_template.definition["agents"]) == 3
    assert len(stored_template.definition["capabilities"]) == 4
    assert "must-not-enter-template" not in json.dumps(stored_template.definition)

    restored_response = await env.client.post(
        "/api/projects/from-template",
        json={
            "template_id": str(stored_template.id),
            "name": "Legacy member template target",
            "visibility": "private",
        },
    )
    assert restored_response.status_code == 201, restored_response.text
    restored = restored_response.json()
    assert restored["template_setup_summary"]["restored_digital_employee_count"] == 3
    assert restored["template_setup_summary"]["restored_tool_count"] == 1
    assert restored["template_setup_summary"]["restored_connection_count"] == 1
    target_project_id = uuid.UUID(restored["id"])
    target_agents = (await env.client.get(f"/api/projects/{target_project_id}/agents")).json()
    assert len(target_agents) == 3
    assert {item["name"] for item in target_agents} == {"Leader", "Worker", "Reviewer"}
    target_agent_ids = {uuid.UUID(item["id"]) for item in target_agents}
    restored_bindings = list(
        (
            await env.db.execute(
                select(ProjectCapabilityBinding).where(
                    ProjectCapabilityBinding.project_id == target_project_id,
                )
            )
        ).scalars()
    )
    assert len(restored_bindings) == 4
    restored_assignments = list(
        (
            await env.db.execute(select(AgentTool).where(AgentTool.agent_id.in_(target_agent_ids)))
        ).scalars()
    )
    restored_dependency_assignments = [
        item for item in restored_assignments if item.tool_id in {common_tool.id, mcp_tool.id}
    ]
    assert len(restored_dependency_assignments) == 4
    assert sum(item.enabled for item in restored_dependency_assignments) == 3
    assert all(item.config in ({}, None) for item in restored_dependency_assignments)


async def test_template_export_preserves_member_override_of_shared_tool(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectTemplate
    from app.models.tool import AgentTool

    env = project_api
    shared_tool = Tool(
        name=f"template-shared-{uuid.uuid4().hex[:8]}",
        display_name="Shared project tool",
        description="A shared tool with a member-level override.",
        type="builtin",
        category="project",
        parameters_schema={"type": "object", "properties": {}},
        enabled=True,
        source="builtin",
    )
    env.db.add(shared_tool)
    await env.db.commit()
    project = await _create_project(env, name="Shared tool override source")
    project_id = project["id"]
    capability_response = await env.client.post(
        f"/api/projects/{project_id}/capabilities",
        json={
            "capability_type": "tool",
            "capability_id": str(shared_tool.id),
            "capability_name": shared_tool.display_name,
            "source": "shared",
            "is_enabled": True,
        },
    )
    assert capability_response.status_code == 201, capability_response.text
    members = (await env.client.get(f"/api/projects/{project_id}/members")).json()
    disabled_member = members[-1]
    disable_response = await env.client.put(
        f"/api/projects/{project_id}/members/{disabled_member['id']}/tools",
        json=[{"tool_id": str(shared_tool.id), "enabled": False}],
    )
    assert disable_response.status_code == 200, disable_response.text
    enable_response = await env.client.put(
        f"/api/projects/{project_id}/members/{disabled_member['id']}/tools",
        json=[{"tool_id": str(shared_tool.id), "enabled": True}],
    )
    assert enable_response.status_code == 200, enable_response.text
    disable_response = await env.client.put(
        f"/api/projects/{project_id}/members/{disabled_member['id']}/tools",
        json=[{"tool_id": str(shared_tool.id), "enabled": False}],
    )
    assert disable_response.status_code == 200, disable_response.text

    template_response = await env.client.post(
        f"/api/projects/{project_id}/templates",
        json={"name": "Shared tool override template"},
    )
    assert template_response.status_code == 201, template_response.text
    template = await env.db.get(ProjectTemplate, uuid.UUID(template_response.json()["id"]))
    assert template is not None
    exported = [
        item
        for item in template.definition["capabilities"]
        if item["capability_type"] == "tool" and item["capability_id"] == str(shared_tool.id)
    ]
    assert len(exported) == len(members)
    assert {item["source"] for item in exported} == {"inherited"}
    assert sum(bool(item["is_enabled"]) for item in exported) == len(members) - 1

    restored_response = await env.client.post(
        "/api/projects/from-template",
        json={"template_id": str(template.id), "name": "Shared tool override target"},
    )
    assert restored_response.status_code == 201, restored_response.text
    restored_agents = (await env.client.get(f"/api/projects/{restored_response.json()['id']}/agents")).json()
    restored_assignments = list(
        (
            await env.db.execute(
                select(AgentTool).where(
                    AgentTool.agent_id.in_([uuid.UUID(item["id"]) for item in restored_agents]),
                    AgentTool.tool_id == shared_tool.id,
                )
            )
        ).scalars()
    )
    assert len(restored_assignments) == len(members)
    assert sum(item.enabled for item in restored_assignments) == len(members) - 1


@pytest.mark.parametrize("invalid_agents", [None, {}, "not-a-list"])
async def test_generic_project_template_rejects_non_list_agent_assets(
    project_api: ProjectApiEnv,
    invalid_agents: object,
):
    response = await project_api.client.post(
        "/api/projects/templates",
        json={
            "name": "Invalid project Agent template",
            "definition": {"agents": invalid_agents},
        },
    )
    assert response.status_code == 422, response.text
    assert response.json()["detail"] == "Project template Agents must be a list"


async def test_generic_project_template_persists_only_sanitized_agent_assets(
    project_api: ProjectApiEnv,
):
    from app.models.project import ProjectTemplate

    secret_uuid = uuid.uuid4()
    secret_email = "operator@example.test"
    secret_token = "super-secret-bearer-token"
    secret_password = "database-password"
    response = await project_api.client.post(
        "/api/projects/templates",
        json={
            "name": "Sanitized project Agent template",
            "definition": {
                "goal": "Safe reusable workflow",
                "agents": [
                    {
                        "name": "Release operator",
                        "role_description": f"Coordinate with {secret_email}",
                        "soul": f"Identity {secret_uuid}; password={secret_password}",
                        "core_memory": f"Authorization: Bearer {secret_token}",
                        "workspace_files": [
                            {
                                "path": "release/notes.md",
                                "content": f"api_key={secret_token}\nowner={secret_email}\n",
                            }
                        ],
                    }
                ],
            },
        },
    )
    assert response.status_code == 201, response.text
    definition = response.json()["definition"]
    serialized_definition = json.dumps(definition)
    for secret in (str(secret_uuid), secret_email, secret_token, secret_password):
        assert secret not in serialized_definition
    assert "[redacted-id]" in serialized_definition
    assert "[redacted-email]" in serialized_definition
    assert "[redacted]" in serialized_definition

    stored = await project_api.db.get(ProjectTemplate, uuid.UUID(response.json()["id"]))
    assert stored is not None
    assert stored.definition == definition


async def test_template_market_admin_visibility_and_cross_tenant_management_boundaries(
    project_api: ProjectApiEnv,
):
    env = project_api
    source_project = await _create_project(env, name="Template boundary source")
    company_admin = await _user(
        env.db,
        await env.db.get(Tenant, env.tenant_id),
        "CompanyAdmin",
    )
    company_admin.role = "org_admin"
    platform_admin = await _user(
        env.db,
        await env.db.get(Tenant, env.tenant_id),
        "PlatformAdmin",
    )
    platform_admin.role = "platform_admin"
    other_tenant = Tenant(name="Other company", slug=f"other-{uuid.uuid4().hex[:8]}")
    env.db.add(other_tenant)
    await env.db.flush()
    other_creator = await _user(env.db, other_tenant, "OtherCreator")
    current_private = ProjectTemplate(
        tenant_id=env.tenant_id,
        created_by_user_id=env.viewer_id,
        name="Company private template",
        is_published=False,
        definition={"goal": "Company-only draft"},
    )
    current_public = ProjectTemplate(
        tenant_id=env.tenant_id,
        created_by_user_id=env.owner_id,
        name="Company public template",
        is_published=True,
        definition={"goal": "Published by this company"},
    )
    other_public = ProjectTemplate(
        tenant_id=other_tenant.id,
        created_by_user_id=other_creator.id,
        name="Other public template",
        is_published=True,
        definition={"goal": "Published by another company"},
    )
    other_private = ProjectTemplate(
        tenant_id=other_tenant.id,
        created_by_user_id=other_creator.id,
        name="Other private template",
        is_published=False,
        definition={"goal": "Other company draft"},
    )
    env.db.add_all((current_private, current_public, other_public, other_private))
    await env.db.commit()
    current_private_id = current_private.id
    current_public_id = current_public.id
    other_public_id = other_public.id
    other_private_id = other_private.id
    company_admin_id = company_admin.id
    platform_admin_id = platform_admin.id

    env.authenticate_as(company_admin_id)
    admin_response = await env.client.get("/api/projects/templates")
    assert admin_response.status_code == 200, admin_response.text
    admin_templates = {uuid.UUID(item["id"]): item for item in admin_response.json()}
    assert current_private_id in admin_templates
    assert current_public_id in admin_templates
    assert other_public_id in admin_templates
    assert other_private_id not in admin_templates
    assert admin_templates[current_private_id]["can_edit"] is True
    assert admin_templates[current_private_id]["can_delete"] is True
    assert admin_templates[other_public_id]["can_edit"] is False
    assert admin_templates[other_public_id]["can_delete"] is False

    market_overwrite = await env.client.put(
        f"/api/projects/templates/{other_public_id}/from-project/{source_project['id']}"
    )
    assert market_overwrite.status_code == 404, market_overwrite.text
    company_overwrite = await env.client.put(
        f"/api/projects/templates/{current_public_id}/from-project/{source_project['id']}"
    )
    assert company_overwrite.status_code == 200, company_overwrite.text
    assert company_overwrite.json()["can_edit"] is True

    env.authenticate_as(env.viewer_id)
    creator_response = await env.client.get(f"/api/projects/templates/{current_private_id}")
    assert creator_response.status_code == 200, creator_response.text
    assert creator_response.json()["can_edit"] is True
    creator_delete = await env.client.delete(f"/api/projects/templates/{current_private_id}")
    assert creator_delete.status_code == 204, creator_delete.text
    assert await env.db.get(ProjectTemplate, current_private_id) is None

    env.authenticate_as(platform_admin_id)
    platform_response = await env.client.get("/api/projects/templates")
    assert platform_response.status_code == 200, platform_response.text
    platform_templates = {uuid.UUID(item["id"]): item for item in platform_response.json()}
    assert other_private_id in platform_templates
    assert platform_templates[other_private_id]["can_edit"] is True
    assert platform_templates[other_private_id]["can_delete"] is True
    platform_delete = await env.client.delete(f"/api/projects/templates/{other_private_id}")
    assert platform_delete.status_code == 204, platform_delete.text
    assert await env.db.get(ProjectTemplate, other_private_id) is None


async def test_template_editor_is_hidden_and_overwrite_uses_same_project_snapshot(
    project_api: ProjectApiEnv,
):
    env = project_api
    source = await _create_project(env, name="Template editor source")
    template_response = await env.client.post(
        f"/api/projects/{source['id']}/templates",
        json={
            "name": "Editable project template",
            "description": "Original description",
            "category": "operations",
            "version": "2.0.0",
            "is_published": False,
        },
    )
    assert template_response.status_code == 201, template_response.text
    template = template_response.json()
    template_id = uuid.UUID(template["id"])

    env.authenticate_as(env.viewer_id)
    denied_editor = await env.client.post(f"/api/projects/templates/{template_id}/editor")
    assert denied_editor.status_code == 404

    env.authenticate_as(env.owner_id)
    editor_response = await env.client.post(f"/api/projects/templates/{template_id}/editor")
    assert editor_response.status_code == 201, editor_response.text
    editor = editor_response.json()
    editor_id = uuid.UUID(editor["id"])
    assert editor["settings"]["template_editor"] == {"template_id": str(template_id)}

    ordinary_projects = await env.client.get("/api/projects", params={"scope": "all"})
    assert ordinary_projects.status_code == 200, ordinary_projects.text
    assert editor_id not in {uuid.UUID(item["id"]) for item in ordinary_projects.json()}
    direct_editor = await env.client.get(f"/api/projects/{editor_id}")
    assert direct_editor.status_code == 200, direct_editor.text

    changed = await env.client.patch(
        f"/api/projects/{editor_id}",
        json={"goal": "Updated through the template editor project"},
    )
    assert changed.status_code == 200, changed.text
    overwrite = await env.client.put(
        f"/api/projects/templates/{template_id}/from-project/{editor_id}"
    )
    assert overwrite.status_code == 200, overwrite.text
    updated_template = overwrite.json()
    assert updated_template["name"] == "Editable project template"
    assert updated_template["description"] == "Original description"
    assert updated_template["category"] == "operations"
    assert updated_template["version"] == "2.0.0"
    assert updated_template["is_published"] is False
    assert updated_template["definition"]["goal"] == "Updated through the template editor project"
    assert "template_editor" not in updated_template["definition"].get("settings", {})

    deleted = await env.client.delete(f"/api/projects/templates/{template_id}")
    assert deleted.status_code == 204, deleted.text
    env.db.expire_all()
    persisted_editor = await env.db.get(Project, editor_id)
    assert persisted_editor is not None
    assert persisted_editor.template_id is None
    assert "template_editor" not in persisted_editor.settings
    ordinary_projects = await env.client.get("/api/projects", params={"scope": "all"})
    assert editor_id in {uuid.UUID(item["id"]) for item in ordinary_projects.json()}


async def test_template_editor_normalizes_structured_roles_into_digital_employees(
    project_api: ProjectApiEnv,
):
    env = project_api
    template_response = await env.client.post(
        "/api/projects/templates",
        json={
            "name": "Structured role template",
            "definition": {
                "roles": [
                    {
                        "key": "research_lead",
                        "name": "研究负责人",
                        "description": "统筹研究范围、证据质量与交付结论。",
                    },
                    {
                        "key": "fact_reviewer",
                        "name": "事实核查员",
                        "description": "核验关键事实与引用来源。",
                    },
                ]
            },
        },
    )
    assert template_response.status_code == 201, template_response.text
    template_id = template_response.json()["id"]

    editor_response = await env.client.post(f"/api/projects/templates/{template_id}/editor")
    assert editor_response.status_code == 201, editor_response.text
    editor_id = editor_response.json()["id"]
    members_response = await env.client.get(f"/api/projects/{editor_id}/members")
    assert members_response.status_code == 200, members_response.text
    members = members_response.json()
    assert [(item["name_snapshot"], item["role_snapshot"]) for item in members] == [
        ("研究负责人", "统筹研究范围、证据质量与交付结论。"),
        ("事实核查员", "核验关键事实与引用来源。"),
    ]
    assert members[0]["is_leader"] is True

    invalid_response = await env.client.post(
        "/api/projects/templates",
        json={"name": "Invalid role template", "definition": {"roles": "负责人"}},
    )
    assert invalid_response.status_code == 422
    assert invalid_response.json()["detail"] == "项目模板角色配置必须为列表。"
