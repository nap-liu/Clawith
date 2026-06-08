"""read_document must NOT silently truncate — oversized docs overflow to a file
through the shared tool_output_store, exactly like every other large tool output.

Root cause (round 2): read_document hard-truncated its own output at 8000 chars
INSIDE the tool and appended "[truncated]" with NO file reference — bypassing the
platform's overflow-to-file safety net (`finalize_tool_output`). The agent could
not retrieve the rest: read_document's schema doesn't even expose `max_chars`, so
there was no escape hatch. The data was destroyed at the tool layer before the
materialize-to-file path ever saw it.

Fix (root cause): return the full extracted text (bounded only by the structural
caps — file size, page/row/cell limits); let `finalize_tool_output` materialize
oversized output to `.tool_results/` with a preview + `read_file` pointer.
"""

from __future__ import annotations

import uuid
from pathlib import Path

from app.services.agent_tools import _read_document_sync
from app.services.llm import tool_output_store as tos


def _make_wide_xlsx(path: Path, n_rows: int, sql_cols: int = 35) -> list[str]:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(["模块", "子模块", "指标", "口径", "限制", "BI", "备注", "数据集查询语句"], 1):
        ws.cell(1, c, h)
    sqls = []
    for r in range(2, 2 + n_rows):
        sql = "SELECT " + ", ".join(f"t.col{i}" for i in range(sql_cols)) + f" FROM yk.sched WHERE rank={r}"
        sqls.append(sql)
        ws.cell(r, 1, "当日数据")
        ws.cell(r, 8, sql)
    wb.save(path)
    return sqls


def test_read_document_returns_full_content_no_truncation(tmp_path):
    # 30 rows of long SQL → real content exceeds the old 8000-char cap even after
    # trailing-empty trimming, so the old code would have truncated mid-document.
    sqls = _make_wide_xlsx(tmp_path / "doc.xlsx", n_rows=30)
    out = _read_document_sync(tmp_path, "doc.xlsx")
    assert len(out) > 8000, "should return far more than the old 8000 cap"
    assert "[truncated" not in out, "read_document must not truncate at the tool layer"
    for sql in sqls:
        assert sql in out, "every SQL block must survive in full"


def test_read_document_budget_is_40k():
    assert tos.budget_for("read_document") == 40_000


def test_oversized_read_document_overflows_to_file(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_DATA_DIR", str(tmp_path))
    from app.config import get_settings

    get_settings.cache_clear()
    try:
        agent_id = str(uuid.uuid4())
        # > 40K (read_document budget) but well under the 100K default — proves the
        # per-tool budget, not the default, drives materialization.
        big = "数据集查询语句\n" + ("SELECT * FROM yk.activity WHERE k=1;\n" * 2000)
        assert 40_000 < len(big) < 100_000
        view = tos.finalize_tool_output(
            big,
            tool_name="read_document",
            agent_id=agent_id,
            session_id="s1",
            tool_call_id="tc1",
        )
        # The LLM sees a recoverable pointer, NOT a dead-end "[truncated]".
        assert "<persisted-output>" in view
        assert "read_file" in view
        assert ".tool_results/" in view
        # The full content is on disk and byte-for-byte recoverable.
        store = tmp_path / agent_id / ".tool_results" / "s1"
        files = list(store.glob("read_document_*.txt"))
        assert len(files) == 1
        assert files[0].read_text(encoding="utf-8") == big
    finally:
        get_settings.cache_clear()
